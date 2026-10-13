"""Genera la receta del caso integral (contrato v1) y su plantilla de Excel: vacía (lo que el usuario debe llenar) y llena con
los datos sintéticos del caso. Las corridas de los diseños se escriben en la plantilla con el mismo orden que D-MIKE creará
(semillas 7 y 11), para que el usuario solo capture las respuestas."""
import csv, json, sys, datetime
from pathlib import Path
RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ)); sys.path.insert(0, str(RAIZ / "herramientas"))
from app.api import Api, VERSION   # noqa: E402
from generar_ejemplos import perfil  # noqa: E402

CASO = sys.argv[1] if len(sys.argv) > 1 else "/tmp/caso/caso_integral_uniones_adhesivas.dmike"
DATOS = Path(sys.argv[2] if len(sys.argv) > 2 else "/tmp/caso/datos")
SALIDA = Path(sys.argv[3] if len(sys.argv) > 3 else "/tmp/receta")
SALIDA.mkdir(parents=True, exist_ok=True)
MODULOS = [("CHARTER", "D", "N"), ("SIPOC", "D", "N"), ("KANO", "D", "R"), ("QFD", "D", "R"), ("CTQ-TREE", "D", "N"), ("DEF-PROBLEMA", "D", "N"), ("INTERESADOS", "D", "R"),
           ("HOJA-DATOS", "M", "N"), ("MSA-VAR", "M", "N"), ("DISTRIB", "M", "N"), ("CAPACIDAD", "M", "N"), ("SPC", "C", "N"), ("DESEMPENO", "M", "R"),
           ("PARETO", "A", "N"), ("ISHIKAWA", "A", "N"), ("CINCO-PQ", "A", "R"), ("ANOVA", "A", "N"), ("CORREL", "A", "R"), ("REGRESION", "A", "R"),
           ("RSM", "I", "N"), ("AMEF", "I", "N"), ("POKAYOKE", "I", "R"), ("PLAN-CONTROL", "C", "N"), ("TABLERO", "C", "N"), ("CHECKLIST-FASE", "C", "R"), ("A3", "C", "R")]
COMPL = [("RACI", "D"), ("EDT", "D"), ("CRONOGRAMA", "D"), ("MONTECARLO", "D"), ("RIESGOS", "D"), ("JUSTIFICACION", "D"), ("BITACORA", "D"), ("ENFOQUE", "D"),
         ("HIPOTESIS", "A"), ("DISENO-RSM", "I"), ("FACTORIAL", "I"), ("OPTIMIZADOR", "I"), ("VALOR-GANADO", "C"), ("COSTO-CALIDAD", "C"), ("LECCIONES", "C")]
COL = lambda n, t, u, d, ob=True: {"nombre": n, "tipo": t, "unidad": u, "descripcion": d, "obligatoria": ob}
HOJAS = [
    {"hoja": "linea_base", "titulo": "Situación actual (Medir)", "columnas": [
        COL("fecha", "fecha", "", "Día de la prueba"), COL("turno", "texto", "", "Turno A, B o C"), COL("lote_adhesivo", "texto", "", "Número de lote del adhesivo"),
        COL("temp_real", "num", "°C", "Temperatura real de curado del panel"), COL("tiempo_real", "num", "min", "Tiempo real de curado"),
        COL("resistencia_base", "num", "MPa", "Resistencia al cortante de la probeta"), COL("porosidad_base", "num", "%", "Porosidad de la unión")]},
    {"hoja": "defectos", "titulo": "Registro de defectos (Analizar)", "columnas": [COL("tipo_defecto", "texto", "", "Un renglón por defecto encontrado")]},
    {"hoja": "gage_rr", "titulo": "Estudio del sistema de medición (Medir)", "columnas": [
        COL("parte_rr", "texto", "", "Probeta medida"), COL("operador_rr", "texto", "", "Quién midió"), COL("replica_rr", "num", "", "1 o 2"),
        COL("medicion_rr", "num", "MPa", "Resultado de la medición")]},
    {"hoja": "cribado", "titulo": "Experimento de cribado (Mejorar)", "columnas": [
        COL("Temperatura", "num", "°C", "Ya viene escrita: no se cambia"), COL("Tiempo", "num", "min", "Ya viene escrita"), COL("Presion", "num", "bar", "Ya viene escrita"),
        COL("Limpieza", "num", "nivel", "Ya viene escrita"), COL("Resistencia_cribado", "num", "MPa", "Resultado de cada corrida")]},
    {"hoja": "dcc", "titulo": "Diseño central compuesto (Mejorar)", "columnas": [
        COL("Temp_curado", "num", "°C", "Ya viene escrita"), COL("Tiempo_curado", "num", "min", "Ya viene escrita"),
        COL("Resistencia", "num", "MPa", "Resultado de cada corrida"), COL("Porosidad", "num", "%", "Resultado de cada corrida")]},
    {"hoja": "despues", "titulo": "Proceso mejorado (Controlar)", "columnas": [COL("fecha_despues", "fecha", "", "Día de la prueba"), COL("resistencia_despues", "num", "MPa", "Resistencia con el ciclo nuevo")]},
]
a = Api(); a.abrir_proyecto(CASO)
P = a.proyecto
res = [a.resultado(r["uid"]) for r in a.resultados()]
disenos, analisis, ids = [], [], {}
nombres = {"DISENO-RSM": ["cribado", "dcc"], "FACTORIAL": ["factorial"], "RSM": ["rsm_resistencia", "rsm_porosidad"], "OPTIMIZADOR": ["optimizador"],
           "MSA-VAR": ["gage_rr"], "DISTRIB": ["normalidad"], "CAPACIDAD": ["capacidad_antes", "capacidad_despues"], "SPC": ["carta_antes", "carta_despues"],
           "PARETO": ["pareto"], "ANOVA": ["anova_lote"], "CORREL": ["correlacion"], "REGRESION": ["regresion"], "HIPOTESIS": ["prueba_t"]}
uid_a_id = {}
for r in res:
    d = r["definicion"]; m = d["analisis"]; cnt = sum(1 for x in disenos + analisis if x["modulo"] == m)
    rid = nombres.get(m, [m.lower()])[min(cnt, len(nombres.get(m, [m])) - 1)]
    uid_a_id[r["uid"]] = rid
    params = dict(d["parametros"])
    for k, v in (d.get("columnas") or {}).items(): params[k] = [f"@{x}" for x in v] if isinstance(v, list) else f"@{v}"
    for k, v in list(params.items()):
        if isinstance(v, str) and v in uid_a_id: params[k] = "#" + uid_a_id[v]
    if m == "DISENO-RSM":
        disenos.append({"id": rid, "modulo": m, "parametros": params, "hoja": rid, "respuestas": ["Resistencia_cribado"] if rid == "cribado" else ["Resistencia", "Porosidad"]})
    else:
        analisis.append({"id": rid, "modulo": m, "parametros": params, "si_faltan_datos": "dejar_pendiente"})
formatos = [{"id": fid, "cabecera": d.get("cabecera", {}), "tablas": d.get("tablas", {})} for fid, d in P.formatos.items() if not fid.startswith("_")]
pf = perfil(P.fijos["nombre_proyecto"]["valor"], P.fijos["cliente_area"]["valor"], P.fijos["responsable"]["valor"], "C", MODULOS)
receta = {"tipo": "dmike-receta", "receta_version": 1, "dmike_minimo": "0.23", "generado_en": datetime.datetime.now().isoformat(timespec="seconds"),
          "origen": "Caso maestro de D-MIKE: uniones adhesivas (datos sintéticos)", "arbol_version": pf.get("arbol_version", "1.3"),
          "datos_fijos": {k: P.fijos[k]["valor"] for k in ("nombre_proyecto", "cliente_area", "responsable", "notas")},
          "perfil": pf, "hoja_sugerida": False, "complementarios": COMPL, "hojas": HOJAS, "disenos": disenos, "formatos": formatos, "analisis": analisis,
          "avance": {mid: m["avance"] for mid, m in P.modulos.items() if m.get("avance") and m["avance"] != "pendiente"},
          "recomendaciones": ["Mide primero el sistema de medición (Gage R&R) y la estabilidad antes de calcular la capacidad.",
                              "Corre los experimentos en el orden de la plantilla (columna corrida) y no cambies los niveles de los factores."]}

# ---------- plantilla de Excel (vacía y llena) ----------
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from app.modulos import doe as D
def leer_csv(n): return list(csv.reader(open(DATOS / n, encoding="utf-8-sig")))
CSV = {"linea_base": "01_linea_base.csv", "defectos": "02_defectos.csv", "gage_rr": "03_gage_rr.csv", "cribado": "04_cribado_fraccionado.csv", "dcc": "05_diseno_central_compuesto.csv", "despues": "06_despues.csv"}
# corridas que D-MIKE creará con las mismas semillas (para escribir los factores en la plantilla)
aux = Api()
corr = {}
for d in disenos:
    r = aux.ejecutar_modulo("DISENO-RSM", d["parametros"]); inf = r["resultado"]["datos"]["diseno"]
    fs = [f["nombre"] for f in inf["factores"]]; js = [aux.hoja.indice(n) for n in fs]
    corr[d["id"]] = [[aux.hoja.columnas[j].valor(i)[2] for j in js] for i in range(aux.hoja.n) if aux.hoja.columnas[js[0]].valor(i)[2] is not None]
for h in HOJAS:
    if h["hoja"] in corr: h["corridas"] = corr[h["hoja"]]          # factores ya escritos en la plantilla (mismo orden que creará D-MIKE)
json.dump(receta, open(SALIDA / "receta_caso_integral.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
NAR = PatternFill("solid", fgColor="DD5813"); CREMA = PatternFill("solid", fgColor="FBE8DD"); GRIS = PatternFill("solid", fgColor="F1ECE2")
def libro(lleno):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "LEEME"
    ws["A1"] = "D-MIKE · Plantilla de datos de la receta"; ws["A1"].font = Font(bold=True, size=14, color="B3460D")
    filas = ["Proyecto: " + (receta["datos_fijos"]["nombre_proyecto"] or ""), "Llena cada hoja con tus datos, una fila por medición. No cambies los nombres de las columnas.",
             "Celdas grises: ya vienen escritas (corridas de los experimentos). Celdas color crema: las llenas tú.", "Deja vacío lo que no tengas: D-MIKE lo marcará como pendiente.",
             "Cuando termines, abre este archivo en D-MIKE (Archivo › Abrir proyecto) y confirma «Ejecutar la receta».", ""]
    for k, t in enumerate(filas): ws.cell(3 + k, 1, t)
    r0 = 3 + len(filas); ws.cell(r0, 1, "Hoja"); ws.cell(r0, 2, "Columna"); ws.cell(r0, 3, "Tipo"); ws.cell(r0, 4, "Unidad"); ws.cell(r0, 5, "Para qué sirve")
    for c in range(1, 6): ws.cell(r0, c).font = Font(bold=True, color="FFFFFF"); ws.cell(r0, c).fill = NAR
    k = r0
    for h in HOJAS:
        for c in h["columnas"]:
            k += 1; ws.cell(k, 1, h["hoja"]); ws.cell(k, 2, c["nombre"]); ws.cell(k, 3, c["tipo"]); ws.cell(k, 4, c["unidad"]); ws.cell(k, 5, c["descripcion"])
    for col, w in zip("ABCDE", (16, 22, 9, 9, 60)): ws.column_dimensions[col].width = w
    for h in HOJAS:
        wh = wb.create_sheet(h["hoja"]); cols = [c["nombre"] for c in h["columnas"]]
        for j, n in enumerate(cols, 1):
            c = wh.cell(1, j, n); c.font = Font(bold=True, color="FFFFFF"); c.fill = NAR; wh.column_dimensions[c.column_letter].width = max(14, len(n) + 4)
        datos_csv = leer_csv(CSV[h["hoja"]])[1:]
        if h["hoja"] in corr:
            nf = len(corr[h["hoja"]][0])
            for i, fila in enumerate(corr[h["hoja"]]):
                for j, v in enumerate(fila, 1): c = wh.cell(i + 2, j, v); c.fill = GRIS
                for j in range(nf + 1, len(cols) + 1):
                    c = wh.cell(i + 2, j, float(datos_csv[i][j - 1]) if lleno else None); c.fill = CREMA
        elif lleno:
            for i, f in enumerate(datos_csv):
                for j, v in enumerate(f, 1):
                    try: v2 = float(v) if h["columnas"][j - 1]["tipo"] == "num" else v
                    except ValueError: v2 = v
                    wh.cell(i + 2, j, v2)
        wh.freeze_panes = "A2"
    wr = wb.create_sheet("_receta"); texto = json.dumps(receta, ensure_ascii=False)
    for i in range(0, len(texto), 30000): wr.cell(i // 30000 + 1, 1, texto[i:i + 30000])
    wr.sheet_state = "hidden"
    return wb
libro(False).save(SALIDA / "caso_integral_plantilla.xlsx"); libro(True).save(SALIDA / "caso_integral_datos_llenos.xlsx")
print(len(disenos), "diseños ·", len(analisis), "análisis ·", len(formatos), "formatos ·", len(json.dumps(receta)) // 1024, "KB")
