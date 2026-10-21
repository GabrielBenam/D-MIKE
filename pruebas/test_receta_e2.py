"""D-MIKE 0.27, Entrega 2 (ESPEC_receta_v2.md secciones 12 a 15): variables, tareas, estados de análisis, revisiones y paquete de avance (E2-R1 a E2-R10).
Datos SINTÉTICOS: las recetas se arman aquí con valores inventados; el caso maestro es el de siempre."""
import copy, gzip, hashlib, json, os, shutil, subprocess, sys
from pathlib import Path
import openpyxl, pytest
from app.api import Api, VERSION
from app.nucleo import archivo, menus, modo as MD, receta as RC, seguimiento as SG
from app import modulos

S = Path(__file__).parent / "sinteticos"
REG = modulos.cargar()
UUID_A, UUID_B = "3f2b8c1e-5d4a-4f6b-9a7c-1e2d3c4b5a69", "9b1deb4d-3b7d-4bad-9bdd-2b0d7b3dcb6d"
DATOS_ESPESOR = [[0.512, "L1"], [0.498, "L1"], [0.507, "L2"], [0.521, "L2"], [0.503, "L3"], [0.5, "L3"]]       # sintéticos


@pytest.fixture(autouse=True)
def _recientes_aislados(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))


# ---------- ayudas ----------
def rec(revision=None, rev_base=None, modo="real", idp=UUID_A, **extra):
    """Receta v2 mínima de un proyecto real: una hoja vacía con dos columnas, una variable pendiente de recolectar y un análisis que necesita Espesor."""
    rc = {"tipo": "dmike-receta", "receta_version": 2, "dmike_minimo": "0.27", "modo_proyecto": modo, "id_proyecto": idp, "origen": "Diagnóstico en línea",
          "generado_en": "2026-10-14T00:00:00", "datos_fijos": {"nombre_proyecto": "Proyecto E2"},
          "hojas": [{"hoja": "mediciones", "origen_datos": "vacia", "columnas": [{"nombre": "Espesor", "tipo": "num", "unidad": "mm"}, {"nombre": "Lote", "tipo": "texto"}]}],
          "variables": [{"nombre": "Espesor", "hoja": "mediciones", "unidad": "mm", "unidad_analisis": "panel", "fuente": "Calibrador", "responsable": "Calidad",
                         "estado_evidencia": "PENDIENTE_RECOLECCION", "ensayo_destructivo": False}],
          "analisis": [{"id": "estad", "modulo": "ESTADISTICOS", "parametros": {"y": "@Espesor"}, "requiere": ["@Espesor"], "estado": "pendiente"}], "tareas": []}
    if revision: rc["revision"] = revision
    if rev_base is not None: rc["revision_base"] = rev_base
    rc.update(extra); rc["huella"] = RC.huella_receta(rc)
    return rc


def json_(tmp, rc, nombre="r.json"):
    ruta = tmp / nombre; ruta.write_text(json.dumps(rc, ensure_ascii=False), encoding="utf-8"); return str(ruta)


def libro(tmp, rc, datos, nombre="r.xlsx"):
    """Plantilla de Excel como la que arma el diagnóstico: hojas de datos y la hoja oculta «_receta» en trozos."""
    wb = openpyxl.Workbook(); wb.active.title = "LEEME"
    for h, filas in datos.items():
        w = wb.create_sheet(h)
        for f in filas: w.append(f)
    wr = wb.create_sheet("_receta"); t = json.dumps(rc, ensure_ascii=False)
    for i in range(0, len(t), 30000): wr.cell(i // 30000 + 1, 1, t[i:i + 30000])
    wr.sheet_state = "hidden"; ruta = tmp / nombre; wb.save(ruta); return str(ruta)


def destino(a): return a.proyectos()["proyectos"][0]["id"]
def act(a, ruta, **kw): return a.ejecutar_receta(ruta, "actualizar", destino(a), **kw)
def tarea(a, tid): return next((t for t in a.proyecto.formatos.get("_tareas", []) if t["id"] == tid), None)
def textos(r): return " | ".join(m["texto"] for m in r["sesion"])


# ---------- E2-R1: validación v2 ----------
def err(**campos):
    rc = rec(); rc.update(campos); rc.pop("huella"); return RC.validar(rc, REG, VERSION)


def test_r1_receta_valida_y_claves_desconocidas_toleradas():
    assert RC.validar(rec(), REG, VERSION) == []
    assert RC.validar(rec(clave_futura={"a": 1}, variables=[{"nombre": "x", "campo_nuevo": 7}], tareas=[{"id": "t", "descripcion": "d", "extra": 1}]), REG, VERSION) == []


def test_r1_variables():
    assert any("no tiene nombre" in e for e in err(variables=[{"nombre": ""}])) and any("no tiene nombre" in e for e in err(variables=[{"estado_evidencia": "DESCONOCIDO"}]))
    assert any("no tiene nombre" in e for e in err(variables=[{"nombre": 5}])) and any("debe ser un objeto" in e for e in err(variables=["x"]))
    assert any("variables debe ser una lista" in e for e in err(variables={"nombre": "x"}))
    assert any("ensayo_destructivo" in e and "«x»" in e for e in err(variables=[{"nombre": "x", "ensayo_destructivo": "si"}]))
    assert err(variables=[{"nombre": "x", "ensayo_destructivo": True, "estado_evidencia": "NO_APLICA"}]) == []
    assert any("estado_evidencia" in e for e in err(variables=[{"nombre": "x", "estado_evidencia": "INVENTADO"}]))


def test_r1_tareas():
    assert err(tareas=[{"id": "a", "estado": "abierta"}, {"id": "b", "estado": "cerrada"}, {"id": "c"}]) == []
    assert any("repetido" in e and "«a»" in e for e in err(tareas=[{"id": "a"}, {"id": "a"}]))
    assert any("no tiene id" in e for e in err(tareas=[{"descripcion": "x"}])) and any("no tiene id" in e for e in err(tareas=[{"id": "  "}]))
    assert any("estado" in e and "abierta o cerrada" in e for e in err(tareas=[{"id": "a", "estado": "hecha"}]))
    assert any("tareas debe ser una lista" in e for e in err(tareas="x"))


def test_r1_revision_y_base():
    assert err(revision=1) == [] and err(revision=3, revision_base=2) == [] and err(revision=3, revision_base=None) == [] and err(revision=2, revision_base=1) == []
    for malo in (0, -1, "2", 1.5, True):
        assert any("revision debe ser un entero" in e for e in err(revision=malo)), malo
    for base in (3, 4):
        assert any("revision_base" in e and "menor" in e for e in err(revision=3, revision_base=base)), base
    assert any("revision_base" in e for e in err(revision=3, revision_base="2")) and any("revision_base" in e for e in err(revision=3, revision_base=True))
    assert any("revision_base" in e and "revision" in e for e in err(revision_base=1))                       # base sin revisión


def test_r1_estado_y_requiere_de_analisis():
    ok = {"id": "a", "modulo": "ESTADISTICOS", "parametros": {"y": "@Espesor"}}
    for est in ("pendiente", "disponible", "analizado", "aprobado"): assert err(analisis=[dict(ok, estado=est, requiere=["@Espesor"])]) == [], est
    assert any("estado" in e and "«a»" in e for e in err(analisis=[dict(ok, estado="listo")]))
    for malo in ("Espesor", ["Espesor"], ["@"], [3], "@Espesor"):
        assert any("requiere" in e and "@" in e for e in err(analisis=[dict(ok, requiere=malo)])), malo
    assert err(analisis=[dict(ok, requiere=[])]) == []


# ---------- E2-R2: variables ----------
def test_r2_se_guardan_y_se_fusionan_por_nombre(tmp_path):
    a = Api(); assert a.ejecutar_receta(json_(tmp_path, rec(revision=1)))["error"] is None
    v = a.proyecto.formatos["_variables"]
    assert v == [{"nombre": "Espesor", "hoja": "mediciones", "unidad": "mm", "unidad_analisis": "panel", "fuente": "Calibrador", "responsable": "Calidad",
                  "estado_evidencia": "PENDIENTE_RECOLECCION", "ensayo_destructivo": False}]
    nuevas = [{"nombre": "Espesor", "estado_evidencia": "REAL_NO_VALIDADO", "responsable": "Metrología"}, {"nombre": "Dureza", "estado_evidencia": "DESCONOCIDO", "ensayo_destructivo": True}]
    rc2 = rec(revision=2, rev_base=1, variables=nuevas)
    assert act(a, json_(tmp_path, rc2, "r2.json"))["error"] is None
    por = {x["nombre"]: x for x in a.proyecto.formatos["_variables"]}
    assert por["Espesor"]["estado_evidencia"] == "REAL_NO_VALIDADO" and por["Espesor"]["responsable"] == "Metrología" and por["Espesor"]["fuente"] == "Calibrador"      # lo no enviado se conserva
    assert por["Dureza"]["ensayo_destructivo"] is True and len(por) == 2
    # una variable que la receta ya no menciona se conserva
    rc3 = rec(revision=3, rev_base=2, variables=[{"nombre": "Dureza", "estado_evidencia": "REAL_VERIFICADO"}])
    assert act(a, json_(tmp_path, rc3, "r3.json"))["error"] is None
    por = {x["nombre"]: x for x in a.proyecto.formatos["_variables"]}
    assert set(por) == {"Espesor", "Dureza"} and por["Dureza"]["estado_evidencia"] == "REAL_VERIFICADO" and por["Espesor"]["estado_evidencia"] == "REAL_NO_VALIDADO"


def test_r2_un_estado_verificado_nunca_baja_en_silencio(tmp_path):
    ver = [{"nombre": "Espesor", "estado_evidencia": "REAL_VERIFICADO"}]
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1, variables=ver)))
    rc2 = rec(revision=2, rev_base=1, variables=[{"nombre": "Espesor", "estado_evidencia": "REAL_NO_VALIDADO"}]); ruta = json_(tmp_path, rc2, "r2.json")
    antes = json.dumps(a.proyecto.a_dict(), sort_keys=True)
    r = act(a, ruta)                                                                   # con la revisión en orden, igual es un conflicto: el usuario decide
    assert r["error"] == "conflicto_revision" and r["revision"]["variables"] == [{"nombre": "Espesor", "actual": "REAL_VERIFICADO", "recibido": "REAL_NO_VALIDADO"}]
    assert json.dumps(a.proyecto.a_dict(), sort_keys=True) == antes                      # nada cambió
    assert act(a, ruta, resolucion="cancelar")["error"] == "cancelado" and json.dumps(a.proyecto.a_dict(), sort_keys=True) == antes
    r = act(a, ruta, resolucion="solo_nuevo"); assert r["error"] is None and "REAL_VERIFICADO" in textos(r)
    assert a.proyecto.formatos["_variables"][0]["estado_evidencia"] == "REAL_VERIFICADO" and a.proyecto.formatos["_receta"]["revision"] == 2
    # verificado -> verificado, y cualquier estado -> verificado, no son conflicto
    rc3 = rec(revision=3, rev_base=2, variables=[{"nombre": "Espesor", "estado_evidencia": "REAL_VERIFICADO", "unidad": "mm"}, {"nombre": "Otra", "estado_evidencia": "DESCONOCIDO"}])
    assert act(a, json_(tmp_path, rc3, "r3.json"))["error"] is None
    # la función de fusión por sí sola tampoco baja el estado
    p = Api().proyecto; SG.fusionar_variables(p, ver)
    assert "REAL_VERIFICADO" in SG.fusionar_variables(p, [{"nombre": "Espesor", "estado_evidencia": "DESCONOCIDO"}])[0] and p.formatos["_variables"][0]["estado_evidencia"] == "REAL_VERIFICADO"


# ---------- E2-R3: tareas ----------
def test_r3_pendientes_pasan_a_tareas_auto_con_el_mismo_texto(tmp_path):
    a = Api(); r = a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    pend = r["receta_resultado"]["pendientes"]
    assert "Llenar la hoja «mediciones» (Espesor, Lote)." in pend and any(p.startswith("Análisis «estad» (ESTADISTICOS): sin datos en Espesor") for p in pend)
    t = a.proyecto.formatos["_tareas"]; por = {x["id"]: x for x in t}
    assert set(por) == {"auto:hoja:mediciones", "auto:analisis:estad"} and all(x["origen"] == "auto" and x["estado"] == "abierta" for x in t)
    assert sorted(x["descripcion"] for x in t) == sorted(pend)                                   # el mismo texto del registro
    assert all(set(x) == {"id", "descripcion", "responsable", "accion", "liga", "estado", "origen"} for x in t)
    assert por["auto:hoja:mediciones"]["liga"] == "hoja:mediciones" and por["auto:analisis:estad"]["liga"] == "analisis:estad"
    assert [m["texto"] for m in r["sesion"] if m["texto"].startswith("Pendiente: ")] == ["Pendiente: " + p for p in pend]       # el registro de la sesión no cambia


def test_r3_las_auto_se_cierran_solas_al_llegar_los_datos_a6(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    rc2 = rec(revision=2, rev_base=1)
    ruta = libro(tmp_path, rc2, {"mediciones": [["Espesor", "Lote"]] + DATOS_ESPESOR})
    r = act(a, ruta); assert r["error"] is None, textos(r)
    assert all(x["estado"] == "cerrada" for x in a.proyecto.formatos["_tareas"]) and "Tareas cerradas porque ya se cumplieron: 2" in textos(r)
    e = a.proyecto.formatos["_analisis"]["estad"]; assert e["estado"] == "analizado" and e["uid"] in {x["uid"] for x in a.resultados()} and e["motivo"] == ""
    assert a.seguimiento()["analisis"][0]["estado"] == "analizado" and a.seguimiento()["tareas_abiertas"] == []


def test_r3_una_tarea_cerrada_no_se_reabre_y_las_de_la_receta_se_conservan(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1, tareas=[{"id": "t1", "descripcion": "Pedir el certificado del calibrador", "responsable": "Calidad", "accion": "Pedir", "liga": "variable:Espesor", "estado": "abierta"}])))
    assert tarea(a, "t1")["origen"] == "receta" and tarea(a, "t1")["estado"] == "abierta"
    tarea(a, "auto:hoja:mediciones")["estado"] = "cerrada"                                 # cerrada a mano
    r = act(a, json_(tmp_path, rec(revision=2, rev_base=1), "r2.json")); assert r["error"] is None            # la receta nueva ya no trae t1 y la hoja sigue vacía
    assert tarea(a, "auto:hoja:mediciones")["estado"] == "cerrada" and tarea(a, "t1")["estado"] == "abierta" and tarea(a, "t1")["origen"] == "receta"
    assert tarea(a, "auto:analisis:estad")["estado"] == "abierta"
    # la receta puede cerrar una tarea suya; después no la reabre aunque la vuelva a mandar abierta
    r3 = rec(revision=3, rev_base=2, tareas=[{"id": "t1", "descripcion": "Certificado recibido", "estado": "cerrada"}]); assert act(a, json_(tmp_path, r3, "r3.json"))["error"] is None
    assert tarea(a, "t1")["estado"] == "cerrada" and tarea(a, "t1")["descripcion"] == "Certificado recibido"
    r4 = rec(revision=4, rev_base=3, tareas=[{"id": "t1", "descripcion": "Certificado recibido", "estado": "abierta"}]); assert act(a, json_(tmp_path, r4, "r4.json"))["error"] is None
    assert tarea(a, "t1")["estado"] == "cerrada"


def test_r3_ids_auto_de_disenos_y_respuestas_en_el_caso_maestro():
    a = Api(); a.ejecutar_receta(str(S / "caso_integral_plantilla.xlsx"))
    ids = {t["id"] for t in a.proyecto.formatos["_tareas"]}
    assert {"auto:respuesta:cribado:Resistencia_cribado", "auto:respuesta:dcc:Resistencia", "auto:respuesta:dcc:Porosidad", "auto:hoja:linea_base", "auto:analisis:factorial"} <= ids
    destino_id = destino(a); r = a.ejecutar_receta(str(S / "caso_integral_datos_llenos.xlsx"), "actualizar", destino_id); assert r["error"] is None
    assert not [t for t in a.proyecto.formatos["_tareas"] if t["estado"] == "abierta"]
    assert all(e["estado"] == "analizado" for e in a.proyecto.formatos["_analisis"].values()) and len(a.proyecto.formatos["_analisis"]) == 15


# ---------- E2-R4: estado de cada análisis ----------
def test_r4_pendiente_con_motivo_y_sin_correr(tmp_path):
    a = Api(); r = a.ejecutar_receta(json_(tmp_path, rec(revision=1, analisis=[
        {"id": "estad", "modulo": "ESTADISTICOS", "parametros": {"y": "@Espesor"}, "requiere": ["@Espesor"]},
        {"id": "otro", "modulo": "ESTADISTICOS", "parametros": {"y": "@Espesor"}, "requiere": ["@Espesor", "@Dureza"]},      # Dureza ni existe en la hoja
        {"id": "lista", "modulo": "ESTADISTICOS", "parametros": {"y": "#estad"}}])))                                       # depende de un resultado previo
    an = a.proyecto.formatos["_analisis"]
    assert a.resultados() == [] and all(e["estado"] == "pendiente" for e in an.values())                                  # nunca se corrió sin datos
    assert an["estad"]["motivo"] == "Sin datos en Espesor." and an["estad"]["requiere"] == ["@Espesor"] and an["estad"]["uid"] is None
    assert "Dureza" in an["otro"]["motivo"] and an["otro"]["requiere"] == ["@Espesor", "@Dureza"]
    assert "estad" in an["lista"]["motivo"] and an["lista"]["motivo"].startswith("Falta el resultado")


def test_r4_disponible_si_hay_datos_y_analizado_solo_con_resultado(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    assert a.seguimiento()["analisis"][0]["estado"] == "pendiente"
    j = a.hoja.indice("Espesor"); a.pegar(0, j, "\n".join(str(f[0]) for f in DATOS_ESPESOR))                              # el usuario captura a mano
    s = a.seguimiento()["analisis"][0]; assert s["estado"] == "disponible" and "falta correr" in s["motivo"]
    assert a.proyecto.formatos["_analisis"]["estad"]["estado"] == "pendiente"                                            # lo guardado no cambia solo al leer
    # corre con la receta y queda analizado; si el resultado desaparece, ya no se declara analizado
    assert act(a, json_(tmp_path, rec(revision=2, rev_base=1), "r2.json"))["error"] is None
    assert a.seguimiento()["analisis"][0]["estado"] == "analizado"
    a.borrar_resultado(a.proyecto.formatos["_analisis"]["estad"]["uid"])
    assert a.seguimiento()["analisis"][0]["estado"] == "disponible"
    e = {"modulo": "X", "estado": "aprobado", "requiere": [], "motivo": "", "uid": "no-existe"}
    a.proyecto.formatos["_analisis"]["x"] = e
    assert next(x for x in a.seguimiento()["analisis"] if x["id"] == "x")["estado"] != "aprobado"                        # aprobado tampoco sin resultado


# ---------- E2-R5: revisiones ----------
def test_r5_misma_huella_ya_aplicada_a2(tmp_path):
    a = Api(); ruta = json_(tmp_path, rec(revision=1)); a.ejecutar_receta(ruta)
    antes = json.dumps(a.proyecto.a_dict(), sort_keys=True)
    r = act(a, ruta); assert r["error"] == "ya_aplicada" and "ya aplicada" in textos(r) and r["revision"]["accion"] == "ya_aplicada"
    assert json.dumps(a.proyecto.a_dict(), sort_keys=True) == antes and "_revisiones" not in a.proyecto.formatos
    # misma receta en la hoja de Excel de la plantilla vacía: igual
    assert act(a, libro(tmp_path, rec(revision=1), {"mediciones": [["Espesor", "Lote"]]}))["error"] == "ya_aplicada"
    # la preparación de la pantalla lo anuncia antes de ejecutar
    er = a.abrir_receta(ruta)["receta"]["estado_revision"]; assert er["accion"] == "ya_aplicada"


def test_r5_misma_receta_con_datos_nuevos_si_se_aplica(tmp_path):
    """Una plantilla con la misma receta pero con datos que el proyecto no tiene sigue siendo la forma de cargar las mediciones: no es «ya aplicada»."""
    a = Api(); rc = rec(revision=1); a.ejecutar_receta(json_(tmp_path, rc))
    llena = libro(tmp_path, rc, {"mediciones": [["Espesor", "Lote"]] + DATOS_ESPESOR})
    r = act(a, llena); assert r["error"] is None and a.proyecto.formatos["_analisis"]["estad"]["estado"] == "analizado"
    assert a.proyecto.formatos["_receta"]["revision"] == 1 and "_revisiones" not in a.proyecto.formatos           # misma receta: nada que archivar
    assert act(a, llena)["error"] == "ya_aplicada"                                                              # y ahora sí, ya aplicada
    n = a.hoja.n
    cambiada = libro(tmp_path, rc, {"mediciones": [["Espesor", "Lote"]] + [[9.9, "Z"]] * 6}, "c.xlsx")             # otros valores en columnas que ya tenían datos: no se sobrescriben
    assert act(a, cambiada)["error"] == "ya_aplicada" and a.hoja.n == n and a.hoja.columnas[a.hoja.indice("Espesor")].valor(0)[2] == 0.512


def test_r5_revision_mas_antigua_no_se_aplica_salvo_forzar_a3(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    assert act(a, json_(tmp_path, rec(revision=2, rev_base=1), "r2.json"))["error"] is None
    antigua = rec(revision=1, tareas=[{"id": "viejo", "descripcion": "tarea de la revisión 1 distinta"}]); ruta = json_(tmp_path, antigua, "vieja.json")
    antes = json.dumps(a.proyecto.a_dict(), sort_keys=True)
    r = act(a, ruta); assert r["error"] == "revision_antigua" and "más antigua" in textos(r) and "1" in r["revision"]["motivo"] and r["revision"]["aplicada"] == 2
    assert json.dumps(a.proyecto.a_dict(), sort_keys=True) == antes
    assert a.abrir_receta(ruta)["receta"]["estado_revision"]["accion"] == "mas_antigua"
    r = act(a, ruta, forzar=True); assert r["error"] is None and a.proyecto.formatos["_receta"]["revision"] == 1 and tarea(a, "viejo")
    assert [e["revision"] for e in a.proyecto.formatos["_revisiones"]] == [None, 2] or [e["revision"] for e in a.proyecto.formatos["_revisiones"]] == [1, 2]
    assert a.proyecto.formatos["_revisiones"][-1]["revision"] == 2                                              # la 2 pasó al historial antes de reemplazarse
    # la misma revisión (igual número) con otro contenido tampoco es posterior
    mism = rec(revision=1, tareas=[{"id": "otra", "descripcion": "x"}]); assert act(a, json_(tmp_path, mism, "m.json"))["error"] == "revision_antigua"


def test_r5_conflicto_de_revision_a3(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    assert act(a, json_(tmp_path, rec(revision=2, rev_base=1), "r2.json"))["error"] is None
    conf = rec(revision=4, rev_base=1, tareas=[{"id": "nueva", "descripcion": "tarea de la revisión 4"}]); ruta = json_(tmp_path, conf, "r4.json")
    estado = a.abrir_receta(ruta)["receta"]["estado_revision"]; assert estado["accion"] == "conflicto" and estado["base"] == 1 and estado["aplicada"] == 2
    antes = json.dumps(a.proyecto.a_dict(), sort_keys=True)
    r = act(a, ruta); assert r["error"] == "conflicto_revision" and "Conflicto de revisión" in textos(r) and r["revision"]["recibida"] == 4
    assert json.dumps(a.proyecto.a_dict(), sort_keys=True) == antes and tarea(a, "nueva") is None                     # no se aplica sin elegir
    assert act(a, ruta, resolucion="cancelar")["error"] == "cancelado" and json.dumps(a.proyecto.a_dict(), sort_keys=True) == antes
    assert act(a, ruta, resolucion="otra cosa")["error"] == "conflicto_revision"                                        # solo vale cancelar o solo_nuevo
    r = act(a, ruta, resolucion="solo_nuevo"); assert r["error"] is None and tarea(a, "nueva") and a.proyecto.formatos["_receta"]["revision"] == 4
    assert a.proyecto.formatos["_revisiones"][-1]["revision"] == 2
    # base nula contra una revisión aplicada, y base distinta sin ninguna aplicada, también son conflicto
    assert act(a, json_(tmp_path, rec(revision=5), "r5.json"))["error"] == "conflicto_revision"
    b = Api(); b.ejecutar_receta(json_(tmp_path, rec(), "sin.json")); assert b.proyecto.formatos["_receta"]["revision"] is None
    assert act(b, json_(tmp_path, rec(revision=3, rev_base=2), "x.json"))["error"] == "conflicto_revision"
    assert act(b, json_(tmp_path, rec(revision=1), "y.json"))["error"] is None                                         # ninguna aplicada y base nula: se aplica


def test_r5_receta_sin_revision_se_aplica_como_hoy(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=3)))
    sin = rec(tareas=[{"id": "s", "descripcion": "sin revisión"}]); r = act(a, json_(tmp_path, sin, "s.json"))
    assert r["error"] is None and a.proyecto.formatos["_receta"]["revision"] is None and tarea(a, "s")
    assert a.proyecto.formatos["_revisiones"][-1]["revision"] == 3


def test_r5_la_receta_anterior_pasa_a_revisiones_sin_perder_nada_a4(tmp_path):
    a = Api(); rc1 = rec(revision=1); a.ejecutar_receta(json_(tmp_path, rc1))
    assert act(a, libro(tmp_path, rec(revision=2, rev_base=1), {"mediciones": [["Espesor", "Lote"]] + DATOS_ESPESOR}, "l2.xlsx"))["error"] is None
    uid, valores = a.proyecto.formatos["_analisis"]["estad"]["uid"], [a.hoja.columnas[a.hoja.indice("Espesor")].valor(i)[2] for i in range(a.hoja.n)]
    # revisión 3 trae otros valores en la misma columna: una columna con datos nunca se sobrescribe
    r3 = libro(tmp_path, rec(revision=3, rev_base=2), {"mediciones": [["Espesor", "Lote"]] + [[7.77, "Z"]] * 6}, "l3.xlsx")
    r = act(a, r3); assert r["error"] is None and "ya tenía datos" in textos(r)
    assert [a.hoja.columnas[a.hoja.indice("Espesor")].valor(i)[2] for i in range(a.hoja.n)] == valores
    assert uid in {x["uid"] for x in a.resultados()} and a.proyecto.formatos["_analisis"]["estad"]["uid"] == uid
    rev = a.proyecto.formatos["_revisiones"]; assert [e["revision"] for e in rev] == [1, 2]
    assert rev[0]["huella"] == rc1["huella"] and json.loads(rev[0]["receta"]) == rc1 and set(rev[0]) == {"revision", "huella", "aplicada_en", "origen", "receta"}
    assert rev[0]["origen"] == "Diagnóstico en línea" and rev[0]["aplicada_en"]
    assert a.proyecto.formatos["_receta"]["revision"] == 3 and a.proyecto.formatos["_receta"]["huella"] == RC.huella_receta(json.loads(a.proyecto.formatos["_receta"]["json"]))


def test_r5_tope_de_20_revisiones_con_aviso(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    avisos = []
    for n in range(2, 24):
        r = act(a, json_(tmp_path, rec(revision=n, rev_base=n - 1, origen=f"Diagnóstico en línea {n}"), f"r{n}.json")); assert r["error"] is None, (n, textos(r))
        avisos += [m["texto"] for m in r["sesion"] if "tope de 20" in m["texto"]]
    rev = a.proyecto.formatos["_revisiones"]
    assert len(rev) == 20 and [e["revision"] for e in rev] == list(range(3, 23)) and a.proyecto.formatos["_receta"]["revision"] == 23
    assert len(avisos) == 2 and "descartó" in avisos[0]                                                          # se avisó al descartar (revisiones 1 y 2)
    assert SG.aplicada(a.proyecto)["huella"] == a.proyecto.formatos["_receta"]["huella"]


def test_r5_antes_de_archivar_se_ignora_la_misma_receta(tmp_path):
    p = Api().proyecto; assert SG.archivar_receta_previa(p, rec()) == [] and "_revisiones" not in p.formatos        # sin receta previa no hay nada que archivar


# ---------- E2-R6: proyecto real sin mediciones ----------
def test_r6_real_sin_mediciones_a5(tmp_path):
    rc = rec(revision=1, variables=[{"nombre": "Espesor", "hoja": "mediciones", "estado_evidencia": "PENDIENTE_RECOLECCION"},
                                    {"nombre": "Dureza", "hoja": "mediciones", "estado_evidencia": "DESCONOCIDO", "ensayo_destructivo": True}])
    a = Api(); r = a.ejecutar_receta(json_(tmp_path, rc)); assert r["error"] is None, textos(r)
    assert MD.valor(a.proyecto) == "real" and a.proyecto.formatos["_modo"]["valor"] == "real"
    assert [c.nombre for c in a.hoja.columnas] == ["Espesor", "Lote"] and a.hoja.n == 0 and all(not any(c.celdas) for c in a.hoja.columnas)      # sin datos sintéticos ni cifras
    assert a.resultados() == [] and a.proyectos()["proyectos"][0]["sintetico"] is False
    assert {x["nombre"]: x["estado_evidencia"] for x in a.proyecto.formatos["_variables"]} == {"Espesor": "PENDIENTE_RECOLECCION", "Dureza": "DESCONOCIDO"}
    an = a.proyecto.formatos["_analisis"]["estad"]; assert an["estado"] == "pendiente" and an["motivo"] and an["uid"] is None
    assert {t["id"] for t in a.proyecto.formatos["_tareas"]} == {"auto:hoja:mediciones", "auto:analisis:estad"}
    assert a.exportar_avance(str(tmp_path / "av.json"))["error"] is None
    pkg = json.loads((tmp_path / "av.json").read_text(encoding="utf-8")); assert pkg["estado_ruta"] == "Sin datos"
    # Datos aleatorios sigue bloqueado en un proyecto real
    assert a.dialogo_modulo("ALEATORIOS")["error"] == "entrada"
    # una receta real que trae datos sintéticos se rechaza
    mala = rec(revision=1, hojas=[{"hoja": "mediciones", "origen_datos": "sintetico", "columnas": [{"nombre": "Espesor"}]}])
    assert Api().ejecutar_receta(json_(tmp_path, mala, "mala.json"))["error"] == "receta"


# ---------- E2-R7: paquete de avance ----------
CLAVES_AVANCE = {"tipo", "avance_version", "id_proyecto", "modo_proyecto", "revision_aplicada", "generado_en", "huella", "fase_actual", "estado_ruta", "variables", "analisis", "disenos", "tareas_abiertas"}


def leer_avance(a, tmp_path, nombre="av.json"):
    ruta = tmp_path / nombre; r = a.exportar_avance(str(ruta)); assert r["error"] is None, textos(r)
    return json.loads(ruta.read_text(encoding="utf-8")), ruta.read_text(encoding="utf-8")


def test_r7_paquete_de_avance_solo_estado_a7(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    llena = libro(tmp_path, rec(revision=2, rev_base=1), {"mediciones": [["Espesor", "Lote"]] + DATOS_ESPESOR})
    act(a, llena)
    pkg, texto = leer_avance(a, tmp_path)
    assert set(pkg) == CLAVES_AVANCE and pkg["tipo"] == "dmike-avance" and pkg["avance_version"] == 1 and pkg["id_proyecto"] == UUID_A
    assert pkg["modo_proyecto"] == "real" and pkg["revision_aplicada"] == 2 and pkg["estado_ruta"] == "Historicos cargados" and pkg["disenos"] == []
    assert pkg["variables"] == [{"nombre": "Espesor", "estado_evidencia": "PENDIENTE_RECOLECCION"}] and pkg["analisis"] == [{"id": "estad", "modulo": "ESTADISTICOS", "estado": "analizado"}]
    assert pkg["tareas_abiertas"] == []
    # ni datos de hojas, ni valores medidos, ni resultados numéricos
    res = json.dumps(a.resultado(a.proyecto.formatos["_analisis"]["estad"]["uid"])["bloques"], ensure_ascii=False)
    for v in DATOS_ESPESOR:
        assert str(v[0]) not in texto and f'"{v[1]}"' not in texto
    media = sum(v[0] for v in DATOS_ESPESOR) / len(DATOS_ESPESOR); assert f"{media:.3f}" not in texto and "0.50" not in texto
    assert "resultados" not in pkg and "bloques" not in texto and len(res) > 50
    # la huella se verifica igual que la de la receta
    assert pkg["huella"] == RC.huella_receta(pkg) and SG.huella_avance_valida(pkg)
    alterado = dict(pkg, revision_aplicada=9); assert not SG.huella_avance_valida(alterado)
    assert SG.huella_avance_valida(json.loads(texto))


def test_r7_sin_id_proyecto_no_se_exporta(tmp_path):
    ruta = tmp_path / "sin.json"
    for a in (Api(), ):                                                                                           # proyecto sin receta
        r = a.exportar_avance(str(ruta)); assert r["error"] == "sin_id_proyecto" and not ruta.exists()
        assert "id_proyecto" in textos(r) and "receta v1" in textos(r) and "sin receta" in textos(r) and "no se exportó" in textos(r).lower()
    b = Api(); b.ejecutar_receta(json_(tmp_path, rec(revision=1, idp=None), "sinid.json"))
    assert b.exportar_avance(str(ruta))["error"] == "sin_id_proyecto" and not ruta.exists()                          # receta sin id_proyecto
    c = Api(); c.ejecutar_receta(str(S / "caso_integral_datos_llenos.xlsx")); assert c.exportar_avance(str(ruta))["error"] == "sin_id_proyecto" and not ruta.exists()      # receta v1
    d = Api(); d.ejecutar_receta(json_(tmp_path, rec(revision=1), "ok.json")); d._elegir = lambda *x, **k: None
    assert d.exportar_avance()["error"] is None and not ruta.exists()                                                # el usuario cancela el diálogo: no hay archivo
    d._elegir = lambda *x, **k: str(ruta); assert d.exportar_avance()["error"] is None and ruta.exists()


def test_r7_estado_ruta_y_disenos_con_el_caso_maestro(tmp_path):
    rc, _ = RC.leer(str(S / "caso_integral_plantilla.xlsx")); rc = copy.deepcopy(rc)
    rc.update(receta_version=2, dmike_minimo="0.27", modo_proyecto="demostracion", id_proyecto=UUID_B, revision=1); rc.pop("folio", None); rc["huella"] = RC.huella_receta(rc)
    vacia = libro(tmp_path, rc, {})
    a = Api(); r = a.ejecutar_receta(vacia); assert r["error"] is None, textos(r)
    pkg, texto = leer_avance(a, tmp_path)
    assert pkg["estado_ruta"] == "DOE disenado" and pkg["modo_proyecto"] == "demostracion"
    d = {x["id"]: x for x in pkg["disenos"]}; assert set(d) == {"cribado", "dcc"} and all(x["estado"] == "creado" and x["respuestas_capturadas"] == 0 and x["corridas"] > 0 for x in d.values())
    assert {t["id"] for t in pkg["tareas_abiertas"]} >= {"auto:diseno:cribado", "auto:diseno:dcc", "auto:analisis:factorial"}
    # con las respuestas capturadas: DOE ejecutado
    _, datos = RC.leer(str(S / "caso_integral_datos_llenos.xlsx"))
    wb = openpyxl.load_workbook(str(S / "caso_integral_datos_llenos.xlsx")); wb["_receta"].delete_rows(1, wb["_receta"].max_row)
    t = json.dumps(rc, ensure_ascii=False)
    for i in range(0, len(t), 30000): wb["_receta"].cell(i // 30000 + 1, 1, t[i:i + 30000])
    llena = str(tmp_path / "llena.xlsx"); wb.save(llena)
    r = act(a, llena); assert r["error"] is None, textos(r)
    pkg2, texto2 = leer_avance(a, tmp_path, "av2.json")
    assert pkg2["estado_ruta"] == "DOE ejecutado" and {x["id"]: (x["estado"], x["respuestas_capturadas"]) for x in pkg2["disenos"]} == {"cribado": ("ejecutado", 1), "dcc": ("ejecutado", 2)}
    assert pkg2["tareas_abiertas"] == [] and all(x["estado"] == "analizado" for x in pkg2["analisis"]) and len(pkg2["analisis"]) == 15
    assert "resultados" not in pkg2 and len(texto2) < 6000                                                         # solo estado
    # nada de lo medido: ningún valor de las columnas de respuesta aparece en el paquete
    vals = {str(c.celdas[i]) for c in a.hoja.columnas if c.nombre in ("Resistencia", "Porosidad", "Resistencia_cribado", "resistencia_base") for i in range(a.hoja.n) if c.celdas[i] not in (None, "")}
    vals = {v for v in vals if "." in v and len(v) >= 4}; assert len(vals) > 20
    assert not [v for v in vals if v in texto2]


def test_r7_estado_ruta_no_determinado_y_resto_de_valores():
    a = Api(); assert SG.estado_ruta(a.proyecto, a.hoja) == "Sin datos"
    a.pegar(-1, 0, "\n1\n2\n3")                                                                                     # datos en una columna sin nombre: no se puede decir más
    assert SG.estado_ruta(a.proyecto, a.hoja) == "no_determinado"
    b = Api(); b.pegar(-1, 0, "x\n1\n2\n3"); assert SG.estado_ruta(b.proyecto, b.hoja) == "Historicos cargados"
    valores = {SG.estado_ruta(b.proyecto, b.hoja), "Sin datos", "DOE disenado", "DOE ejecutado", SG.NO_DETERMINADO}
    assert valores <= set(SG.RUTA) | {SG.NO_DETERMINADO}                                                           # nunca "línea base", "causas", etc.


def test_r7_tarea_auto_de_analisis_no_filtra_mensajes_del_modulo(tmp_path):
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    t = tarea(a, "auto:analisis:estad"); t["descripcion"] = "Análisis «estad»: valor 123.456 en la celda B7"           # un mensaje del módulo con un dato
    pkg, texto = leer_avance(a, tmp_path); assert "123.456" not in texto and "B7" not in texto
    assert {x["id"]: x["descripcion"] for x in pkg["tareas_abiertas"]}["auto:analisis:estad"].startswith("Análisis «estad»")


# ---------- E2-R8: menú y comandos ----------
def test_r8_menu_y_catalogo_de_comandos():
    archivo_m = dict(menus.MENUS)["Archivo"]
    assert ("Importar receta del diagnóstico…", "importar_receta") in archivo_m and ("Exportar paquete de avance…", "exportar_avance") in archivo_m
    assert archivo_m.index(("Importar receta del diagnóstico…", "importar_receta")) == archivo_m.index(("Abrir proyecto…", "abrir_proyecto")) + 1
    assert ("Seguimiento del proyecto…", "seguimiento") in dict(menus.MENUS)["Proyecto"]
    nombres = {c["n"] for c in Api()._comandos()}
    assert {"Importar receta del diagnóstico", "Exportar paquete de avance", "Seguimiento del proyecto"} <= nombres


# ---------- E2-R10: compatibilidad en ambos sentidos ----------
@pytest.fixture(scope="module")
def dmike_026(tmp_path_factory):
    ruta = tmp_path_factory.mktemp("m26") / "caso_026.dmike"
    with gzip.open(S / "caso_integral_v0.26.0.dmike.gz", "rb") as g, open(ruta, "wb") as f: shutil.copyfileobj(g, f)
    return str(ruta), json.loads((S / "caso_integral_v0.26.0.esperado.json").read_text(encoding="utf-8"))


def test_r10_dmike_026_abre_en_027_sin_cambios_a8(dmike_026):
    ruta, esp = dmike_026
    a = Api(); r = a.abrir_proyecto(ruta); assert r["error"] is None
    res = [a.resultado(x["uid"]) for x in a.resultados()]
    assert len(res) == esp["n_resultados"] == 17 and a.hoja.n == esp["n_filas"] and [c.nombre for c in a.hoja.columnas] == esp["columnas"]
    assert [x["definicion"]["congelado"]["huella"] for x in res] == esp["huellas"]
    assert [hashlib.sha256(json.dumps(x["bloques"], sort_keys=True, ensure_ascii=False).encode()).hexdigest() for x in res] == esp["sha_bloques"]
    assert sorted(a.proyecto.formatos) == esp["claves_formatos"]                                                    # abrir no escribe ninguna clave nueva
    assert all(a.reproducir_resultado(x["uid"])["error"] is None for x in a.resultados())
    assert a.proyecto.formatos["_receta"]["folio"] == esp["folio_receta"] == "DMK-PRUEBA-0001"
    # y todo lo nuevo funciona sobre él sin escribir nada
    s = a.seguimiento(); assert s["revision"] is None and s["variables"] == [] and s["tareas_abiertas"] == [] and s["analisis"] == []
    assert a.exportar_avance(str(Path(ruta).parent / "x.json"))["error"] == "sin_id_proyecto"
    assert sorted(a.proyecto.formatos) == esp["claves_formatos"]


def test_r10_actualizar_un_dmike_026_conserva_folio_y_archiva_la_receta(dmike_026, tmp_path):
    ruta, esp = dmike_026
    a = Api(); a.abrir_proyecto(ruta); n = len(a.resultados())
    rc = rec(revision=1, modo="demostracion"); r = act(a, json_(tmp_path, rc))
    assert r["error"] is None, textos(r)
    assert a.proyecto.formatos["_receta"]["folio"] == "DMK-PRUEBA-0001" and len(a.resultados()) == n + 0
    rev = a.proyecto.formatos["_revisiones"]; assert len(rev) == 1 and rev[0]["revision"] is None and rev[0]["huella"].startswith("sha256:")
    assert json.loads(rev[0]["receta"])["origen"].startswith("Caso maestro")                                       # la receta v1 anterior quedó guardada completa
    assert [c.nombre for c in a.hoja.columnas][:31] == esp["columnas"][:31]                                         # no se perdió ninguna columna


def _codigo_026(tmp_path):
    """Extrae el código de la 0.26.0 desde el historial de git. Se omite si el historial no está disponible (clon superficial)."""
    raiz = Path(__file__).resolve().parent.parent
    for rev in ("ca81621", "17cd6a8", "7e1fc2d"):
        try:
            if 'VERSION = "0.26.0"' not in subprocess.run(["git", "show", f"{rev}:app/api.py"], cwd=raiz, capture_output=True, text=True, check=True).stdout: continue
            dest = tmp_path / "viejo"; dest.mkdir()
            tar = subprocess.run(["git", "archive", rev, "app"], cwd=raiz, capture_output=True, check=True).stdout
            subprocess.run(["tar", "-x", "-C", str(dest)], input=tar, check=True)
            return dest
        except (subprocess.CalledProcessError, FileNotFoundError): continue
    return None


def test_r10_dmike_027_abre_en_026_sin_error_a8(tmp_path):
    viejo = _codigo_026(tmp_path)
    if viejo is None: pytest.skip("no hay historial de git con la 0.26.0")
    a = Api(); a.ejecutar_receta(json_(tmp_path, rec(revision=1, variables=[{"nombre": "Espesor", "estado_evidencia": "REAL_VERIFICADO"}], tareas=[{"id": "t1", "descripcion": "x"}])))
    assert act(a, libro(tmp_path, rec(revision=2, rev_base=1, variables=[{"nombre": "Espesor", "estado_evidencia": "REAL_VERIFICADO"}]), {"mediciones": [["Espesor", "Lote"]] + DATOS_ESPESOR}))["error"] is None
    ruta = str(tmp_path / "nuevo.dmike"); archivo.guardar(ruta, a.hoja, a.proyecto, VERSION)
    claves = sorted(k for k in a.proyecto.formatos if k.startswith("_")); assert {"_variables", "_tareas", "_analisis", "_revisiones", "_receta"} <= set(claves)
    con = __import__("sqlite3").connect(ruta); assert dict(con.execute("SELECT clave, valor FROM meta"))["version_formato"] == "1" == str(archivo.VERSION_FORMATO); con.close()
    codigo = ("import sys, json; sys.path.insert(0, sys.argv[1]); from app.api import Api; import app.api as A; from app.nucleo import receta as RC; from app import modulos\n"
              "a = Api(); r = a.abrir_proyecto(sys.argv[2]); rc = json.load(open(sys.argv[3]))\n"
              "print(A.VERSION, r['error'], a.hoja.n, len(a.hoja.columnas), len(a.resultados()), sorted(k for k in a.proyecto.formatos if k.startswith('_')))\n"
              "print(RC.validar(rc, modulos.cargar(), A.VERSION))")
    nueva = json_(tmp_path, rec(revision=1), "para026.json")
    out = subprocess.run([sys.executable, "-I", "-c", codigo, str(viejo), ruta, nueva], capture_output=True, text=True, timeout=180)
    assert out.returncode == 0, out.stderr[-600:]
    l1, l2 = out.stdout.splitlines()[:2]
    assert l1.split()[:2] == ["0.26.0", "None"] and "_variables" in l1 and "_revisiones" in l1 and str(len(a.resultados())) in l1
    assert "requiere D-MIKE 0.27" in l2                                                                              # una receta con campos de la Entrega 2 la rechaza la 0.26
    # y al guardarlo desde la 0.26 las claves nuevas siguen ahí (la 0.26 las conserva sin entenderlas)
    codigo2 = ("import sys; sys.path.insert(0, sys.argv[1]); from app.api import Api\n"
               "a = Api(); a.abrir_proyecto(sys.argv[2]); a._elegir = lambda *x, **k: sys.argv[3]; r = a.guardar(True); print(r['error'])")
    re = str(tmp_path / "re026.dmike")
    out = subprocess.run([sys.executable, "-I", "-c", codigo2, str(viejo), ruta, re], capture_output=True, text=True, timeout=180)
    assert out.returncode == 0 and out.stdout.split()[0] == "None", out.stderr[-600:]
    b = Api(); assert b.abrir_proyecto(re)["error"] is None and b.proyecto.formatos["_receta"]["revision"] == 2 and len(b.proyecto.formatos["_tareas"]) == len(a.proyecto.formatos["_tareas"])


def test_r10_receta_027_del_diagnostico_se_declara_con_dmike_minimo(tmp_path):
    mayor, menor = VERSION.split(".")[:2]
    siguiente = f"{mayor}.{int(menor) + 1}"                                       # la versión vigente más uno: la prueba no se rompe al subir VERSION
    rc = rec(); rc["dmike_minimo"] = siguiente; rc["huella"] = RC.huella_receta({k: v for k, v in rc.items() if k != "huella"})
    assert any(f"requiere D-MIKE {siguiente}" in e for e in RC.validar(rc, REG, VERSION))     # una receta que pide una versión posterior a la vigente se rechaza


# ---------- A1: el caso maestro v1 no cambia ----------
def test_a1_caso_maestro_v1_17_resultados_y_optimos():
    a = Api(); r = a.ejecutar_receta(str(S / "caso_integral_datos_llenos.xlsx")); assert r["error"] is None
    res = [a.resultado(x["uid"]) for x in a.resultados()]; assert len(res) == 17
    opt = next(x for x in res if x["definicion"]["analisis"] == "OPTIMIZADOR")
    t = next(b for b in opt["bloques"] if b["tipo"] == "tabla" and b["titulo"] == "Ajuste óptimo de los factores")
    assert [f[1] for f in t["filas"]] == ["129.42323", "108.63155"]
    assert r["receta_resultado"]["log"][-1] == "Análisis corridos: 15 de 15." and r["receta_resultado"]["tareas_cerradas"] == []
    esp = json.loads((S / "caso_integral_v0.26.0.esperado.json").read_text(encoding="utf-8"))
    assert [x["definicion"]["congelado"]["huella"] for x in res] == esp["huellas"]                                  # mismas huellas que con la 0.26.0
    assert a.proyecto.formatos["_receta"]["revision"] is None and a.proyecto.formatos["_receta"]["huella"].startswith("sha256:")


# ---------- interfaz ----------
def _abrir_por_menu(pagina, ruta):
    pagina.api._elegir = lambda *x, **k: ruta
    pagina.evaluate("ACCIONES_UI.importar_receta()")
    pagina.wait_for_selector(".dialogo:has-text('¿Ejecutar la receta?')")


def test_interfaz_importar_receta_y_seguimiento(pagina, tmp_path):
    ruta = json_(tmp_path, rec(revision=1, tareas=[{"id": "t1", "descripcion": "Pedir el certificado", "responsable": "Calidad", "accion": "Pedir"}]))
    _abrir_por_menu(pagina, ruta)
    pagina.click(".dialogo .primario")
    pagina.wait_for_function("document.querySelector('#sesion').textContent.includes('Receta ejecutada')", timeout=60000)
    assert pagina.api.proyecto.formatos["_receta"]["revision"] == 1
    pagina.evaluate("ACCIONES_UI.seguimiento()")
    pagina.wait_for_selector(".dialogo:has-text('Seguimiento del proyecto')")
    cuerpo = pagina.inner_text("#seguimiento")
    assert "Revisión aplicada: 1" in cuerpo and "Espesor" in cuerpo and "PENDIENTE_RECOLECCION" in cuerpo and "Pedir el certificado" in cuerpo
    assert "Llenar la hoja" in cuerpo and "estad" in cuerpo and "pendiente" in cuerpo and "Sin datos en Espesor" in cuerpo
    assert pagina.locator("#seguimiento input, #seguimiento textarea, #seguimiento select").count() == 0          # solo lectura: sin editor
    pagina.click(".dialogo-botones button:not(.primario)")
    pagina.evaluate("ACCIONES_UI.exportar_avance()")
    pagina.api._elegir = lambda *x, **k: str(tmp_path / "av.json")
    pagina.evaluate("ACCIONES_UI.exportar_avance()")
    pagina.wait_for_function("document.querySelector('#sesion').textContent.includes('Paquete de avance exportado')", timeout=30000)
    assert (tmp_path / "av.json").exists()


def test_interfaz_conflicto_de_revision_en_la_apertura_de_la_receta(pagina, tmp_path):
    a = pagina.api; a.ejecutar_receta(json_(tmp_path, rec(revision=1)))
    act(a, json_(tmp_path, rec(revision=2, rev_base=1), "r2.json"))
    pagina.evaluate("pintarDocs()")
    conf = rec(revision=4, rev_base=1, tareas=[{"id": "nueva", "descripcion": "tarea de la revisión 4"}])
    _abrir_por_menu(pagina, json_(tmp_path, conf, "r4.json"))
    assert "Conflicto de revisión" in pagina.inner_text("#receta-revision")                                          # el aviso ya está en el diálogo de apertura
    pagina.click("#receta-actualizar")
    pagina.wait_for_selector(".dialogo:has-text('Conflicto de revisión') #conflicto-revision")
    assert "No hay fusión automática" in pagina.inner_text("#conflicto-revision") and a.proyecto.formatos["_receta"]["revision"] == 2
    pagina.click(".dialogo[aria-label='Conflicto de revisión'] .dialogo-botones button:not(.primario)")           # cancelar: no cambia nada
    assert a.proyecto.formatos["_receta"]["revision"] == 2 and tarea(a, "nueva") is None
    _abrir_por_menu(pagina, json_(tmp_path, conf, "r4.json")); pagina.click("#receta-actualizar")
    pagina.wait_for_selector("#conflicto-revision")
    pagina.click(".dialogo[aria-label='Conflicto de revisión'] .primario")                                           # aplicar solo lo nuevo
    pagina.wait_for_function("document.querySelector('#sesion').textContent.includes('Receta ejecutada')", timeout=60000)
    assert a.proyecto.formatos["_receta"]["revision"] == 4 and tarea(a, "nueva") and a.proyecto.formatos["_revisiones"][-1]["revision"] == 2


def test_interfaz_receta_mas_antigua_y_ya_aplicada(pagina, tmp_path):
    a = pagina.api; ruta = json_(tmp_path, rec(revision=1)); a.ejecutar_receta(ruta)
    act(a, json_(tmp_path, rec(revision=2, rev_base=1), "r2.json")); pagina.evaluate("pintarDocs()")
    _abrir_por_menu(pagina, ruta)                                                                                   # revisión 1, ya aplicada (su huella está en el historial)
    assert "ya estaba aplicada" in pagina.inner_text("#receta-revision")
    pagina.click("#receta-actualizar")
    pagina.wait_for_function("document.querySelector('#sesion').textContent.includes('ya aplicada')", timeout=30000)
    antigua = rec(revision=1, tareas=[{"id": "viejo", "descripcion": "x"}])
    _abrir_por_menu(pagina, json_(tmp_path, antigua, "vieja.json")); pagina.click("#receta-actualizar")
    pagina.wait_for_selector(".dialogo[aria-label='Receta más antigua']")
    assert a.proyecto.formatos["_receta"]["revision"] == 2
    pagina.click(".dialogo[aria-label='Receta más antigua'] .primario")
    pagina.wait_for_function("document.querySelector('#sesion').textContent.includes('Receta ejecutada')", timeout=60000)
    assert a.proyecto.formatos["_receta"]["revision"] == 1 and tarea(a, "viejo")


def test_interfaz_menu_archivo_trae_las_dos_opciones(pagina):
    pagina.click(".menu > button:text-is('Archivo')")
    textos_ = pagina.locator(".menu.abierto button[role='menuitem']").all_inner_texts()
    assert any("Importar receta del diagnóstico" in t for t in textos_) and any("Exportar paquete de avance" in t for t in textos_)
    pagina.keyboard.press("Escape")
    assert pagina.evaluate("['importar_receta','exportar_avance','seguimiento'].every(k => typeof ACCIONES_UI[k] === 'function')")
