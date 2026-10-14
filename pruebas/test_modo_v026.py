"""D-MIKE 0.26, Entrega 1: modo real/demostración, recetas v1 y v2, id_proyecto, marca sintética y migración (E1-R1 a E1-R10).
Datos SINTÉTICOS (caso maestro y recetas armadas aquí)."""
import copy, os, gzip, hashlib, json, shutil, sqlite3, subprocess, sys
from pathlib import Path
import openpyxl, pytest
from app.api import Api
from app.nucleo import archivo, modo as MD, receta as RC
from app import modulos

S = Path(__file__).parent / "sinteticos"
REG = modulos.cargar()
UUID_A, UUID_B = "3f2b8c1e-5d4a-4f6b-9a7c-1e2d3c4b5a69", "9b1deb4d-3b7d-4bad-9bdd-2b0d7b3dcb6d"
DATOS = "y\tg\n" + "\n".join(f"{10 + (i % 3) * 2 + (i % 5) * .37123}\t{'ABC'[i % 3]}" for i in range(30))


@pytest.fixture(autouse=True)
def _recientes_aislados(tmp_path, monkeypatch):
    """Los proyectos recientes viven en un archivo de usuario (APPDATA): cada prueba usa el suyo para no leer ni ensuciar el real."""
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))


# ---------- ayudas ----------
def base():
    return copy.deepcopy(RC.leer(str(S / "caso_integral_plantilla.xlsx"))[0])


def escribir(tmp, nombre, rc):
    ruta = tmp / nombre; ruta.write_text(json.dumps(rc, ensure_ascii=False), encoding="utf-8"); return str(ruta)


def v1(tmp, origen, folio=None, nombre="r1.json"):
    rc = base(); rc.pop("folio", None); rc["origen"] = origen
    if folio: rc["folio"] = folio
    return escribir(tmp, nombre, rc)


def v2(tmp, modo, idp=UUID_A, hojas_sinteticas=False, variables=None, nombre="r2.json", huella=True, revision=None, rev_base=None):
    rc = base(); rc.pop("folio", None)
    rc.update(receta_version=2, dmike_minimo="0.26", modo_proyecto=modo, origen="Diagnóstico en línea")
    if idp: rc["id_proyecto"] = idp
    if hojas_sinteticas: rc["hojas"][0]["origen_datos"] = "sintetico"
    if variables: rc["variables"] = variables
    if revision: rc["revision"] = revision
    if rev_base is not None: rc["revision_base"] = rev_base
    if huella: rc["huella"] = RC.huella_receta(rc)
    return escribir(tmp, nombre, rc)


def formato_modo(a): return a.proyecto.formatos.get("_modo")


# ---------- E1-R1, E1-R2: versión y validación v2 ----------
def test_r1_version_soportada_2_y_mayor_se_rechaza():
    assert RC.VERSION_SOPORTADA == 2
    rc = base(); rc["receta_version"] = 3
    e = RC.validar(rc, REG, "0.26.0"); assert any("más nueva" in x for x in e)
    assert RC.validar(base(), REG, "0.26.0") == []          # v1 intacta


def test_r2_validacion_v2(tmp_path):
    ok = RC.leer(v2(tmp_path, "real"))[0]; assert RC.validar(ok, REG, "0.26.0") == []
    r = dict(ok, modo_proyecto="otro"); r["huella"] = RC.huella_receta(r); assert any("modo_proyecto" in x for x in RC.validar(r, REG, "0.26.0"))
    r = copy.deepcopy(ok); r.pop("modo_proyecto"); r["huella"] = RC.huella_receta(r); assert any("debe declarar modo_proyecto" in x for x in RC.validar(r, REG, "0.26.0"))
    for malo in ("DMK-ABCD-EFGH-JKLM", UUID_A.upper(), UUID_A.replace("-4", "-1", 1), "123"):
        r = dict(ok, id_proyecto=malo); r["huella"] = RC.huella_receta(r); assert any("UUID v4" in x for x in RC.validar(r, REG, "0.26.0")), malo
    r = dict(ok, revision=9); assert any("huella" in x for x in RC.validar(r, REG, "0.26.0"))           # alterada tras calcular la huella
    r = dict(ok); r["huella"] = r["huella"].removeprefix("sha256:"); assert RC.validar(r, REG, "0.26.0") == []
    sin = {k: v for k, v in ok.items() if k != "huella"}; assert RC.validar(sin, REG, "0.26.0") == []     # la huella es opcional
    assert RC.huella_receta(ok) == RC.huella_receta(sin)                                                  # no depende del campo huella
    assert RC.huella_receta(ok) == "sha256:" + hashlib.sha256(json.dumps(sin, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


def test_r2_sinteticos_solo_con_demostracion(tmp_path):
    hs = {"hojas_sinteticas": True, "variables": [{"nombre": "x", "hoja": "h", "estado_evidencia": "SINTETICO_DEMO"}]}
    real = RC.leer(v2(tmp_path, "real", **hs))[0]; e = RC.validar(real, REG, "0.26.0")
    assert any("hoja" in x and "sintética" in x for x in e) and any("SINTETICO_DEMO" in x for x in e)
    demo = RC.leer(v2(tmp_path, "demostracion", **hs))[0]; assert RC.validar(demo, REG, "0.26.0") == []
    mala = RC.leer(v2(tmp_path, "demostracion", variables=[{"nombre": "x", "estado_evidencia": "INVENTADO"}]))[0]
    assert any("estado_evidencia" in x for x in RC.validar(mala, REG, "0.26.0"))


# ---------- E1-R3: inferencia ----------
@pytest.mark.parametrize("origen,esperado", [
    ("Caso maestro de D-MIKE: uniones adhesivas (datos sintéticos)", "demostracion"), ("CASO MAESTRO", "demostracion"),
    ("Datos SINTETICOS de práctica", "demostracion"), ("sintético", "demostracion"), ("Prueba Sintética", "demostracion"),
    ("Diagnóstico en línea", "real"), ("Línea 3, planta Querétaro", "real"), ("", "real"), (None, "real")])
def test_r3_inferencia_por_origen(origen, esperado):
    assert MD.inferir_por_origen(origen) == esperado
    assert MD.detectar({"origen": origen})[:2] == (esperado, "inferido_v1")


def test_r3_huellas_de_recetas_sinteticas_distribuidas(monkeypatch):
    """Refuerzo opcional: la lista del repositorio coincide con las recetas sintéticas que distribuye y las marca aunque el origen no lo diga."""
    esperadas = {RC.huella_sintetica(RC.leer(str(S / f))[0]) for f in ("caso_integral_plantilla.xlsx", "caso_integral_datos_llenos.xlsx")}
    assert esperadas <= MD.huellas_conocidas()
    rc = base(); rc["origen"] = "Planta X"
    assert MD.detectar(rc, RC.huella_sintetica(rc))[0] == "real"
    monkeypatch.setattr(MD, "huellas_conocidas", lambda: {RC.huella_sintetica(rc)})
    assert MD.detectar(rc, RC.huella_sintetica(rc)) == ("demostracion", "inferido_v1")
    assert RC.huella_sintetica(dict(rc, folio="DMK-X", generado_en="otra fecha")) == RC.huella_sintetica(rc)      # ni folio ni fecha cuentan


# ---------- A1: caso maestro v1 ----------
def test_a1_caso_maestro_v1_demostracion_inferida_y_17_resultados():
    a = Api(); r = a.abrir_proyecto(str(S / "caso_integral_datos_llenos.xlsx"))
    assert r["receta"]["modo"] == "demostracion" and r["receta"]["deteccion"] == "inferido_v1" and r["receta"]["puede_cambiar_modo"]
    r = a.ejecutar_receta(r["receta"]["ruta"]); assert r["error"] is None
    assert formato_modo(a) == {"valor": "demostracion", "deteccion": "inferido_v1", "cambiado": False}
    res = [a.resultado(x["uid"]) for x in a.resultados()]; assert len(res) == 17
    opt = next(x for x in res if x["definicion"]["analisis"] == "OPTIMIZADOR")
    t = next(b for b in opt["bloques"] if b["tipo"] == "tabla" and b["titulo"] == "Ajuste óptimo de los factores")
    assert [f[1] for f in t["filas"]] == ["129.42323", "108.63155"]
    assert all(x["etiqueta"] == "SINTÉTICO / NO EMPRESARIAL" for x in res)


# ---------- A2, E1-R4: cambio único ----------
def test_a2_v1_otro_origen_es_real_y_permite_un_solo_cambio(tmp_path):
    ruta = v1(tmp_path, "Proyecto planta 7")
    a = Api(); r = a.abrir_proyecto(ruta)
    assert r["receta"]["modo"] == "real" and r["receta"]["puede_cambiar_modo"]
    r = a.ejecutar_receta(ruta, "nuevo", None, "demostracion"); assert r["error"] is None          # el cambio único, antes de cargar datos
    assert formato_modo(a) == {"valor": "demostracion", "deteccion": "inferido_v1", "cambiado": True}
    assert a.cambiar_modo("real")["error"] == "modo" and MD.valor(a.proyecto) == "demostracion"       # segundo intento rechazado
    assert a.modo_proyecto()["puede_cambiar"] is False


def test_r4_cambio_en_proyecto_vacio_se_consume_una_vez_y_no_tras_cargar_datos():
    a = Api(); assert a.modo_proyecto()["valor"] == "real" and a.modo_proyecto()["puede_cambiar"]
    assert a.cambiar_modo("real")["error"] is None and a.modo_proyecto()["puede_cambiar"]            # mismo modo: no consume el cambio
    assert a.cambiar_modo("demostracion")["error"] is None and MD.es_demo(a.proyecto)
    assert a.cambiar_modo("real")["error"] == "modo"
    b = Api(); b.pegar(-1, 0, DATOS)
    r = b.cambiar_modo("demostracion"); assert r["error"] == "modo" and "cargar datos" in r["sesion"][0]["texto"] and not MD.es_demo(b.proyecto)
    assert Api().cambiar_modo("raro")["error"] == "modo"


def test_r4_v2_declarado_no_se_cambia(tmp_path):
    ruta = v2(tmp_path, "real"); a = Api()
    r = a.ejecutar_receta(ruta, "nuevo", None, "demostracion")
    assert r["error"] == "modo" and "declarado" in r["sesion"][0]["texto"] and a.hoja.columnas == [] and "_modo" not in a.proyecto.formatos
    assert a.abrir_receta(ruta)["receta"]["puede_cambiar_modo"] is False
    assert a.ejecutar_receta(ruta)["error"] is None
    assert formato_modo(a) == {"valor": "real", "deteccion": "declarado", "cambiado": False}
    assert a.cambiar_modo("demostracion")["error"] == "modo"


# ---------- A3, E1-R6: real rechaza sintéticos ----------
def test_a3_v2_real_con_hoja_sintetica_se_rechaza_sin_cargar_nada(tmp_path):
    ruta = v2(tmp_path, "real", hojas_sinteticas=True); a = Api()
    r = a.abrir_receta(ruta); assert r["error"] == "receta" and "sintética" in r["sesion"][0]["texto"] and "receta" not in r
    for modo in ("nuevo", "datos"):
        r = a.ejecutar_receta(ruta, modo)
        assert r["error"] == "receta" and "sintética" in r["sesion"][0]["texto"], modo
        assert a.hoja.columnas == [] and a.hoja.n == 0 and len(a.proyectos()["proyectos"]) == 1 and a.proyecto.formatos == {}
    ruta = v2(tmp_path, "real", variables=[{"nombre": "x", "estado_evidencia": "SINTETICO_DEMO"}], nombre="r2b.json")
    assert a.ejecutar_receta(ruta, "datos")["error"] == "receta" and a.hoja.columnas == []


def test_r6_v1_inferida_real_con_hoja_sintetica_se_rechaza_al_ejecutar(tmp_path):
    rc = base(); rc.pop("folio", None); rc["origen"] = "Planta"; rc["hojas"][0]["origen_datos"] = "sintetico"
    ruta = escribir(tmp_path, "x.json", rc); a = Api()
    assert a.abrir_receta(ruta)["receta"]["sintetico_en_receta"] is True            # se muestra; todavía se puede pasar a demostración
    assert a.ejecutar_receta(ruta)["error"] == "receta" and a.hoja.columnas == []
    assert a.ejecutar_receta(ruta, "nuevo", None, "demostracion")["error"] is None


def test_r6_v2_demostracion_acepta_sinteticos(tmp_path):
    ruta = v2(tmp_path, "demostracion", hojas_sinteticas=True, variables=[{"nombre": "x", "estado_evidencia": "SINTETICO_DEMO"}]); a = Api()
    assert a.abrir_receta(ruta)["error"] is None and a.ejecutar_receta(ruta)["error"] is None and a.hoja.columnas


def test_r6_actualizar_exige_el_mismo_modo(tmp_path):
    a = Api(); a.ejecutar_receta(v2(tmp_path, "real", nombre="real.json")); id_real = a.proyectos()["proyectos"][0]["id"]
    ruta = v2(tmp_path, "demostracion", idp=None, nombre="demo.json")
    r = a.ejecutar_receta(ruta, "actualizar", id_real); assert r["error"] == "modo" and "real" in r["sesion"][0]["texto"]
    assert a.cambiar_proyecto(id_real)["error"] is None and a.hoja.columnas and MD.valor(a.proyecto) == "real"


# ---------- A4: Datos aleatorios ----------
PARAMS_ALEA = {"dest": "n", "n": "12", "dist": "normal", "a": "10", "b": "2", "semilla": "5"}


def test_a4_datos_aleatorios_bloqueado_en_real_y_con_marca_en_demostracion():
    a = Api()                                                       # proyecto nuevo sin receta = real
    r = a.ejecutar_modulo("ALEATORIOS", PARAMS_ALEA)
    assert r["error"] == "entrada" and "bloqueado" in r["sesion"][0]["texto"] and a.hoja.columnas == []
    assert a.dialogo_modulo("ALEATORIOS")["error"] == "entrada"
    from app.modulos.calc_hoja import Aleatorios
    with pytest.raises(Exception, match="bloqueado"): Aleatorios.ejecutar(a.hoja, PARAMS_ALEA, a.proyecto)      # también por la vía directa
    d = Api(); d.cambiar_modo("demostracion")
    r = d.ejecutar_modulo("ALEATORIOS", PARAMS_ALEA)
    assert r["error"] is None and "SINTÉTICO / NO EMPRESARIAL" in r["sesion"][0]["texto"]
    assert d.hoja.n == 12 and "SINTÉTICO / NO EMPRESARIAL" in d.hoja.columnas[0].desc


# ---------- A5, E1-R7: marca en pestaña, resultados y exportaciones ----------
def proyecto_con_resultado(demo):
    a = Api()
    if demo: a.cambiar_modo("demostracion")
    a.pegar(-1, 0, DATOS)
    uid = a.ejecutar_modulo("BOXPLOT", {"y": 0, "grupo": 1})["resultado"]["uid"]
    return a, uid


def textos_xlsx(ruta):
    wb = openpyxl.load_workbook(ruta); return " | ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None), wb


ETQ = "SINTÉTICO / NO EMPRESARIAL"


@pytest.mark.parametrize("demo", [True, False])
def test_a5_marca_en_pestana_y_resultados(demo):
    a, uid = proyecto_con_resultado(demo)
    p = a.proyectos()["proyectos"][0]
    assert p["sintetico"] is demo and p["etiqueta"] == (ETQ if demo else None)
    assert a.resultado(uid).get("etiqueta") == (ETQ if demo else None)
    assert a.resultados()[0]["etiqueta"] == (ETQ if demo else None)
    assert (ETQ in a.texto_resultado(uid)) is demo
    r = a.repetir_resultado(uid); assert r["resultado"].get("etiqueta") == (ETQ if demo else None)
    assert "etiqueta" not in a.proyecto.resultados[0]                      # no se guarda: se deriva del modo, así los .dmike siguen idénticos


@pytest.mark.parametrize("demo", [True, False])
def test_a5_marca_en_exportaciones(demo, tmp_path):
    a, uid = proyecto_con_resultado(demo)
    a.pegar(-1, len(a.hoja.columnas), "x\tz\n1\t2\n2\t4\n3\t5\n4\t9\n5\t10")
    # .xlsx del resultado
    a._elegir = lambda *x, **k: str(tmp_path / "r.xlsx"); assert a.exportar_resultado(uid)["error"] is None
    txt, wb = textos_xlsx(tmp_path / "r.xlsx"); assert (ETQ in txt) is demo
    # .png del resultado
    a._elegir = lambda *x, **k: str(tmp_path / "g.png"); assert a.exportar_resultado(uid)["error"] is None
    from PIL import Image
    import base64, io
    original = Image.open(io.BytesIO(base64.b64decode(a.proyecto.resultados[0]["png"]))).size
    est = Image.open(tmp_path / "g.png").size
    assert (est[1] > original[1]) is demo and est[0] == original[0]
    if demo: assert Image.open(tmp_path / "g.png").convert("RGB").getpixel((3, 3)) == (0xB3, 0x26, 0x1E)       # franja de la etiqueta
    # formato (Excel)
    a.abrir_formato("AMEF"); a._elegir = lambda *x, **k: str(tmp_path / "f.xlsx"); assert a.exportar_formato("AMEF")["error"] is None
    assert (ETQ in textos_xlsx(tmp_path / "f.xlsx")[0]) is demo
    # Minitab: .xlsx (hoja AVISO, Datos limpia) y .csv (archivo vecino)
    a._elegir = lambda *x, **k: str(tmp_path / "m.xlsx"); a.exportar()
    txt, wb = textos_xlsx(tmp_path / "m.xlsx")
    assert (ETQ in txt) is demo and wb.sheetnames == (["Datos", "AVISO"] if demo else ["Datos"]) and ETQ not in " ".join(str(c.value) for row in wb["Datos"].iter_rows() for c in row)
    a._elegir = lambda *x, **k: str(tmp_path / "m.csv"); a.exportar()
    assert (tmp_path / "m_AVISO.txt").exists() is demo and ETQ not in (tmp_path / "m.csv").read_text(encoding="utf-8-sig")
    # Power BI
    a._elegir = lambda *x, **k: str(tmp_path); assert a.exportar_powerbi()["error"] is None
    carpeta = next(tmp_path.glob("D-MIKE_PowerBI_*"))
    csvs = list(carpeta.glob("*.csv")); assert csvs
    for f in csvs: assert (f.read_text(encoding="utf-8-sig").splitlines()[0].split(",")[-1] == "marca") is demo, f.name
    assert (ETQ in (carpeta / "LEEME.txt").read_text(encoding="utf-8-sig")) is demo
    assert (ETQ in (carpeta / "dim_proyecto.csv").read_text(encoding="utf-8-sig")) is demo


@pytest.mark.parametrize("demo", [True, False])
def test_a5_marca_en_hoja_de_verificacion(demo, tmp_path):
    a = Api()
    if demo: a.cambiar_modo("demostracion")
    a.pegar(-1, 0, DATOS)
    uid = a.ejecutar_modulo("ESTADISTICOS", {"y": 0})["resultado"]["uid"]
    a._elegir = lambda *x, **k: str(tmp_path / "v.xlsx"); r = a.exportar_verificacion(uid); assert r["error"] is None, r["sesion"]
    assert (ETQ in textos_xlsx(tmp_path / "v.xlsx")[0]) is demo


# ---------- E1-R5: persistencia y .dmike sin _modo ----------
def test_r5_modo_se_guarda_y_dmike_sin_modo_se_infiere_de_la_receta(tmp_path):
    a = Api(); a.ejecutar_receta(v1(tmp_path, "Proyecto planta 7"), "nuevo", None, "demostracion")
    ruta = str(tmp_path / "p.dmike"); archivo.guardar(ruta, a.hoja, a.proyecto, "0.26.0")
    b = Api(); b.abrir_proyecto(ruta); assert formato_modo(b)["valor"] == "demostracion" and formato_modo(b)["cambiado"] is True
    # sin «_modo» (como lo guardó la 0.25): se infiere de «_receta.origen» (o de su JSON)
    for origen, esperado in (("Caso maestro de D-MIKE", "demostracion"), ("Planta 7", "real")):
        p = Api().proyecto; rc = base(); rc["origen"] = origen
        p.formatos["_receta"] = {"folio": "", "ids": {}, "json": json.dumps(rc)}                     # forma de 0.25: sin clave «origen»
        assert MD.estado(p) == {"valor": esperado, "deteccion": "inferido_v1", "cambiado": False}
        p.formatos["_receta"]["origen"] = origen; assert MD.valor(p) == esperado                     # forma de 0.26
    assert MD.estado(Api().proyecto) == {"valor": "real", "deteccion": "inferido_v1", "cambiado": False}   # sin receta: real
    assert MD.valor(Api().proyecto) == "real"


# ---------- E1-R8: folio ----------
def test_r8_v1_con_folio_lo_ignora_y_avisa(tmp_path):
    ruta = v1(tmp_path, "Caso maestro (sintético)", folio="DMK-ABCD-EFGH-JKLM"); a = Api()
    r = a.abrir_receta(ruta)
    assert "folio" not in r["receta"] and r["receta"]["trae_folio"] and any("llave" in m["texto"] for m in r["sesion"]) and "actualizable" not in r["receta"]
    r = a.ejecutar_receta(ruta); assert r["error"] is None and any("folio" in m["texto"] and "llave" in m["texto"] for m in r["sesion"])
    rec = a.proyecto.formatos["_receta"]
    assert "folio" not in rec and "DMK-ABCD" not in json.dumps(rec) and "DMK-ABCD" not in json.dumps(a.proyecto.a_dict())
    b = Api(); b.ejecutar_receta(ruta, "datos"); assert "DMK-ABCD" not in json.dumps(b.proyecto.a_dict())


def test_a6_mismo_folio_sin_id_proyecto_no_actualiza_solo(tmp_path):
    ruta1, ruta2 = v1(tmp_path, "Caso maestro", folio="DMK-AAAA-BBBB-CCCC", nombre="a.json"), v1(tmp_path, "Caso maestro", folio="DMK-AAAA-BBBB-CCCC", nombre="b.json")
    a = Api(); a.ejecutar_receta(ruta1)
    r = a.abrir_receta(ruta2); assert "actualizable" not in r["receta"] and r["receta"]["abiertos"]
    assert a.ejecutar_receta(ruta2, "actualizar", None)["error"] == "proyecto"        # sin confirmación manual no se actualiza


# ---------- E1-R9: id_proyecto ----------
def test_a6_id_proyecto_se_guarda_y_se_propone_para_actualizar(tmp_path):
    ruta = v2(tmp_path, "demostracion"); a = Api()
    a.ejecutar_receta(ruta); assert a.proyecto.formatos["_receta"]["id_proyecto"] == UUID_A
    a.nuevo_proyecto()                                                              # otro proyecto abierto, sin id
    r = a.abrir_receta(ruta)["receta"]
    assert r["actualizable"]["id"] == a.proyectos()["proyectos"][0]["id"] and r["id_proyecto"] == UUID_A and len(r["abiertos"]) == 2
    n = len(a.proyectos()["proyectos"])
    # 0.27 (12.5, regla 2): la misma receta, byte por byte, ya está aplicada; para actualizar se usa una revisión nueva de la receta
    assert a.ejecutar_receta(ruta, "actualizar", r["actualizable"]["id"])["error"] == "ya_aplicada"
    ruta2 = v2(tmp_path, "demostracion", revision=1, nombre="r2b.json")
    out = a.ejecutar_receta(ruta2, "actualizar", r["actualizable"]["id"]); assert out["error"] is None and len(a.proyectos()["proyectos"]) == n
    assert a.proyecto.formatos["_receta"]["id_proyecto"] == UUID_A and a.proyecto.formatos["_receta"]["origen"] == "Diagnóstico en línea"


def test_a6_confirmacion_manual_con_v1_o_sin_coincidencia(tmp_path):
    a = Api(); a.ejecutar_receta(v1(tmp_path, "Caso maestro", nombre="a.json")); destino = a.proyectos()["proyectos"][0]["id"]
    ruta = v2(tmp_path, "demostracion", idp=UUID_B)                                 # v2 con id que no coincide con ningún proyecto abierto
    r = a.abrir_receta(ruta)["receta"]; assert "actualizable" not in r and r["abiertos"][0]["id"] == destino
    assert a.ejecutar_receta(ruta, "actualizar", destino)["error"] is None           # el usuario confirma a mano
    assert a.proyecto.formatos["_receta"]["id_proyecto"] == UUID_B                   # el proyecto elegido queda con ese id
    # v1: nunca se propone nada; se confirma a mano
    r1 = v1(tmp_path, "Caso maestro (segunda versión)", nombre="c.json"); assert "actualizable" not in a.abrir_receta(r1)["receta"]
    assert a.ejecutar_receta(r1, "actualizar", destino)["error"] is None
    assert a.proyecto.formatos["_receta"]["id_proyecto"] == UUID_B                   # una v1 no borra el id ya guardado


def test_a6_id_duplicado_se_rechaza_o_se_crea_sin_id(tmp_path):
    ruta = v2(tmp_path, "demostracion"); a = Api(); a.ejecutar_receta(ruta)
    r = a.ejecutar_receta(ruta); assert r["error"] == "id_conflicto" and len(a.proyectos()["proyectos"]) == 1
    r = a.ejecutar_receta(ruta, "nuevo", None, None, True); assert r["error"] is None and len(a.proyectos()["proyectos"]) == 2
    assert a.proyecto.formatos["_receta"]["id_proyecto"] is None
    # actualizar un proyecto que tiene OTRO id, o con un id que ya tiene otro proyecto abierto
    ruta_b = v2(tmp_path, "demostracion", idp=UUID_B, nombre="b.json"); a.ejecutar_receta(ruta_b)
    ids = {p["nombre"]: p["id"] for p in a.proyectos()["proyectos"]}
    docs = [d for d in a._docs if (d["proyecto"].formatos.get("_receta") or {}).get("id_proyecto") == UUID_A]
    assert a.ejecutar_receta(ruta_b, "actualizar", docs[0]["id"])["error"] == "id_conflicto"


def test_e13_r3_solo_datos_no_comprueba_ni_asigna_id(tmp_path):
    ruta = v2(tmp_path, "demostracion"); a = Api(); a.ejecutar_receta(ruta)
    r = a.ejecutar_receta(ruta, "datos"); assert r["error"] is None and len(a.proyectos()["proyectos"]) == 2
    assert not (a.proyecto.formatos.get("_receta") or {}).get("id_proyecto")


def test_e13_r4_id_proyecto_de_no_lanza_excepcion(tmp_path):
    import sqlite3
    ruta = v2(tmp_path, "demostracion"); a = Api(); a.ejecutar_receta(ruta)
    dm = str(tmp_path / "ok.dmike"); archivo.guardar(dm, a.hoja, a.proyecto, "0.26.0")
    assert archivo.id_proyecto_de(dm) == UUID_A
    os.remove(dm); assert archivo.id_proyecto_de(dm) is None                        # borrado
    (tmp_path / "vacio.dmike").write_bytes(b""); assert archivo.id_proyecto_de(str(tmp_path / "vacio.dmike")) is None
    (tmp_path / "basura.dmike").write_bytes(b"no es sqlite" * 50); assert archivo.id_proyecto_de(str(tmp_path / "basura.dmike")) is None
    for nombre, js in (("jsondanado", "{no es json"), ("sinreceta", json.dumps({"formatos": {}})), ("lista", "[1,2]"), ("formatos_nulo", json.dumps({"formatos": None}))):
        ruta_x = str(tmp_path / (nombre + ".dmike")); con = sqlite3.connect(ruta_x)
        con.execute("CREATE TABLE proyecto (clave TEXT PRIMARY KEY, json TEXT)"); con.execute("INSERT INTO proyecto VALUES ('estado', ?)", (js,)); con.commit(); con.close()
        assert archivo.id_proyecto_de(ruta_x) is None, nombre
    ruta_y = str(tmp_path / "sintabla.dmike"); sqlite3.connect(ruta_y).close(); assert archivo.id_proyecto_de(ruta_y) is None


def test_a6_id_de_proyecto_reciente_tambien_se_rechaza(tmp_path, monkeypatch):
    ruta = v2(tmp_path, "demostracion"); a = Api(); a.ejecutar_receta(ruta)
    dm = str(tmp_path / "guardado.dmike"); archivo.guardar(dm, a.hoja, a.proyecto, "0.26.0")
    assert archivo.id_proyecto_de(dm) == UUID_A and archivo.id_proyecto_de(str(tmp_path / "no_existe.dmike")) is None
    b = Api(); monkeypatch.setattr(b, "recientes", lambda: [{"ruta": dm, "nombre": "Guardado"}])
    r = b.abrir_receta(ruta)["receta"]; assert r["conflicto_id"]["nombre"] == "Guardado"
    out = b.ejecutar_receta(ruta); assert out["error"] == "id_conflicto" and "reciente" in out["sesion"][0]["texto"] and b.hoja.columnas == []
    assert b.ejecutar_receta(ruta, "nuevo", None, None, True)["error"] is None


# ---------- E1-R10, A7, A8: migración ----------
@pytest.fixture(scope="module")
def dmike_025(tmp_path_factory):
    ruta = tmp_path_factory.mktemp("m") / "caso_025.dmike"
    with gzip.open(S / "caso_integral_v0.25.0.dmike.gz", "rb") as g, open(ruta, "wb") as f: shutil.copyfileobj(g, f)
    return str(ruta), json.loads((S / "caso_integral_v0.25.0.esperado.json").read_text(encoding="utf-8"))


def test_a7_dmike_025_abre_con_resultados_huellas_y_disenos_iguales(dmike_025):
    ruta, esp = dmike_025
    a = Api(); r = a.abrir_proyecto(ruta); assert r["error"] is None
    res = [a.resultado(x["uid"]) for x in a.resultados()]
    assert len(res) == esp["n_resultados"] == 17 and a.hoja.n == esp["n_filas"] and [c.nombre for c in a.hoja.columnas] == esp["columnas"]
    assert [x["definicion"]["congelado"]["huella"] for x in res] == esp["huellas"]
    assert [hashlib.sha256(json.dumps(x["bloques"], sort_keys=True, ensure_ascii=False).encode()).hexdigest() for x in res] == esp["sha_bloques"]
    assert sorted(a.proyecto.formatos) == esp["claves_formatos"] and "_modo" not in a.proyecto.formatos      # abrir no escribe nada
    assert all(a.reproducir_resultado(x["uid"])["error"] is None for x in a.resultados())                     # la huella sigue comprobando
    assert MD.estado(a.proyecto) == {"valor": "demostracion", "deteccion": "inferido_v1", "cambiado": False}
    assert all(x["etiqueta"] == ETQ for x in res)


def test_a8_folio_previo_en_receta_se_conserva(dmike_025, tmp_path):
    ruta, esp = dmike_025; assert esp["folio_receta"] == "DMK-PRUEBA-0001"
    a = Api(); a.abrir_proyecto(ruta)
    assert a.proyecto.formatos["_receta"]["folio"] == "DMK-PRUEBA-0001"
    nuevo = str(tmp_path / "re.dmike"); archivo.guardar(nuevo, a.hoja, a.proyecto, "0.26.0")
    b = Api(); b.abrir_proyecto(nuevo); assert b.proyecto.formatos["_receta"]["folio"] == "DMK-PRUEBA-0001"
    # actualizar con una receta nueva (sin folio) tampoco lo quita ni lo reescribe
    ruta_r = v1(tmp_path, "Caso maestro (segunda versión)", nombre="act.json"); id_doc = b.proyectos()["proyectos"][0]["id"]
    assert b.ejecutar_receta(ruta_r, "actualizar", id_doc)["error"] is None
    assert b.proyecto.formatos["_receta"]["folio"] == "DMK-PRUEBA-0001"
    # una receta con otro folio: el de la receta se ignora, el que ya estaba en el .dmike queda como estaba
    ruta_f = v1(tmp_path, "Caso maestro (tercera versión)", folio="DMK-OTRO-OTRO-OTRO", nombre="otro.json")
    assert b.ejecutar_receta(ruta_f, "actualizar", id_doc)["error"] is None and b.proyecto.formatos["_receta"]["folio"] == "DMK-PRUEBA-0001"
    assert "DMK-OTRO" not in json.dumps(b.proyecto.a_dict())


def test_r10_migracion_sin_cambio_de_tabla(tmp_path):
    a = Api(); a.ejecutar_receta(v2(tmp_path, "demostracion"))
    ruta = str(tmp_path / "n.dmike"); archivo.guardar(ruta, a.hoja, a.proyecto, "0.26.0")
    con = sqlite3.connect(ruta)
    assert dict(con.execute("SELECT clave, valor FROM meta"))["version_formato"] == "1" == str(archivo.VERSION_FORMATO)
    assert {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")} == {"meta", "columnas", "celdas", "proyecto"}
    d = json.loads(con.execute("SELECT json FROM proyecto").fetchone()[0])
    assert set(d) == {"perfil", "version", "fijos", "modulos", "resultados", "historial", "formatos"}                 # mismas claves de nivel superior
    assert {"_modo", "_receta"} <= set(d["formatos"])                                                                  # las nuevas viven dentro de formatos


def _codigo_025(tmp_path):
    """Extrae el código de la 0.25.0 desde el historial de git. Se omite si el historial no está disponible (clon superficial)."""
    raiz = Path(__file__).resolve().parent.parent
    for rev in ("025edde", "f05296d"):
        try:
            if "0.25.0" not in subprocess.run(["git", "show", f"{rev}:app/api.py"], cwd=raiz, capture_output=True, text=True, check=True).stdout.split("VERSION = ")[1][:12]: continue
            dest = tmp_path / "viejo"; dest.mkdir()
            tar = subprocess.run(["git", "archive", rev, "app"], cwd=raiz, capture_output=True, check=True).stdout
            subprocess.run(["tar", "-x", "-C", str(dest)], input=tar, check=True)
            return dest
        except (subprocess.CalledProcessError, FileNotFoundError, IndexError): continue
    return None


def test_a7_dmike_026_abre_en_025_sin_error(tmp_path):
    viejo = _codigo_025(tmp_path)
    if viejo is None: pytest.skip("no hay historial de git con la 0.25.0")
    a = Api(); a.ejecutar_receta(v2(tmp_path, "demostracion"))
    ruta = str(tmp_path / "nuevo.dmike"); archivo.guardar(ruta, a.hoja, a.proyecto, "0.26.0")
    codigo = ("import sys; sys.path.insert(0, sys.argv[1]); from app.api import Api; import app.api as A\n"
              "a = Api(); r = a.abrir_proyecto(sys.argv[2]); print(A.VERSION, r['error'], a.hoja.n, len(a.hoja.columnas), sorted(k for k in a.proyecto.formatos if k.startswith('_')))")
    out = subprocess.run([sys.executable, "-I", "-c", codigo, str(viejo), ruta], capture_output=True, text=True, timeout=120)
    assert out.returncode == 0, out.stderr[-500:]
    assert out.stdout.split()[:2] == ["0.25.0", "None"] and "_modo" in out.stdout and "_receta" in out.stdout


# ---------- interfaz ----------
def test_interfaz_muestra_modo_y_permite_un_cambio(pagina, tmp_path):
    ruta = v1(tmp_path, "Proyecto planta 7", folio="DMK-ABCD-EFGH-JKLM")
    pagina.api._elegir = lambda *x, **k: ruta
    pagina.evaluate("ACCIONES_UI.abrir_proyecto()")
    pagina.wait_for_selector(".dialogo:has-text('¿Ejecutar la receta?')")
    cuerpo = pagina.inner_text(".dialogo")
    assert "Modo del proyecto: Real" in cuerpo and "inferido" in cuerpo and "ignora" in cuerpo and "DMK-ABCD" not in cuerpo
    assert pagina.locator("#receta-modo").count() == 1
    pagina.select_option("#receta-modo", "demostracion")
    pagina.click(".dialogo .primario")
    pagina.wait_for_function("document.querySelector('#sesion').textContent.includes('Receta ejecutada')", timeout=60000)
    assert MD.valor(pagina.api.proyecto) == "demostracion" and formato_modo(pagina.api)["cambiado"] is True
    pagina.wait_for_selector("#docs .etiqueta-sint")
    et = pagina.locator("#docs .doc-tab.activo .etiqueta-sint")          # E1.3-R2: insignia corta; el texto completo en title y aria-label
    assert et.inner_text() == "SINTÉTICO" and et.get_attribute("title") == ETQ and et.get_attribute("aria-label") == ETQ
    assert float(et.evaluate("e => getComputedStyle(e).fontSize").replace("px", "")) <= 9


def test_interfaz_v2_declarado_sin_selector_y_actualizacion_manual(pagina, tmp_path):
    ruta = v2(tmp_path, "real"); pagina.api.ejecutar_receta(ruta); pagina.evaluate("pintarDocs()")
    assert pagina.locator("#docs .etiqueta-sint").count() == 0                       # un proyecto real nunca lleva la etiqueta
    pagina.api._elegir = lambda *x, **k: v1(tmp_path, "Planta", nombre="v1.json")
    pagina.evaluate("ACCIONES_UI.abrir_proyecto()")
    pagina.wait_for_selector(".dialogo:has-text('¿Ejecutar la receta?')")
    assert pagina.locator("#receta-destino").count() == 1 and pagina.locator("#receta-actualizar").count() == 1
    assert "confirma a mano" in pagina.inner_text(".dialogo")
    pagina.api._elegir = lambda *x, **k: ruta
    pagina.evaluate("document.querySelectorAll('.velo').forEach(v => v.remove())")
    pagina.evaluate("ACCIONES_UI.abrir_proyecto()")
    pagina.wait_for_selector(".dialogo:has-text('¿Ejecutar la receta?')")
    assert pagina.locator("#receta-modo").count() == 0 and "declarado por la receta" in pagina.inner_text(".dialogo")
    t = pagina.inner_text(".dialogo"); assert "mismo id_proyecto" in t, t; assert pagina.inner_text(".dialogo .primario") == "Crear nuevo sin id_proyecto"


def test_interfaz_resultado_de_demostracion_lleva_etiqueta(pagina):
    pagina.api.cambiar_modo("demostracion"); pagina.api.pegar(-1, 0, DATOS)
    r = pagina.api.ejecutar_modulo("BOXPLOT", {"y": 0, "grupo": 1})["resultado"]
    pagina.evaluate("(res) => abrirPestana('res:' + res.uid, res)", r)
    pagina.wait_for_selector("#vista .vista-barra .etiqueta-sint")
    assert pagina.inner_text("#vista .vista-barra .etiqueta-sint") == ETQ
    real = Api(); real.pegar(-1, 0, DATOS); rr = real.ejecutar_modulo("BOXPLOT", {"y": 0, "grupo": 1})["resultado"]
    assert "etiqueta" not in rr
