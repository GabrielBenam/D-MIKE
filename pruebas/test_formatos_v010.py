"""Formatos de v0.10, catálogo de indicadores, tokens y exportación a Power BI. Datos SINTÉTICOS salvo los
ejemplos del curso citados."""
import csv, json, re
from pathlib import Path
from openpyxl import load_workbook
from app.api import Api
from app.nucleo.tokens import TOKENS, tema_power_bi
from test_modulos import SINT


def fmt(api, fid, cab=None, tablas=None):
    d = api.abrir_formato(fid)["datos"]
    if cab: d["cabecera"].update(cab)
    for t, filas in (tablas or {}).items():
        for i, f in enumerate(filas):
            if i < len(d["tablas"][t]): d["tablas"][t][i].update(f)
            else: d["tablas"][t].append(f)
    return api.guardar_formato(fid, d)["calc"]


def test_oee_formato_ejemplos_del_curso():
    c = fmt(Api(), "OEE", tablas={"horas": [
        {"periodo": "7-8", "total": "60", "paro": "10", "muerto": "10", "capacidad": "5", "produccion": "195", "rechazo": "2"},
        {"periodo": "8-9", "total": "60", "paro": "0", "muerto": "35", "capacidad": "5", "produccion": "120", "rechazo": "1"},
        {"periodo": "9-10", "total": "60", "paro": "0", "muerto": "5", "capacidad": "7", "produccion": "370", "rechazo": "3"}]})
    assert [f["oee"] for f in c["calc"]["horas"][:3]] == ["77.2%", "39.7%", "87.4%"]          # EJ-TABLERO-OEE
    assert c["clases"]["horas"][1]["disponibilidad"] == "no" and c["cab"]["oee_calc"].startswith("67.5% del tiempo disponible")
    c = fmt(Api(), "OEE", tablas={"horas": [{"total": "480", "paro": "30", "muerto": "60", "capacidad": "2", "produccion": "650", "rechazo": "50"}]})
    assert c["calc"]["horas"][0]["oee"] == "66.7%" and c["cab"]["oee_calc"].startswith("66.7% del tiempo disponible")       # EJ-OEE-CURSO


def test_takt_y_kanban():
    c = fmt(Api(), "TAKT", cab={"disp_simple": "720", "demanda_simple": "45"})
    assert c["cab"]["takt_simple_calc"].startswith("16 ")                                    # EJ-TAKT
    a = Api(); d = a.abrir_formato("KANBAN")["datos"]
    assert [f["mes"] for f in d["tablas"]["demanda"]][:2] == ["Enero", "Febrero"] and len(d["tablas"]["demanda"]) == 12
    demandas = [1000, 1100, 950, 1200, 1050, 980, 1150, 1020, 990, 1300, 900, 1008]           # sintéticas
    for f, v in zip(d["tablas"]["demanda"], demandas): f["demanda"] = str(v)
    d["cabecera"].update(entrega="2", ubicaciones="1", capacidad="50")
    c = a.guardar_formato("KANBAN", d)["calc"]
    assert c["cab"]["cont_calc"].isdigit() and c["cab"]["prom_calc"] == "1,054"


def test_charter_sipoc_cinco_pq():
    a = Api()
    c = fmt(a, "CHARTER", cab={"caso_negocio": "Caso sintético"}, tablas={"objetivos": [{"ctq": "Tiempo", "ahorro": "1500"}],
                                                                         "equipo": [{"rol": "Líder", "nombre": "Persona A"}]})
    assert any("champion o sponsor" in x for x in c["avisos"]) and "1,500" in c["bloques"][0]["texto"]
    c = fmt(a, "SIPOC", tablas={"entradas": [{"proveedor": "Almacén", "entrada": "Material", "tipo": "X"}]})
    assert any("sin requisito" in x for x in c["avisos"])
    c = fmt(a, "CINCO-PQ", cab={"problema": "La máquina está descompuesta"},
            tablas={"porques": [{"respuesta": "Se fundieron los fusibles"}, {"respuesta": "Las máquinas se sobrecalientan"},
                                {"respuesta": "No existe programa de mantenimiento"}]})              # EJ-5PQ, recortado
    assert c["bloques"][0]["texto"].startswith("Lectura inversa: No existe programa de mantenimiento → entonces")


def test_ishikawa_diagrama_y_exportacion(tmp_path):
    a = Api(); a.fijar_datos({"nombre_proyecto": "Proyecto sintético"})
    c = fmt(a, "ISHIKAWA", cab={"efecto": "Datos mal capturados"}, tablas={"causas": [
        {"categoria": "Métodos", "causa": "Formato ambiguo", "verificada": "Sí, es causa"}, {"categoria": "Mano de obra", "causa": "Sin entrenamiento"},
        {"causa": "Causa sin categoría"}]})
    assert "<svg" in c["bloques"][0]["svg"] and "png" not in c and any("sin categoría" in x for x in c["avisos"])
    ruta = str(tmp_path / "ish.xlsx"); a.exportar_formato("ISHIKAWA", ruta)
    assert len(load_workbook(ruta)["ISHIKAWA"]._images) == 1


def test_amef_escalas_desplegables(tmp_path):
    a = Api(); esp = a.abrir_formato("AMEF")["espec"]
    sev = next(c for c in esp["tablas"][0]["columnas"] if c["campo"] == "sev")
    assert sev["tipo"] == "opcion" and sev["opciones"][9] == ["9", "9 · muy alta"]
    c = fmt(a, "AMEF", tablas={"filas": [{"falla": "F", "sev": "9", "occ": "2", "det": "3"}]})
    assert c["calc"]["filas"][0]["rpn"] == 54
    ruta = str(tmp_path / "a.xlsx"); a.exportar_formato("AMEF", ruta)
    vals = [x.value for fila in load_workbook(ruta)["AMEF"].iter_rows() for x in fila]
    assert 9 in vals and 54 in vals                                                          # calificaciones como número


def test_catalogo_kpis():
    cat = Api().catalogo_kpis()
    ids = [k["id"] for k in cat]
    assert len(cat) > 100 and len(set(ids)) == len(ids) and "OEE" in ids
    assert all(k["sentido"] in ("<=", ">=") for k in cat)
    base = [k for k in cat if k["fuente"].startswith("BASE")]
    assert len(base) >= 20 and all(re.search(r"\d\d_|PL\d\d", k["fuente"]) for k in base)


def test_tablero_tarjetas_y_advertencia():
    a = Api()
    c = fmt(a, "TABLERO", cab={"semanas": "3"}, tablas={"kpi": [
        {"kpi_id": "OEE", "indicador": "OEE", "unidad": "%", "sentido": ">=", "meta": "85", "advertencia": "80", "s1": "70", "s2": "82", "s3": "88"},
        {"indicador": "Scrap", "sentido": "<=", "meta": "2", "advertencia": "3", "tarjeta": "No", "s1": "2.5"}]})
    k = c["clases"]["kpi"][0]
    assert (k["s1"], k["s2"], k["s3"]) == ("no", "advertencia", "ok") and c["clases"]["kpi"][1]["s1"] == "advertencia"
    t = c["tarjetas"]
    assert len(t) == 1 and t[0]["valor"] == "88" and t[0]["estado"] == "ok" and t[0]["cambio"] == "+6" and t[0]["meta"] == "≥ 85"
    c = fmt(a, "TABLERO", tablas={"kpi": [{"indicador": "X", "sentido": ">=", "meta": "85", "advertencia": "90"}]})
    assert any("límite de advertencia" in x for x in c["avisos"])


def test_tokens_coinciden_con_estilos():
    css = (Path(__file__).parent.parent / "app" / "ui" / "estilos.css").read_text(encoding="utf-8")
    raiz = css[css.index(":root"):css.index("}", css.index(":root"))]
    for k in ("papel", "panel", "superficie", "linea", "tinta", "tinta-suave", "acento", "acento-suave"):
        assert f"--{k}: {TOKENS[k]}" in raiz
    t = tema_power_bi()
    assert t["dataColors"][0] == TOKENS["acento"] and "Montserrat" in t["textClasses"]["title"]["fontFace"]


def test_exportar_powerbi(tmp_path):
    a = Api(); a.pegar(-1, 0, SINT); a.fijar_datos({"nombre_proyecto": "Proyecto sintético/1"})
    fmt(a, "TABLERO", cab={"semanas": "3", "inicio": "05/10/2026", "col_fecha": "fecha"}, tablas={
        "kpi": [{"kpi_id": "TMU", "indicador": "Tiempo medio", "sentido": "<=", "meta": "13", "col": "tiempo"},
                {"indicador": "Manual", "sentido": ">=", "meta": "25", "s1": "21", "s2": "26"}],
        "wip": [{"unidad": "U-1", "estado": "En revisión"}]})
    fmt(a, "AMEF", tablas={"filas": [{"falla": "F", "sev": "5", "occ": "4", "det": "2"}]})
    fmt(a, "OEE", tablas={"horas": [{"total": "480", "paro": "30", "muerto": "60", "capacidad": "2", "produccion": "650", "rechazo": "50"}]})
    a.ejecutar_modulo("BOXPLOT", {"y": 4, "grupo": 2})
    r = a.exportar_powerbi(str(tmp_path))
    assert r["error"] is None
    d = tmp_path / "D-MIKE_PowerBI_Proyecto sintético_1"
    leer = lambda n: list(csv.DictReader(open(d / n, encoding="utf-8-sig")))
    hechos = leer("hecho_indicador.csv")
    assert {h["indicador_id"] for h in hechos} == {"TMU", "KPI2"} and hechos[0]["inicio_semana"] == "2026-10-05"
    assert leer("dim_indicador.csv")[0]["sentido"] == "menor_mejor"
    assert leer("oee.csv")[0]["oee"].startswith("0.666")
    assert leer("amef.csv")[0]["rpn"] == "40" and leer("datos_hoja.csv")[0]["fecha"] == "2026-10-05"
    assert leer("trabajos_en_proceso.csv")[0]["estado"] == "En revisión" and len(leer("resultados.csv")) == 1
    assert json.loads((d / "tema_dmike.json").read_text(encoding="utf-8"))["name"] == "D-MIKE"
    assert "Obtener datos" in (d / "LEEME.txt").read_text(encoding="utf-8-sig")
    assert Api().exportar_powerbi.__self__ is not None
    b = Api(); b._elegir = lambda *x, **k: None
    assert "cancelada" in b.exportar_powerbi()["sesion"][0]["texto"]


def test_formatos_v013():
    a = Api()
    c = fmt(a, "SMED", cab={"unidad": "min"}, tablas={"ops": [
        {"operacion": "Retirar molde", "real": "5", "pot": "3", "clase": "Interna"}, {"operacion": "Buscar herramienta", "real": "4", "pot": "0", "clase": "Desperdicio"},
        {"operacion": "Limpiar", "real": "3", "pot": "4", "clase": "Externa"}]})
    assert c["calc"]["ops"][2]["acumulado"] == "12" and c["cab"]["total_real"] == "12 min" and c["cab"]["total_pot"] == "7 min"
    assert c["cab"]["desp_real"] == "4 min" and c["cab"]["ahorro"].startswith("5 min") and any("potencial supera" in x for x in c["avisos"])
    c = fmt(a, "GANTT", cab={"semana1": "05/10/2026"}, tablas={"definicion": [
        {"actividad": "Charter", "inicio": "05/10/2026", "fin": "16/10/2026"}, {"actividad": "Largo", "inicio": "12/10/2026", "fin": "31/12/2026"}]})
    assert c["calc"]["definicion"][0]["dias"] == "12" and c["calc"]["definicion"][0]["semanas"] == "1 a 2" and any("semana 12" in x for x in c["avisos"])
    assert "<svg" in c["bloques"][0]["svg"]
    c = fmt(a, "MUDAS", tablas={"mudas": [{"impacto": "Alto"}, {"impacto": "Medio"}]})
    assert any("Impacto alto: Sobreproducción" in x for x in c["avisos"]) and len(a.abrir_formato("MUDAS")["datos"]["tablas"]["mudas"]) == 8
    c = fmt(a, "KATA", cab={"objetivo": "producir mejor"}, tablas={"problemas": [{"problema": "p1", "estado": "Haciendo"}, {"problema": "p2", "estado": "Hecho"}]})
    assert c["cab"]["resumen"].startswith("Por resolver 0 · Haciendo 1") and any("medible" in x for x in c["avisos"])
    assert any("folios repetidos" in x for x in fmt(a, "TARJETAS", tablas={"op": [{"folio": "1", "oportunidad": "x"}, {"folio": "1", "oportunidad": "y"}]})["avisos"])
    assert any("Faltan 6 de las 7" in x for x in fmt(a, "A3", cab={"antecedentes": "algo"})["avisos"])
    assert fmt(a, "KAIZEN", cab={"fecha_inicio": "05/10/2026", "fecha_promesa": "10/10/2026", "fecha_fin": "20/10/2026"})["avisos"]
    assert fmt(a, "CUATRO-Q", cab={"area": "A"}, tablas={"plan": [{"accion": "hacer"}]})["avisos"]


def test_formatos_v014():
    a = Api()
    c = fmt(a, "KANO", tablas={"atributos": [{"requisito": "Entrega a tiempo", "respuesta_con": "Neutral", "respuesta_sin": "Insatisfecho"},
                                             {"requisito": "Rastreo", "respuesta_con": "Satisfecho", "respuesta_sin": "No le importa"},
                                             {"requisito": "Raro", "respuesta_con": "Insatisfecho", "respuesta_sin": "Satisfecho"}]})
    assert [x.get("clasificacion") for x in c["calc"]["atributos"][:3]] == ["Básico", "Deleite", "Revisar"] and c["cab"]["resumen"] == "Básico 1 · Deleite 1"
    d = a.abrir_formato("QFD")["datos"]                                     # v0.21: QFD completo (importancia = peso relativo % × relación)
    d["tablas"]["tecnicos"][0]["nombre"], d["tablas"]["tecnicos"][1]["nombre"] = "Tiempo", "Costo"
    d["tablas"]["requisitos_cliente"] = [{"requisito": "Rápido", "importancia": "5", "rel_1": "9", "rel_2": "1"}, {"requisito": "Barato", "importancia": "3", "rel_2": "9"}, {"requisito": "Bonito", "importancia": "2"}]
    c = a.guardar_formato("QFD", d)["calc"]
    assert [x.get("imp_abs") for x in c["calc"]["tecnicos"][:2]] == ["450", "320"] and any("Bonito" in x for x in c["avisos"])
    c = fmt(a, "CTQ-TREE", cab={"necesidad": "Servicio excelente"}, tablas={"ramas": [{"impulsor": "Espera", "requisito": "Menos de 5 min"}, {"impulsor": "Amable", "requisito": "Saludo cordial"}]})
    assert c["clases"]["ramas"][1]["requisito"] == "advertencia" and "<svg" in c["bloques"][0]["svg"]
    c = fmt(a, "VSM", tablas={"modelos": [{"modelo": "M1", "tc_proceso_1": "30", "tc_proceso_2": "x", "vol_ene": "100", "vol_feb": "50"}, {"modelo": "M2", "tc_proceso_1": "20", "tc_proceso_2": "x"},
                                          {"modelo": "M3", "tc_proceso_1": "10", "tc_proceso_2": "5"}],
                              "resultados": [{"antes": "100", "despues": "60"}, {}, {}, {}, {"antes": "10", "despues": "4"}, {"antes": "6", "despues": "2"}, {"antes": "4", "despues": "2"}]})
    assert [x.get("familia") for x in c["calc"]["modelos"][:3]] == ["A", "A", "B"] and c["calc"]["modelos"][0]["vol_total"] == "150"
    assert c["calc"]["resultados"][0]["ahorro"] == "40 (40%)" and c["cab"]["vueltas"] == "13.2 / 33"
    c = fmt(a, "TIEMPOS", tablas={"elementos": [{"elemento": "Tomar", **{f"m{k}": v for k, v in enumerate(["5", "4", "4", "6", "5", "5"], 1)}}, {"elemento": "Soltar", "m1": "3", "m2": "2"}]})
    assert c["calc"]["elementos"][0]["repetido"] == "4" and c["calc"]["elementos"][0]["moda"] == "5" and c["calc"]["elementos"][1]["repetido"] == "no se repite" and c["cab"]["ciclo"] == "6 s"
    c = fmt(a, "BALANCE", cab={"t_disponible": "27000", "demanda": "450"}, tablas={"operaciones": [{"codigo": "OP10", "tiempo": "55"}, {"codigo": "OP20", "tiempo": "70"}]})
    assert c["cab"]["takt"] == "60 s" and c["cab"]["operadores"] == "2.083" and c["clases"]["operaciones"][1]["pct"] == "no" and any("OP20" in x for x in c["avisos"])
    c = fmt(a, "TPM", cab={"t_operacion": "400"}, tablas={"fallas": [{"descripcion": "Banda", "reparacion": "2", "tipo": "Mecánica"}, {"descripcion": "Sensor", "reparacion": "4", "tipo": "Eléctrica"}]})
    assert c["cab"]["mtbf"] == "200" and c["cab"]["mttr"] == "3" and c["cab"]["disp"].startswith("98.5% = MTBF / (MTBF + MTTR)")
    assert "reactivos 1" in fmt(a, "POKAYOKE", tablas={"casos": [{"error": "Pieza invertida", "fase": "Reactiva", "dispositivo": "Sensor"}]})["cab"]["resumen"]
    c = fmt(a, "CINCO-S", tablas={"s": [{"calif": "4"}, {"calif": "3"}, {"calif": "3"}, {"calif": "1"}, {"calif": "1"}]})
    assert c["cab"]["total"] == "12 de 20 (60%)" and any("3S" in x for x in c["avisos"]) and len(a.abrir_formato("CINCO-S")["datos"]["tablas"]["s"]) == 5
    assert "1 anormales" in fmt(a, "GEMBA", cab={"tema": "Flujo"}, tablas={"obs": [{"observacion": "Espera de material", "condicion": "Anormal", "desperdicio": "Espera"}]})["cab"]["resumen"]
    c = fmt(a, "SCRUM", cab={"fecha_inicio": "05/10/2026", "fecha_entrega": "16/10/2026", "horas_disponibles": "40"}, tablas={"historias": [{"historia": "x", "horas": "50"}],
                                                                                                                       "burndown": [{"dia": "1", "restante": "50"}, {"dia": "2", "restante": "48"}]})
    assert c["cab"]["dias"] == "12" and "125%" in c["cab"]["carga"] and c["calc"]["burndown"][9]["ideal"] == "0" and "<svg" in c["bloques"][0]["svg"]
    assert fmt(a, "LSW", tablas={"diario": [{"tarea": "Revisar tablero"}]})["clases"]["diario"][0]["hora"] == "falta"
    c = fmt(a, "TRABAJO-EST", cab={"takt": "30"}, tablas={"pasos": [{"paso": "Cargar", "manual": "10", "auto": "15", "caminata": "3"}, {"paso": "Inspeccionar", "manual": "20", "caminata": "2"}]})
    assert c["calc"]["pasos"][1]["acumulado"] == "35" and any("supera el takt" in x for x in c["avisos"]) and "polyline" in c["bloques"][0]["svg"]


def test_formatos_v015_contra_el_curso():
    a = Api()
    # Muestreo (15_GB MUE-DISCRETA y MUE-CONTINUA)
    c = fmt(a, "MUESTREO", cab={"tipo": "discreto", "poblacion": "400", "p": "0.09", "d": "0.01", "confianza": "90"})
    assert c["cab"]["n_sin"].startswith("2,203") and c["cab"]["n"].startswith("339")
    c = fmt(a, "MUESTREO", cab={"tipo": "continuo", "poblacion": "3200", "s": "0.458", "d": "0.100", "confianza": "95", "periodo": "480", "p": ""})
    assert c["cab"]["n"] == "81" and "5.93" in c["cab"]["frecuencia"]
    # Finanzas (23_BB ejercicios 1 a 7; PL18: 1,560 / 1,680 / 2,175 / 4,079; 10.0% a 23.1%)
    ent = {"demanda": ["239", "239", "272", "272"], "capacidad": ["239", "239", "272", "272"], "precio": ["65"] * 4, "material": ["50", "50", "50", "43"],
           "conversion": ["50625", "47625", "47625", "47625"], "dias": ["25"] * 4}
    d = a.abrir_formato("FINANZAS")["datos"]
    for f, (k, _) in zip(d["tablas"]["estados"], __import__("app.modulos.formatos5", fromlist=["x"]).INDICADORES):
        for j in range(4): f[f"e{j + 1}"] = ent[k][j]
    c = a.guardar_formato("FINANZAS", d)["calc"]
    res = {r: f for r, f in zip(("vendidas", "ingresos", "c_material", "c_conv", "c_conv_pza", "utilidad", "margen", "equilibrio"), c["calc"]["resultados"])}
    assert [res["utilidad"][f"e{j}"] for j in range(1, 5)] == ["1,560.00", "1,680.00", "2,175.00", "4,079.00"]
    assert res["margen"]["e1"] == "10.0%" and res["margen"]["e4"] == "23.1%"
    assert res["equilibrio"]["e1"] == "135" and res["equilibrio"]["e2"] == "127"                        # ejercicios 1 y 2
    assert res["c_conv_pza"]["e2"] == "7.97" and res["c_conv_pza"]["e3"] == "7.00"                     # ejercicios 3 y 4
    c = fmt(a, "FINANZAS", cab={"ventas_mv": "100000", "var_mv": "25000"})
    assert c["cab"]["margen_variable"].startswith("75,000.00 (75%")                                    # FIN-MARGEN-VARIABLE
    # Desempeño (16_GB DES-EJEMPLO-INTEGRAL y tabla de referencia)
    from app.modulos.formatos5 import sigma_de_dpmo, dpmo_de_sigma
    lp, cp = sigma_de_dpmo(13986); assert (round(lp, 2), round(cp, 2)) == (2.20, 3.70)
    for dpmo, s in ((66807, 3.0), (22750, 3.5), (6210, 4.0), (1350, 4.5), (233, 5.0), (3.4, 6.0)): assert round(sigma_de_dpmo(dpmo)[1], 1) == s
    assert round(dpmo_de_sigma(6.0), 1) == 3.4
    c = fmt(a, "DESEMPENO", cab={"conv_dato": "rendimiento", "conv_valor": "0.993790"}, tablas={"pasos": [
        {"paso": "Filtro", "unidades": "100", "oportunidades": "2", "defectos": "5"}, {"paso": "Agua", "unidades": "100", "oportunidades": "1", "defectos": "1"}]})
    assert c["calc"]["pasos"][0]["dpmo"] == "25,000" and c["cab"]["dpmo_total"] == "20,000" and c["cab"]["prioridad"].startswith("Filtro")
    assert "sigma corto plazo 4" in c["cab"]["conv_resultado"]


def test_formatos_v015_logica():
    a = Api()
    c = fmt(a, "DEF-PROBLEMA", cab={"sujeto": "la maleta", "predicado": "se perdió en el vuelo"})
    assert c["cab"]["enunciado"] == "La maleta se perdió en el vuelo." and any("presente" in x for x in c["avisos"]) and any("dónde" in x for x in c["avisos"])
    c = fmt(a, "DEF-PROBLEMA", cab={"sujeto": "el piso", "predicado": "está mojado", "donde": "pasillo 3", "cuando": "turno 2", "magnitud": "3 charcos por turno"})
    assert not c["avisos"]
    c = fmt(a, "HOSHIN", tablas={"objetivos": [{"perspectiva": "Finanzas", "nombre": "Incrementar ROI", "metrica": "ROI %", "base": "7", "meta": "12", "actual": "9.5"}],
                                 "proyectos": [{"nombre": "Reducir inventario", "metrica": "M$", "base": "24", "meta": "16", "actual": "20", "vinculo": "3"}]})
    assert c["calc"]["objetivos"][0]["avance"] == "50%" and c["calc"]["proyectos"][0]["avance"] == "50%"
    assert any("proyectos sin estrategias" in x for x in c["avisos"]) and c["clases"]["proyectos"][0]["vinculo"] == "mal"
    c = fmt(a, "PMAPA", tablas={"pasos": [{"paso": "Recibir", "salidas": "orden", "valor": "No"}, {"paso": "Cortar"}],
                                "entradas": [{"paso_n": "1", "entrada": "humedad", "clase": "Ruido (N)"}, {"paso_n": "5", "entrada": "x"}]})
    assert any("número de paso no existe" in x for x in c["avisos"]) and any("Pasos sin salida (Y): 2" in x for x in c["avisos"]) and "<svg" in c["bloques"][0]["svg"]
    c = fmt(a, "TOC", cab={"demanda": "50"}, tablas={"recursos": [{"recurso": "Corte", "capacidad": "60"}, {"recurso": "Pintura", "capacidad": "45"}]})
    assert c["cab"]["restriccion"].startswith("Pintura") and c["cab"]["tipo"].startswith("Interna") and c["clases"]["recursos"][1]["uso"] == "no"
    assert fmt(a, "TOC", cab={"demanda": "30"})["cab"]["tipo"].startswith("Externa")
    c = fmt(a, "ESPAGUETI", tablas={"antes": [{"distancia": "12", "veces": "2", "retorno": "Sí"}, {"distancia": "6"}], "despues": [{"distancia": "9"}]})
    assert c["cab"]["total_antes"] == "30 m" and c["cab"]["reduccion"].startswith("21 m (70%")
    c = fmt(a, "FLUJO", cab={"takt": "60"}, tablas={"pasos": [{"tc": "50", "confiable": "Sí"}, {"tc": "52", "confiable": "Sí"}, {"tc": "30", "confiable": "Sí"}, {"tc": "31", "confiable": "No"}]})
    assert [f.get("conexion") for f in c["calc"]["pasos"][:3]] == ["Flujo continuo", "Supermercado con Kanban", "Supermercado con Kanban"]
    c = fmt(a, "VISUAL", tablas={"senales": [{"anomalia": "Paro", "senal": "Visual"}]})
    assert c["clases"]["senales"][0]["quien"] == "falta" and any("paso 3" in x for x in c["avisos"])
    # checklist: cambia los puntos al cambiar de fase
    d = a.abrir_formato("CHECKLIST-FASE")["datos"]; assert len(d["tablas"]["puntos"]) == 7
    d["tablas"]["puntos"][0]["respuesta"] = "No"; d["cabecera"]["fase"] = "D"
    r = a.guardar_formato("CHECKLIST-FASE", d); assert any("no cumplidos: 1" in x for x in r["calc"]["avisos"])
    d["cabecera"]["fase"] = "M"; r = a.guardar_formato("CHECKLIST-FASE", d)
    assert len(r["datos"]["tablas"]["puntos"]) == 11 and r["calc"]["cab"]["resultado"].startswith("0 de 11")
    assert len(a.abrir_formato("CHECKLIST-FASE")["datos"]["tablas"]["puntos"]) == 11


def test_gestion_pmbok_calculos():
    from app.modulos.formatos6 import zona, evm, vpn, tir, recuperacion
    # PMBOK 6, figura 11-5: puntajes y zonas leídas de la figura
    for p, i, z in ((.90, .80, "Alta"), (.30, .80, "Alta"), (.50, .40, "Media"), (.10, .80, "Media"), (.90, .10, "Media"), (.70, .10, "Baja"), (.30, .10, "Baja")):
        assert zona(p * i) == z, (p, i)
    # PMBOK 7, 2.7.2: valor ganado
    r = evm(1000, 500, 400, 450)
    assert (r["SV"], r["CV"], r["SPI"]) == (-100, -50, .8) and round(r["CPI"], 4) == .8889
    assert round(r["EAC"], 2) == 1125 and round(r["ETC"], 2) == 675 and round(r["VAC"], 2) == -125 and round(r["TCPI"], 4) == 1.0909
    # PMBOK 7, 4.4.1: VPN, TIR, recuperación (ejemplo de anualidad de libro de texto)
    fl = [-1000, 400, 400, 400, 400]
    assert round(vpn(.10, fl), 2) == 267.95 and round(100 * tir(fl), 2) == 21.86 and recuperacion(fl) == 2.5


def test_gestion_pmbok_formatos():
    a = Api()
    c = fmt(a, "INTERESADOS", tablas={"registro": [{"nombre": "Gerente de planta", "rol": "Patrocinador"}, {"nombre": "Supervisor", "rol": "Dueño del proceso"}],
                                      "analisis": [{"poder": "Alto", "interes": "Alto", "urgencia": "Sí", "legitimidad": "Sí", "actual": "Neutral", "deseado": "Líder"},
                                                   {"poder": "Bajo", "interes": "Alto", "actual": "Resistente", "deseado": "De apoyo", "accion": "Taller"}]})
    an = c["calc"]["analisis"]
    assert an[0]["cuadrante"] == "Gestionar de cerca" and an[0]["prominencia"].startswith("Definitivo") and an[0]["brecha"] == "+2"
    assert an[1]["cuadrante"] == "Mantener informado" and c["clases"]["analisis"][0]["accion"] == "falta" and len(c["bloques"]) == 2
    c = fmt(a, "RIESGOS", tablas={"registro": [
        {"tipo": "Amenaza", "riesgo": "Retraso del proveedor", "probabilidad": "Alta", "impacto": "Muy alto", "estrategia": "Explotar", "costo": "10000"},
        {"tipo": "Oportunidad", "riesgo": "Descuento por volumen", "probabilidad": "Media", "impacto": "Moderado", "estrategia": "Mejorar", "acciones": "Negociar", "costo": "2000"}]},
        cab={"reserva_cont": "1000"})
    r = c["calc"]["registro"]
    assert r[0]["puntaje"] == "0.56" and r[0]["prioridad"] == "Alta" and r[0]["vem"] == "-7,000.00" and r[1]["vem"] == "1,000.00"
    assert any("estrategia de oportunidades" in x for x in c["avisos"]) and any("prioridad alta sin dueño" in x for x in c["avisos"])
    assert any("reserva de contingencia" in x for x in c["avisos"]) and c["cab"]["vem_total"] == "-6,000.00" and "<svg" in c["bloques"][0]["svg"]
    c = fmt(a, "RACI", cab={"p1": "Líder", "p2": "Champion", "p3": "Calidad"}, tablas={"actividades": [
        {"actividad": "Charter", "p1": "R", "p2": "A"}, {"actividad": "MSA", "p1": "R", "p3": "A", "p2": "A"}, {"actividad": "Piloto", "p1": "C"}]})
    rv = [x.get("revision") for x in c["calc"]["actividades"][:3]]
    assert rv[0] == "correcta" and "2 con A" in rv[1] and "falta A" in rv[2] and "falta R" in rv[2] and c["encabezados"]["p1"] == "Líder"
    c = fmt(a, "BITACORA", tablas={"supuestos": [{"tipo": "Supuesto", "descripcion": "Habrá datos de 3 turnos", "estado": "Falso"}],
                                   "incidentes": [{"descripcion": "Calibrador fuera de servicio", "fecha": "01/09/2026", "objetivo": "05/09/2026", "estado": "Abierto"}],
                                   "cambios": [{"descripcion": "Ampliar alcance a la línea 2", "decision": "Aprobado"}]})
    assert any("supuestos resultaron falsos" in x for x in c["avisos"]) and c["clases"]["incidentes"][0]["asignado"] == "falta"
    assert c["calc"]["cambios"][0]["id"] == "CC1" and c["clases"]["cambios"][0]["decidio"] == "falta" and any("fecha objetivo" in x for x in c["avisos"])
    c = fmt(a, "LECCIONES", tablas={"objetivos": [{"meta": "2", "real": "1.5", "sentido": "<=", "evidencia": "Carta p"}, {"meta": "95", "real": "90", "sentido": ">="}]})
    assert c["calc"]["objetivos"][0]["cumple"] == "Sí" and c["calc"]["objetivos"][1]["cumple"] == "No" and c["clases"]["objetivos"][1]["variacion"] == "falta"
    assert any("transición" in x for x in c["avisos"]) and any("No hay lecciones" in x for x in c["avisos"])
    assert [f["dimension"] for f in a.abrir_formato("LECCIONES")["datos"]["tablas"]["objetivos"]] == ["Alcance", "Calidad", "Costo", "Cronograma"]
    c = fmt(a, "VALOR-GANADO", cab={"bac": "1000"}, tablas={"cortes": [{"periodo": "S1", "pv": "200", "ev": "200", "ac": "210"}, {"periodo": "S2", "pv": "500", "ev": "400", "ac": "600"}]})
    assert c["calc"]["cortes"][1]["spi"] == "0.8" and c["clases"]["cortes"][1]["gasto"] == "no" and c["clases"]["cortes"][1]["explicacion"] == "falta"
    assert c["cab"]["estado"].startswith("S2: atrasado") and "EAC 1,500.00" in c["cab"]["pronostico"] and any("TCPI" in x or "CPI de" in x for x in c["avisos"])
    c = fmt(a, "COSTO-CALIDAD", cab={"ventas": "100000"}, tablas={"partidas": [{"categoria": "Prevención", "antes": "1000", "despues": "3000"},
                                                                              {"categoria": "Falla interna", "antes": "9000", "despues": "4000"}, {"categoria": "Falla externa", "antes": "2000", "despues": "3000"}]})
    assert c["cab"]["total"].startswith("Antes 12,000.00 (12.0% de ventas)") and c["cab"]["no_conformidad"] == "11,000.00 → 7,000.00"
    assert any("falla externa aumentó" in x for x in c["avisos"]) and any("prevención" in x for x in c["avisos"])
    c = fmt(a, "JUSTIFICACION", cab={"inversion": "1000", "tasa": "10"}, tablas={"flujos": [{"beneficio": "500", "costo": "100"} for _ in range(4)]})
    assert c["cab"]["vpn"].startswith("267.95") and c["cab"]["tir"].startswith("21.86%") and c["cab"]["recuperacion"] == "2.50 años" and c["cab"]["roi"] == "40.0%"
    # Charter con los campos del PMBOK 6 y vínculo al registro de interesados
    c = fmt(a, "CHARTER", cab={"caso_negocio": "x", "alcance": "y"})
    assert any("patrocinador" in x for x in c["avisos"]) and any("criterios de salida" in x for x in c["avisos"]) and "2 en el registro" in c["cab"]["interesados_registro"]


def test_tablero_adelantado_y_explicacion():
    a = Api()
    d = a.abrir_formato("TABLERO")["datos"]
    d["tablas"]["kpi"] = [{"indicador": f"K{i}", "sentido": "<=", "meta": "5", "s1": "8", "tipo_ind": "Atrasado"} for i in range(3)]
    c = a.guardar_formato("TABLERO", d)["calc"]
    assert any("no cumple la meta" in x for x in c["avisos"]) and any("adelantado" in x for x in c["avisos"]) and c["clases"]["kpi"][0]["explicacion"] == "falta"


def test_ruta_critica_figura_6_16_del_pmbok6():
    from app.modulos.formatos7 import cpm, leer_pred
    r, fin, orden = cpm({"A": (5, []), "B": (5, leer_pred("A")), "C": (10, leer_pred("A")), "D": (15, leer_pred("B, C"))})
    assert fin == 30 and [a for a in orden if abs(r[a]["TF"]) < 1e-9] == ["A", "C", "D"]
    assert (r["B"]["ES"] + 1, r["B"]["EF"], r["B"]["LS"] + 1, r["B"]["LF"], r["B"]["TF"], r["B"]["FF"]) == (6, 10, 11, 15, 5, 5)   # convención día 1
    assert (r["D"]["ES"] + 1, r["D"]["EF"]) == (16, 30)
    # relaciones SS, FF y SF con adelantos y retrasos
    r, fin, _ = cpm({"A": (4, []), "B": (3, leer_pred("A SS+1")), "C": (2, leer_pred("B FF+2")), "E": (1, leer_pred("A CF+3"))})
    assert r["B"]["ES"] == 1 and r["C"]["EF"] == r["B"]["EF"] + 2 and r["E"]["EF"] == r["A"]["ES"] + 3
    import pytest
    with pytest.raises(RecursionError): cpm({"A": (1, leer_pred("B")), "B": (1, leer_pred("A"))})


def test_arbol_de_decision_figura_11_15_del_pmbok6():
    a = Api()
    c = fmt(a, "ARBOL-DECISION", cab={"criterio": "max"}, tablas={"nodos": [
        {"tipo": "Decisión", "etiqueta": "Construir o mejorar"},
        {"padre": "1", "tipo": "Azar", "etiqueta": "Construir planta nueva", "costo": "120"},
        {"padre": "2", "tipo": "Resultado", "etiqueta": "Demanda fuerte", "prob": "60", "valor": "200"},
        {"padre": "2", "tipo": "Resultado", "etiqueta": "Demanda débil", "prob": "40", "valor": "90"},
        {"padre": "1", "tipo": "Azar", "etiqueta": "Mejorar planta", "costo": "50"},
        {"padre": "5", "tipo": "Resultado", "etiqueta": "Demanda fuerte", "prob": "60", "valor": "120"},
        {"padre": "5", "tipo": "Resultado", "etiqueta": "Demanda débil", "prob": "40", "valor": "60"}]})
    n = c["calc"]["nodos"]
    assert c["cab"]["resultado"] == "Mejorar planta · VME 46.00" and "neto 36.00" in n[1]["vme"] and "neto 46.00" in n[4]["vme"]
    assert [n[i]["neto"] for i in (2, 3, 5, 6)] == ["80.00", "-30.00", "70.00", "10.00"]


def test_edt_cronograma_montecarlo_y_otros():
    a = Api()
    c = fmt(a, "EDT", tablas={"elementos": [
        {"codigo": "1", "nombre": "Proyecto", "tipo": "Proyecto"}, {"codigo": "1.1", "nombre": "Definir", "tipo": "Cuenta de control"},
        {"codigo": "1.1.1", "nombre": "Charter", "tipo": "Paquete de trabajo", "costo": "100", "responsable": "Líder", "aceptacion": "Firmado"},
        {"codigo": "1.1.2", "nombre": "SIPOC", "tipo": "Paquete de trabajo", "costo": "50"},
        {"codigo": "1.2", "nombre": "Medir", "tipo": "Entregable", "costo": "999"}, {"codigo": "1.2.1", "nombre": "MSA", "tipo": "Paquete de trabajo", "costo": "200"},
        {"codigo": "1.3.1", "nombre": "Huérfano", "tipo": "Paquete de trabajo"}, {"codigo": "1.1.1", "nombre": "Repetido"}]})
    e = c["calc"]["elementos"]; k = c["clases"]["elementos"]
    assert c["cab"]["costo_total"] == "350.00" and e[1]["costo_total"] == "150.00" and k[3]["responsable"] == "falta"
    assert "regla del 100%" in e[4]["revision"] and "una sola cuenta de control" in e[5]["revision"] and k[6]["codigo"] == "mal" and k[7]["codigo"] == "mal"
    assert "<svg" in c["bloques"][0]["svg"]
    c = fmt(a, "CRONOGRAMA", cab={"inicio": "05/10/2026", "calendario": "habiles"}, tablas={"actividades": [
        {"id": "A", "nombre": "Charter", "o": "4", "m": "5", "p": "6", "edt": "1.1.1"}, {"id": "B", "m": "5", "pred": "A"},
        {"id": "C", "o": "8", "m": "10", "p": "12", "pred": "A"}, {"id": "D", "m": "15", "pred": "B, C"}, {"id": "H", "hito": "Sí", "pred": "D"}, {"id": "X", "m": "2", "edt": "9.9"}]})
    ac = c["calc"]["actividades"]
    assert c["cab"]["duracion"] == "30 días hábiles" and c["cab"]["ruta"].startswith("A → C → D") and ac[0]["es"] == "05/10/2026" and ac[0]["ef"] == "09/10/2026"
    assert ac[1]["ht"] == "5" and c["cab"]["fin"] == "13/11/2026" and any("9.9" in x for x in c["avisos"])
    d = a.abrir_formato("CRONOGRAMA")["datos"]; d["cabecera"]["fijar_base"] = "fijar"
    r = a.guardar_formato("CRONOGRAMA", d)
    assert r["datos"]["cabecera"]["base_fecha"] and r["datos"]["cabecera"]["fijar_base"] == "" and r["datos"]["tablas"]["actividades"][3]["base_fin"] == "30.0"
    d = r["datos"]; d["tablas"]["actividades"][2]["m"] = "13"; d["tablas"]["actividades"][2]["p"] = "15"
    c = a.guardar_formato("CRONOGRAMA", d)["calc"]
    assert c["cab"]["desvio"].startswith("+") and any("línea base" in x for x in c["avisos"])
    c = fmt(a, "MONTECARLO", cab={"iteraciones": "3000", "semilla": "7", "meta": "31"})
    assert c["cab"]["p50"] and float(c["cab"]["prob_meta"].rstrip("%")) < 100 and len(c["bloques"]) == 3
    c2 = fmt(a, "MONTECARLO", cab={"iteraciones": "3000", "semilla": "7", "meta": "31"}); assert c2["cab"] == c["cab"]          # reproducible
    c = fmt(a, "FODA", tablas={"debilidades": [{"factor": "Sin MSA", "deriva": "Datos poco confiables"}], "amenazas": [{"factor": "Proveedor único", "impacto": "Alto"}]})
    assert any("sin fortaleza" in x for x in c["avisos"]) and any("Datos poco confiables" in x for x in c["avisos"])
    c = fmt(a, "AFINIDAD", tablas={"ideas": [{"idea": "a", "grupo": "Método"}, {"idea": "b", "grupo": "Método"}, {"idea": "c"}]})
    assert c["cab"]["resumen"] == "2 ideas en 1 grupos" and c["clases"]["ideas"][2]["grupo"] == "falta"
    d = a.abrir_formato("ENFOQUE")["datos"]; assert len(d["tablas"]["variables"]) == 15
    for f in d["tablas"]["variables"]: f["valor"] = "1"
    c = a.guardar_formato("ENFOQUE", d)["calc"]; assert c["cab"]["sugerencia"].startswith("Predictivo")
    d = a.abrir_formato("CHECKLIST-FASE")["datos"]; d["cabecera"].update(fase="D", decision="Continuar a la siguiente fase", vs_caso="Con desviaciones")
    c = a.guardar_formato("CHECKLIST-FASE", d)["calc"]
    assert any("Continuar con modificaciones" in x for x in c["avisos"]) and any("quién tomó" in x for x in c["avisos"])
    d["cabecera"]["fase"] = "I"; r = a.guardar_formato("CHECKLIST-FASE", d); assert r["datos"]["tablas"]["puntos"] == [] and r["calc"]["cab"]["resultado"].startswith("Sin lista")
    c = fmt(a, "TARJETAS", tablas={"op": [{"folio": "1", "oportunidad": "x", "propuesta": "y", "esfuerzo": "1", "valor": "5"}, {"folio": "2", "oportunidad": "z", "propuesta": "w", "esfuerzo": "5", "valor": "1"}]})
    assert [x.get("cuadrante") for x in c["calc"]["op"][:2]] == ["Victoria rápida", "Evitar"] and "<svg" in c["bloques"][0]["svg"]


def test_plantillas_dmaic_tornado_y_enfoque_en_charter():
    a = Api()
    d = a.abrir_formato("CRONOGRAMA")["datos"]; d["cabecera"].update(plantilla="dmaic", calendario="naturales")
    r = a.guardar_formato("CRONOGRAMA", d)
    assert r["calc"]["cab"]["duracion"].startswith("99 días")           # 13 + 16 + 28 + 25 + 17 del curso (13_GB DEF-GANTT-DMAIC)
    assert len(r["datos"]["tablas"]["actividades"]) == 23 and r["datos"]["cabecera"]["plantilla"] == ""
    d = r["datos"]
    for f in d["tablas"]["actividades"]:
        if f.get("m") not in (None, "", "0"): f["o"], f["p"] = str(float(f["m"]) * .8), str(float(f["m"]) * 1.5)
    a.guardar_formato("CRONOGRAMA", d)
    m = a.guardar_formato("MONTECARLO", a.abrir_formato("MONTECARLO")["datos"])["calc"]
    assert len(m["bloques"]) == 3 and "tornado" in m["bloques"][2]["svg"]
    p50 = float(m["cab"]["p50"].split(" / ")[0].replace(",", "")); assert p50 > 99      # la cola pesimista alarga el proyecto
    d = a.abrir_formato("EDT")["datos"]; d["cabecera"]["plantilla"] = "dmaic"
    r = a.guardar_formato("EDT", d); el = r["datos"]["tablas"]["elementos"]
    assert len(el) == 28 and el[0]["codigo"] == "1" and el[-1]["nombre"] == "Revisiones de fase" and "21 paquetes" in r["calc"]["cab"]["resumen"]
    e = a.abrir_formato("ENFOQUE")["datos"]; e["cabecera"]["elegido"] = "Híbrido"; a.guardar_formato("ENFOQUE", e)
    assert a.guardar_formato("CHARTER", a.abrir_formato("CHARTER")["datos"])["calc"]["cab"]["enfoque"] == "Híbrido"


def test_a1_vacio_no_es_cero():
    """A1 (v0.25): un paro vacío no se toma como cero: la fila queda pendiente y no se calcula; un 0 tecleado sí cuenta."""
    a = Api()
    c = fmt(a, "OEE", tablas={"horas": [{"periodo": "1", "total": "60", "muerto": "5", "capacidad": "7", "produccion": "370", "rechazo": "3"}]})
    assert c["clases"]["horas"][0].get("paro") == "pendiente" and "oee" not in c["calc"]["horas"][0] and any("pendiente" in x for x in c["avisos"])
    c = fmt(a, "OEE", tablas={"horas": [{"periodo": "1", "total": "60", "paro": "0", "muerto": "5", "capacidad": "7", "produccion": "370", "rechazo": "3"}]})
    assert "oee" in c["calc"]["horas"][0]
