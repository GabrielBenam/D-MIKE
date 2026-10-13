"""Incremento 5: guardar/abrir .dmike y exportar .xlsx/.csv para Minitab. Datos sintéticos."""
import csv, sqlite3
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from app.api import Api

PERFIL = str(Path(__file__).parent / "sinteticos" / "perfil_prueba_sintetico.json")
SINT = ("fecha\ttiempo\tcondicion\n2026-10-05\t16.2\tSin\n05/10/2026\t12,5 min\tCon\n2026-10-06\t2.7\tCon\n")


def con_datos():
    api = Api(); api.pegar(-1, 0, SINT); return api


def test_guardar_y_abrir_conserva_todo(tmp_path):
    api = Api(); api.abrir_perfil(PERFIL)
    api.pegar(0, 6, "16.2\n2.7\nabc"); api.fijar_datos({"responsable": "Persona sintética"})
    api.marcar_avance("CORRIDA", "en_curso")
    assert api.modificado()
    ruta = str(tmp_path / "prueba.dmike")
    api._elegir = lambda *a, **k: ruta
    r = api.guardar()
    assert r["error"] is None and not api.modificado() and r["archivo"]["nombre"] == "prueba.dmike"
    b = Api(); r2 = b.abrir_proyecto(ruta)
    assert r2["error"] is None and b.hoja.estado() == api.hoja.estado()
    assert b.proyecto.a_dict() == api.proyecto.a_dict() and not b.modificado()
    assert r2["proyecto"]["version"] == 1


def test_guardar_sin_elegir_ruta_no_cambia(tmp_path):
    api = con_datos(); api._elegir = lambda *a, **k: None
    assert "cancelado" in api.guardar()["sesion"][0]["texto"] and api.modificado()


def test_guardar_agrega_extension(tmp_path):
    api = con_datos(); api._elegir = lambda *a, **k: str(tmp_path / "sin_ext")
    api.guardar(); assert (tmp_path / "sin_ext.dmike").exists()


def test_abrir_archivo_ajeno_no_toca_nada(tmp_path):
    api = con_datos(); antes = api.hoja.estado()
    malo = tmp_path / "x.dmike"; malo.write_text("no soy sqlite", encoding="utf-8")
    r = api.abrir_proyecto(str(malo))
    assert r["error"] == "abrir" and api.hoja.estado() == antes
    otro = tmp_path / "y.dmike"; c = sqlite3.connect(otro); c.execute("CREATE TABLE meta (clave, valor)")
    c.execute("INSERT INTO meta VALUES ('formato', 'otro')"); c.commit(); c.close()
    assert "no es un proyecto de D-MIKE" in api.abrir_proyecto(str(otro))["sesion"][0]["texto"]


def test_nuevo_proyecto_limpia():
    api = con_datos(); api.abrir_perfil(PERFIL)
    r = api.nuevo_proyecto()
    assert r["hoja"]["columnas"] == [] and r["proyecto"] is None and not api.modificado()


def test_exportar_xlsx(tmp_path):
    api = con_datos(); ruta = str(tmp_path / "d.xlsx")
    r = api.exportar(ruta)
    ws = load_workbook(ruta)["Datos"]
    assert [c.value for c in ws[1]] == ["fecha", "tiempo", "condicion"]
    assert ws["A2"].value == datetime(2026, 10, 5) and ws["A2"].number_format == "DD/MM/YYYY"
    assert ws["A3"].value == datetime(2026, 10, 5)            # DD/MM/AAAA leído como fecha
    assert ws["B2"].value == 16.2 and ws["B3"].value is None   # «12,5 min» sale vacío
    assert ws["C4"].value == "Con"
    assert "1 celdas que no corresponden" in r["sesion"][1]["texto"]


def test_exportar_csv(tmp_path):
    api = con_datos(); ruta = str(tmp_path / "d.csv"); api.exportar(ruta)
    filas = list(csv.reader(open(ruta, encoding="utf-8-sig")))
    assert filas == [["fecha", "tiempo", "condicion"], ["2026-10-05", "16.2", "Sin"],
                     ["2026-10-05", "", "Con"], ["2026-10-06", "2.7", "Con"]]


def test_nombres_minitab():
    api = Api(); api.pegar(-1, 0, "'tiempo#\t*x\t" + "n" * 40 + "\t" + "n" * 40 + "\n1\t2\t3\t4")
    from app.nucleo.archivo import tabla_limpia
    cols, cambios, _ = tabla_limpia(api.hoja)
    nombres = [c["nombre"] for c in cols]
    assert nombres[:2] == ["tiempo", "x"] and len(nombres[2]) == 31 and nombres[3].endswith("_2") and len(nombres[3]) == 31
    assert len(cambios) == 4


def test_exportar_hoja_vacia():
    assert Api().exportar("x.xlsx")["error"] == "vacia"


def test_ruta_con_espacios_acentos_y_numeral(tmp_path):
    d = tmp_path / "Mis proyectos #1 año"; d.mkdir()
    api = con_datos(); ruta = str(d / "prueba ñ.dmike"); api._elegir = lambda *a, **k: ruta
    assert api.guardar()["error"] is None
    b = Api(); assert b.abrir_proyecto(ruta)["error"] is None and b.hoja.n == 3
