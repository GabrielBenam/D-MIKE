"""E3-4: el diseño sale sin respuestas en un archivo de corridas (RunId, versión, orden de corrida, tipo de punto, bloque, factores y la respuesta en blanco),
en el orden en que se corre, y vuelve con sus respuestas por run_id (criterio 5 de ESPEC_flujo_real_dmaic). Datos SINTÉTICOS."""
import csv
from pathlib import Path
import openpyxl
from app.api import Api
from app.nucleo import modo as MD

DISENO = {"tipo": "fact", "nombres": "Temp, Tiempo", "bajos": "100, 10", "altos": "120, 30", "centros": "2", "replicas": "1",
          "respuesta": "Resistencia", "aleatorio": True, "semilla": "7", "limites": "90, 130; 5, 40"}
CAB = ["RunId", "version_diseno", "OrdenCorrida", "TipoPunto", "Bloque", "Temp", "Tiempo", "Resistencia"]


def api(demo=False, **extra):
    a = Api()
    if demo: MD.fijar(a.proyecto, MD.DEMO, MD.DECLARADO)
    r = a.ejecutar_modulo("DISENO-RSM", {**DISENO, **extra}); assert r["error"] is None, r["sesion"][0]["texto"]
    return a


def col(a, n): return a.hoja.indice(n)


def leer_csv(ruta): return list(csv.reader(open(ruta, encoding="utf-8-sig")))


def exporta(a, ruta, **k):
    r = a.exportar_corridas(ruta=str(ruta), **k); assert r["error"] is None, r["sesion"][-1]["texto"]; return r


def test_csv_trae_las_columnas_el_orden_de_corrida_y_la_respuesta_en_blanco(tmp_path):
    a = api(); r = exporta(a, tmp_path / "c.csv")
    filas = leer_csv(tmp_path / "c.csv")
    assert filas[0] == CAB and len(filas) == 7 and r["corridas"] == {"design_id": "DIS-Resistencia", "version_diseno": 1, "n": 6, "ruta": str(tmp_path / "c.csv")}
    assert [int(f[2]) for f in filas[1:]] == [1, 2, 3, 4, 5, 6]                       # en el orden en que se corre
    assert all(f[1] == "1" and f[-1] == "" for f in filas[1:]) and len({f[0] for f in filas[1:]}) == 6
    # cada fila trae los niveles reales de su corrida
    for f in filas[1:]:
        i = [a.hoja.columnas[col(a, "OrdenCorrida")].valor(k)[0] for k in range(a.hoja.n)].index(f[2])
        assert f[5] == a.hoja.columnas[col(a, "Temp")].valor(i)[0] and f[6] == a.hoja.columnas[col(a, "Tiempo")].valor(i)[0]
        assert f[0] == f"DIS-Resistencia-R{int(a.hoja.columnas[col(a, 'OrdenEst')].valor(i)[2]):03d}"


def test_nunca_precarga_respuestas_aunque_la_hoja_ya_las_tenga(tmp_path):
    a = api()
    for i in range(a.hoja.n): a.editar(i, col(a, "Resistencia"), str(60 + i))
    exporta(a, tmp_path / "s.csv"); assert all(f[-1] == "" for f in leer_csv(tmp_path / "s.csv")[1:])
    exporta(a, tmp_path / "c.csv", con_respuestas=True); assert all(f[-1] != "" for f in leer_csv(tmp_path / "c.csv")[1:])      # solo si se pide


def test_excel_trae_corridas_y_la_hoja_del_diseno(tmp_path):
    a = api(); exporta(a, tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "c.xlsx")
    assert wb.sheetnames == ["Corridas", "Diseño"]
    ws = wb["Corridas"]; assert [c.value for c in ws[1]] == CAB and ws.max_row == 7 and ws["A2"].value.startswith("DIS-Resistencia-R")
    assert isinstance(ws["F2"].value, (int, float)) and ws["H2"].value is None            # niveles como número; respuesta en blanco
    meta = {r[0].value: r[1].value for r in wb["Diseño"].iter_rows() if r[0].value}
    assert meta["design_id"] == "DIS-Resistencia" and meta["version_diseno"] == 1 and len(meta["Huella de la definición (SHA-256)"]) == 64
    assert meta["Temp"] == 100 and "Marca" not in meta                                    # proyecto real: sin marca sintética


def test_marca_sintetica_solo_en_demostracion(tmp_path):
    d = api(demo=True); exporta(d, tmp_path / "d.csv"); exporta(d, tmp_path / "d.xlsx")
    f = leer_csv(tmp_path / "d.csv"); assert f[0][-1] == "marca" and all(x[-1] == MD.ETIQUETA for x in f[1:])
    assert any(r[0].value == "Marca" and r[1].value == MD.ETIQUETA for r in openpyxl.load_workbook(tmp_path / "d.xlsx")["Diseño"].iter_rows())
    r = api(); exporta(r, tmp_path / "r.csv"); assert "marca" not in leer_csv(tmp_path / "r.csv")[0]


def llenar(ruta, destino, valor=lambda oe: 50 + 1.5 * oe):
    filas = leer_csv(ruta); j = filas[0].index("Resistencia")
    for f in filas[1:]: f[j] = valor(int(f[0].rsplit("-R", 1)[1])) if not f[j] else f[j]
    with open(destino, "w", encoding="utf-8", newline="") as h: csv.writer(h).writerows(filas)
    return str(destino)


def test_ida_y_vuelta_csv_conserva_el_run_id_y_verifica_niveles_sin_avisos_de_ruido(tmp_path):
    a = api(); exporta(a, tmp_path / "c.csv")
    r = a.importar_respuestas(llenar(tmp_path / "c.csv", tmp_path / "ll.csv"))
    assert r["error"] is None and r["importacion"]["escritas"] == 6
    assert not any("no trae las columnas" in x or "ignoran" in x for x in r["importacion"]["avisos"])      # trae los factores y sus columnas son las del archivo de corridas
    d = a.disenos()[0]; assert d["respuestas"] == {"n": 6, "version": 1}
    for k in range(a.hoja.n):                                                              # cada respuesta cayó en la corrida de su RunId
        oe = int(a.hoja.columnas[col(a, "OrdenEst")].valor(k)[2]); assert float(a.hoja.columnas[col(a, "Resistencia")].celdas[k]) == 50 + 1.5 * oe


def test_ida_y_vuelta_excel(tmp_path):
    a = api(); exporta(a, tmp_path / "c.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "c.xlsx"); ws = wb["Corridas"]
    for fila in range(2, ws.max_row + 1): ws.cell(fila, 8).value = 40 + int(ws.cell(fila, 1).value.rsplit("-R", 1)[1])
    wb.save(tmp_path / "ll.xlsx")
    r = a.importar_respuestas(str(tmp_path / "ll.xlsx")); assert r["error"] is None and r["importacion"]["escritas"] == 6


def test_ida_y_vuelta_con_la_marca_de_demostracion(tmp_path):
    a = api(demo=True); exporta(a, tmp_path / "c.csv")
    r = a.importar_respuestas(llenar(tmp_path / "c.csv", tmp_path / "ll.csv")); assert r["error"] is None and r["importacion"]["escritas"] == 6


def test_el_archivo_de_corridas_verifica_los_niveles_si_se_altera(tmp_path):
    a = api(); exporta(a, tmp_path / "c.csv"); filas = leer_csv(tmp_path / "c.csv")
    for f in filas[1:]: f[-1] = "5"
    filas[1][5] = "999"                                                                    # un nivel de temperatura que el diseño no tiene en esa corrida
    with open(tmp_path / "x.csv", "w", encoding="utf-8", newline="") as h: csv.writer(h).writerows(filas)
    r = a.importar_respuestas(str(tmp_path / "x.csv"))
    assert r["error"] == "respuestas" and any("Los niveles no coinciden" in e for e in r["importacion"]["errores"])


def test_el_archivo_lleva_la_version_vigente_y_uno_viejo_se_rechaza(tmp_path):
    a = api(); exporta(a, tmp_path / "v1.csv")
    nuevo = {"100": "100", "120": "140", "110": "120"}
    for i, t in enumerate(a.hoja.columnas[col(a, "Temp")].celdas[:a.hoja.n]): a.editar(i, col(a, "Temp"), nuevo[t])
    assert a.registrar_version_diseno("DIS-Resistencia", "Temperatura alta de 140")["estado"] == "nueva"
    exporta(a, tmp_path / "v2.csv")
    assert {f[1] for f in leer_csv(tmp_path / "v2.csv")[1:]} == {"2"} and {f[1] for f in leer_csv(tmp_path / "v1.csv")[1:]} == {"1"}
    r = a.importar_respuestas(llenar(tmp_path / "v1.csv", tmp_path / "l1.csv"))
    assert r["error"] == "respuestas" and any("Versión de diseño incompatible (la vigente es la 2)" in e for e in r["importacion"]["errores"])
    assert a.importar_respuestas(llenar(tmp_path / "v2.csv", tmp_path / "l2.csv"))["error"] is None


def test_errores_de_exportacion(tmp_path):
    vacio = Api(); r = vacio.exportar_corridas(ruta=str(tmp_path / "x.csv"))
    assert r["error"] == "corridas" and "no tiene diseños" in r["sesion"][-1]["texto"] and not (tmp_path / "x.csv").exists()
    a = api(); a.ejecutar_modulo("DISENO-RSM", {**DISENO, "nombres": "P, Q"})
    r = a.exportar_corridas(ruta=str(tmp_path / "y.csv")); assert r["error"] == "corridas" and "Hay varios diseños" in r["sesion"][-1]["texto"]
    assert a.exportar_corridas("DIS-Nada", str(tmp_path / "z.csv"))["error"] == "corridas"
    assert a.exportar_corridas("DIS-Resistencia_2", str(tmp_path / "ok.csv"))["error"] is None
    b = api(); b.editar(1, col(b, "OrdenEst"), b.hoja.columnas[col(b, "OrdenEst")].valor(0)[0])          # OrdenEst repetido
    r = b.exportar_corridas(ruta=str(tmp_path / "w.csv")); assert r["error"] == "corridas" and "OrdenEst repetido" in r["sesion"][-1]["texto"]


def test_diseno_anterior_a_029_tambien_exporta(tmp_path):
    a = api(); info = a.proyecto.formatos["_disenos"]["Resistencia"]
    for k in ("design_id", "version_diseno", "huella_diseno", "historial_versiones", "orden_est", "orden_corrida", "col_runid"): info.pop(k, None)
    a.hoja.eliminar_columna(col(a, "RunId"))
    exporta(a, tmp_path / "v.csv"); f = leer_csv(tmp_path / "v.csv")
    assert f[0] == CAB and f[1][0].startswith("DIS-Resistencia-R") and {x[1] for x in f[1:]} == {"1"}


def test_interfaz_menu_y_boton_exportan_corridas(pagina, tmp_path):
    from playwright.sync_api import expect
    pagina.api.ejecutar_modulo("DISENO-RSM", DISENO); ruta = str(tmp_path / "ui.csv"); pagina.api._elegir = lambda *x, **k: ruta
    pagina.reload(); pagina.wait_for_selector("body[data-listo='1']"); pagina.evaluate("document.getElementById('inicio') && document.getElementById('inicio').remove()")
    pagina.click(".menu > button:text-is('Estadística')"); pagina.click("text=Exportar corridas del diseño…")
    expect(pagina.locator("#sesion")).to_contain_text("Exportadas 6 corridas de DIS-Resistencia")
    assert leer_csv(ruta)[0] == CAB
    pagina.evaluate("ACCIONES_UI.disenos_dialogo()"); ruta2 = str(tmp_path / "ui2.xlsx"); pagina.api._elegir = lambda *x, **k: ruta2
    pagina.click("#disenos button:has-text('Exportar corridas')")
    expect(pagina.locator("#sesion")).to_contain_text("ui2.xlsx"); assert openpyxl.load_workbook(ruta2).sheetnames == ["Corridas", "Diseño"]
