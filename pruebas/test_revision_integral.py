"""Revisión integral de todo lo construido: cada módulo con datos válidos y con entradas malas (nunca debe
lanzar una excepción sin control), cada exportación abierta de nuevo, ida y vuelta del .dmike y cada opción de
menú en la interfaz sin errores de JavaScript. Datos SINTÉTICOS."""
import base64, io, json, random
import pytest
from openpyxl import load_workbook
from PIL import Image
from app.api import Api
from app import modulos as _m
_m.cargar()
from app.modulos.formatos import FORMATOS

rng = random.Random(21)
FILAS = []
for i in range(48):
    c = "Sin" if i % 2 == 0 else "Con"; op = f"Op{1 + (i // 2) % 2}"
    FILAS.append(f"2026-10-{1 + i // 4:02d}\tU-{i:03d}\t{c}\t{op}\t{round(rng.gauss(14 if c == 'Sin' else 11, 1.5), 2)}\t"
                 f"{rng.choice('AAABBCD')}\t{100 + i}\t{round(60 - 0.2 * (100 + i) + rng.gauss(0, 1), 2)}\t{rng.randint(0, 6)}\t{50 + 5 * (i % 3)}")
SINT = "fecha\tunidad\tcondicion\toperador\ttiempo\tdefecto\ttemp\tdureza\tdefectos\tn\n" + "\n".join(FILAS)

VALIDOS = {
    "ESTADISTICOS": {"y": 4, "grupo": 2}, "CORRIDA": {"y": 4, "x": 0}, "BOXPLOT": {"y": 4, "grupo": 2, "separar": 3, "puntos": True},
    "PARETO": {"cat": 5, "otros": "95"}, "ANOVA": {"y": 4, "factor": 2, "bloque": 3, "alfa": "0.05"},
    "DISTRIB": {"y": 4, "lie": "8", "lse": "18"}, "SPC": {"tipo": "xr", "y": 4, "tam": "4", "pruebas": [str(k) for k in range(1, 8)]},
    "CAPACIDAD": {"y": 4, "tam": "4", "lie": "8", "lse": "18", "objetivo": "12", "resumen": True},
    "CORREL": {"vars": [4, 6, 7], "spearman": True}, "REGRESION": {"y": 7, "xs": [6], "grado": "2", "x0": "120"},
}
INF = ("HIPOTESIS", "ANOVA2", "TUKEY", "DIST-ID", "BOXCOX", "POTENCIA", "NOPARAM", "CHI2")   # test_inferencia.py
CALC = ("CALCULADORA", "PATRON", "ALEATORIOS", "RECODIFICAR", "APILAR", "DUPLICADOS")
MALOS = [{}, {"y": None}, {"y": 2}, {"y": 999}, {"y": 4, "lie": "abc"}, {"vars": [4]}, {"xs": []}, {"tipo": "xs", "y": 4, "tam": "1"},
         {"y": 4, "tam": "-3"}, {"cat": 4}, {"y": 4, "factor": 4}]


@pytest.fixture
def api():
    a = Api(); a.pegar(-1, 0, SINT); a.fijar_datos({"nombre_proyecto": "Revisión sintética", "responsable": "Persona sintética"}); return a


def cerrar_todo(pagina):
    """Cierra diálogos, paleta, ventanas de módulos y la pantalla de inicio que haya abierto una opción de menú."""
    for _ in range(6):
        if not pagina.locator(".velo").count(): break
        b = pagina.locator(".velo .dialogo-botones button:not(.primario)")
        if b.count(): b.last.click()
        else: pagina.keyboard.press("Escape"); pagina.evaluate("document.querySelectorAll('.velo').forEach(v => v.remove())")
        pagina.wait_for_timeout(80)
    pagina.evaluate("document.getElementById('inicio') && document.getElementById('inicio').remove(); document.querySelectorAll('.ctx-menu').forEach(m => m.remove())")


def _dialogos(api):
    return [m for m, c in api.registro.items() if getattr(c, "VISTA", "dialogo") == "dialogo" and m not in FORMATOS]


def test_todos_los_dialogos_tienen_caso_valido(api):
    faltan = set(_dialogos(api)) - set(VALIDOS) - {"MSA-ATR", "DISENO-RSM", "RSM", "OPTIMIZADOR", "ORDENAR", "REEMPLAZAR", "MSA-VAR", "FACTORIAL", "PROBABILIDAD", *CALC, *INF}   # cubiertos en test_superficies, test_hoja_e2, test_gage_rr y test_calc_hoja
    assert not faltan, f"Módulos sin caso en la revisión: {faltan}"


@pytest.mark.parametrize("mid", list(VALIDOS))
def test_modulo_valido_exporta_y_repite(api, mid, tmp_path):
    r = api.ejecutar_modulo(mid, VALIDOS[mid])
    assert r["error"] is None, r["sesion"][0]["texto"]
    res = r["resultado"]
    for b in res["bloques"]:
        if b["tipo"] == "tabla":
            assert all(len(f) == len(b["cabeza"]) for f in b["filas"]), f"{mid}: tabla «{b['titulo']}» con filas desiguales"
        if b["tipo"] == "grafica": assert b["svg"].startswith("<svg")
    json.dumps(res)                                                   # todo el resultado debe poder viajar a la interfaz
    x = tmp_path / f"{mid}.xlsx"; assert api.exportar_resultado(res["uid"], str(x))["error"] is None
    wb = load_workbook(x); assert "Reporte" in wb.sheetnames
    if res.get("png"):
        p = tmp_path / f"{mid}.png"; api.exportar_resultado(res["uid"], str(p))
        assert Image.open(p).size[0] > 300
    r2 = api.repetir_resultado(res["uid"])
    assert r2["error"] is None and r2["resultado"]["titulo"] == res["titulo"]
    assert res["titulo"] in api.texto_resultado(res["uid"])


@pytest.mark.parametrize("mid", list(VALIDOS) + ["MSA-ATR", "DISENO-RSM", "RSM", "OPTIMIZADOR", "ORDENAR", "REEMPLAZAR", "MSA-VAR", "FACTORIAL", "PROBABILIDAD", *CALC, *INF])
def test_modulo_entradas_malas_no_revientan(api, mid):
    for p in MALOS:
        r = api.ejecutar_modulo(mid, dict(p))       # cualquier excepción no controlada hace fallar la prueba
        assert r["error"] in (None, "entrada"), (mid, p, r)


def test_hoja_vacia_y_columnas_vacias():
    a = Api()
    for mid in VALIDOS: assert a.dialogo_modulo(mid)["error"] == "vacia"
    for mid in ("DISENO-RSM", "OPTIMIZADOR", "POTENCIA", "HIPOTESIS"): assert a.dialogo_modulo(mid)["error"] is None   # funcionan con la hoja vacía
    a.pegar(-1, 0, "x\ty\n\t\n\t")
    for mid, p in VALIDOS.items():
        r = a.ejecutar_modulo(mid, {k: (0 if isinstance(v, int) else [0, 1] if isinstance(v, list) else v) for k, v in p.items()})
        assert r["error"] in (None, "entrada")


@pytest.mark.parametrize("fid", list(FORMATOS))
def test_formatos_llenar_exportar(api, fid, tmp_path):
    r = api.abrir_formato(fid); esp, d = r["espec"], r["datos"]
    for c in esp["cabecera"]:
        if c.get("fijo") or c.get("calc"): continue
        d["cabecera"][c["campo"]] = c["opciones"][-1][0] if c.get("tipo") == "opcion" else "tiempo" if c.get("tipo") == "col" else "5"
    for t in esp["tablas"]:
        for f in d["tablas"][t["id"]][:2]:
            for c in t["columnas"]:
                if c.get("calc"): continue
                f[c["campo"]] = (c["opciones"][-1][0] if c.get("tipo") == "opcion" else "tiempo" if c.get("tipo") == "col"
                                 else "4" if c.get("tipo") in ("numero", "entero") else "Texto sintético")
    g = api.guardar_formato(fid, d)
    assert g["error"] is None and "png" not in g["calc"]
    json.dumps(g)
    x = tmp_path / f"{fid}.xlsx"
    assert api.exportar_formato(fid, str(x))["error"] is None and load_workbook(x).sheetnames[0] == fid[:31]


def test_proyecto_completo_ida_y_vuelta(api, tmp_path):
    api.abrir_perfil("pruebas/sinteticos/perfil_prueba_sintetico.json")
    for mid, p in VALIDOS.items(): api.ejecutar_modulo(mid, p)
    for fid in FORMATOS: api.guardar_formato(fid, api.abrir_formato(fid)["datos"])
    ruta = str(tmp_path / "todo.dmike"); api._elegir = lambda *a, **k: ruta
    assert api.guardar()["error"] is None
    b = Api(); assert b.abrir_proyecto(ruta)["error"] is None
    assert len(b.resultados()) == len(VALIDOS) and set(b.proyecto.formatos) == set(FORMATOS)
    assert b.hoja.estado() == api.hoja.estado() and not b.modificado()
    assert b.repetir_resultado(b.resultados()[0]["uid"])["error"] is None
    assert b.exportar_powerbi(str(tmp_path))["error"] is None
    for ext in ("xlsx", "csv"): assert b.exportar(str(tmp_path / f"m.{ext}"))["error"] is None


def test_interfaz_cada_opcion_de_menu_sin_errores(pagina, tmp_path):
    errores = []
    pagina.on("pageerror", lambda e: errores.append(str(e)))
    pagina.api.pegar(-1, 0, SINT)
    pagina.api._elegir = lambda *a, **k: None                          # todos los diálogos de archivo se cancelan
    pagina.reload(); pagina.wait_for_selector("body[data-listo='1']")
    menus = pagina.locator(".menu > button")
    for m in range(menus.count()):
        titulo = menus.nth(m).inner_text()
        pagina.click(f".menu > button:text-is('{titulo}')")
        ops = pagina.locator(".menu.abierto button[role='menuitem']:not([disabled])")
        nombres = [ops.nth(i).inner_text() for i in range(ops.count())]
        pagina.keyboard.press("Escape")
        for n in nombres:
            if n in ("Salir",) or n.startswith("Fijar o contraer la cinta"): continue
            pagina.click(f".menu > button:text-is('{titulo}')")
            pagina.click(f".menu.abierto button:text-is('{n}')")
            pagina.wait_for_timeout(250)
            cerrar_todo(pagina)
            pagina.wait_for_timeout(100)
    assert not errores, errores
    assert pagina.locator(".pestana").count() >= 5                     # se abrieron los formatos
