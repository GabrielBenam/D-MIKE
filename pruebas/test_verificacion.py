"""Hoja de verificación contra Excel: las fórmulas de Excel (recalculadas con LibreOffice) deben coincidir con D-MIKE."""
import shutil, subprocess, random
import pytest
from openpyxl import load_workbook
from app.api import Api

pytestmark = pytest.mark.skipif(shutil.which("soffice") is None, reason="se necesita LibreOffice para recalcular las fórmulas")
rng = random.Random(17)
DATOS = "x\ty\n" + "\n".join(f"{i}\t{round(5 + 0.8 * i + rng.gauss(0, 2), 3)}" for i in range(1, 26))


def recalcular(ruta, destino):
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(destino), str(ruta)], check=True, capture_output=True, timeout=120)
    return load_workbook(destino / ruta.name, data_only=True)["Verificación"]


def filas(ws):
    out = {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r[0]: break                                      # la fila en blanco separa la tabla de la nota final
        out[r[0]] = r
    return out


@pytest.mark.parametrize("mid,params", [("ESTADISTICOS", {"y": 1}), ("CORREL", {"vars": [0, 1]}), ("REGRESION", {"y": 1, "xs": [0]}),
                                         ("CAPACIDAD", {"y": 1, "tam": "1", "lie": "0", "lse": "35"}), ("DISTRIB", {"y": 1})])
def test_excel_coincide_con_dmike(mid, params, tmp_path):
    a = Api(); a.pegar(-1, 0, DATOS)
    res = a.ejecutar_modulo(mid, params)["resultado"]
    x = tmp_path / f"{mid}.xlsx"; r = a.exportar_verificacion(res["uid"], str(x))
    assert r["error"] is None, r["sesion"][0]["texto"]
    f = filas(recalcular(x, tmp_path / "rec"))
    comparados = [k for k, v in f.items() if v[2] is not None and v[1] is not None]
    assert len(comparados) >= 4
    malos = [(k, v[1], v[2]) for k, v in f.items() if v[1] is not None and v[2] is not None and abs(v[1] - v[2]) > 1e-4 * max(1, abs(v[1]))]
    assert not malos, malos
    assert all(v[6] in ("Coincide", "sin comparar") for v in f.values())


def test_comparacion_con_minitab_marca_revisar(tmp_path):
    a = Api(); a.pegar(-1, 0, DATOS)
    res = a.ejecutar_modulo("ESTADISTICOS", {"y": 1})["resultado"]
    x = tmp_path / "e.xlsx"; a.exportar_verificacion(res["uid"], str(x))
    wb = load_workbook(x); ws = wb["Verificación"]
    ws["D7"] = ws["B7"].value * 1.5                                    # un valor de Minitab «distinto»
    wb.save(x)
    f = filas(recalcular(x, tmp_path / "rec"))
    assert f["Media"][6] == "Revisar"


def test_no_soportados_y_filtro(tmp_path):
    a = Api(); a.pegar(-1, 0, DATOS + "\n30\t")
    r = a.ejecutar_modulo("CORRIDA", {"y": 1})["resultado"]
    assert a.exportar_verificacion(r["uid"], str(tmp_path / "n.xlsx"))["error"] == "verificacion"
