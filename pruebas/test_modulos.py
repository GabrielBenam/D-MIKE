"""Los 10 módulos del MVP y Estadísticos de columna, a través del puente real. Datos SINTÉTICOS."""
import random
from openpyxl import load_workbook
import pytest
from app.api import Api
from app.modulos.analisis import clases_histograma

random.seed(7)
FILAS = []
for i in range(40):
    cond = "Sin" if i % 2 == 0 else "Con"; op = f"Op{1 + (i // 2) % 2}"
    t = round(random.gauss(15 if cond == "Sin" else 11, 1.8) + (0.6 if op == "Op2" else 0), 1)
    FILAS.append(f"2026-10-{5 + i // 2:02d}\tU-{i + 1:03d}\t{cond}\t{op}\t{t}\t{random.choice(['A', 'A', 'A', 'B', 'B', 'C', 'D'])}")
SINT = "fecha\tunidad\tcondicion\toperador\ttiempo\ttipo_error\n" + "\n".join(FILAS)


@pytest.fixture
def api():
    a = Api(); a.pegar(-1, 0, SINT); return a


def ok(api, mid, p):
    r = api.ejecutar_modulo(mid, p)
    assert r["error"] is None, r["sesion"][0]["texto"]
    return r["resultado"]


def textos(res): return " ".join(b.get("texto", "") for b in res["bloques"])


def test_clases_histograma_ejemplo_logistica():
    # 16_GB, HIST-PROCEDIMIENTO: 200 datos de 150 a 179, K = 10 -> W = 3, primera clase 149.5 a 152.5
    x = [150 + (i % 30) for i in range(200)]
    b, u, w = clases_histograma(x, k=10)
    assert u == 1 and w == 3 and b[0] == 149.5 and b[1] == 152.5


def test_estadisticos_por_grupo(api):
    r = ok(api, "ESTADISTICOS", {"y": 4, "grupo": 2})
    t = r["bloques"][0]
    assert [f[0] for f in t["filas"]] == ["Sin", "Con"] and t["filas"][0][1] == "20"


def test_corrida(api):
    r = ok(api, "CORRIDA", {"y": 4, "x": 0})
    assert r["png"] and "<svg" in r["bloques"][0]["svg"] and "mediana" in textos(r).lower()


def test_boxplot_con_separacion(api):
    r = ok(api, "BOXPLOT", {"y": 4, "grupo": 2, "separar": 3, "puntos": True})
    assert len(r["bloques"][1]["filas"]) == 4


def test_anova_detecta_diferencia(api):
    r = ok(api, "ANOVA", {"y": 4, "factor": 2, "alfa": "0.05"})
    assert "se rechaza H0" in r["bloques"][0]["texto"] and "no se rechaza" not in r["bloques"][0]["texto"]
    assert "equivale a la prueba t" in " ".join(r["avisos"])


def test_anova_con_bloque_balanceado(api):
    r = ok(api, "ANOVA", {"y": 4, "factor": 2, "bloque": 3, "alfa": "0.05"})
    assert "MSBL / MSE" in textos(r) and r["bloques"][2]["filas"][1][0] == "operador"


def test_anova_bloque_desbalanceado_rechaza(api):
    api.editar(0, 3, "Op9")
    r = api.ejecutar_modulo("ANOVA", {"y": 4, "factor": 2, "bloque": 3, "alfa": "0.05"})
    assert r["error"] == "entrada" and "balanceado" in r["sesion"][0]["texto"]


def test_pareto(api):
    r = ok(api, "PARETO", {"cat": 5, "otros": "95"})
    fl = r["bloques"][1]["filas"]
    assert fl[0][0] == "A" and fl[-1][3] == "100"


def test_distrib_con_especificaciones(api):
    r = ok(api, "DISTRIB", {"y": 4, "lie": "8", "lse": "18"})
    assert "Prueba de normalidad" in textos(r) and "Nivel sigma" in textos(r)


def test_distrib_lie_mayor_que_lse(api):
    assert api.ejecutar_modulo("DISTRIB", {"y": 4, "lie": "20", "lse": "10"})["error"] == "entrada"


def test_msa_atributos_desde_hoja():
    a = Api()
    ref = ["P", "P", "P", "P", "P", "F", "F", "F", "F", "F"]
    b1 = ["P", "F", "P", "P", "P", "P", "F", "F", "F", "F"]; b2 = ["P", "F", "F", "P", "P", "P", "F", "F", "F", "F"]
    txt = "ref\tBen_1\tBen_2\tCyn_1\tCyn_2\n" + "\n".join("\t".join(x) for x in zip(ref, b1, b2, ref, ref))
    a.pegar(-1, 0, txt)
    r = ok(a, "MSA-ATR", {"ref": 0, "e1_t1": 1, "e1_t2": 2, "e2_t1": 3, "e2_t2": 4, "pasa": "P"})
    fila_ben = r["bloques"][1]["filas"][0]
    assert fila_ben[0] == "Ben" and fila_ben[1] == "9 de 10 (90.0%)" and fila_ben[4] == "7 de 10 (70.0%)"
    assert fila_ben[6:9] == ["1", "1", "1"]


def test_errores_de_entrada(api):
    assert api.ejecutar_modulo("CORRIDA", {"y": None})["error"] == "entrada"
    assert api.ejecutar_modulo("CORRIDA", {"y": 2})["error"] == "entrada"      # columna de texto
    assert Api().dialogo_modulo("ANOVA")["error"] == "vacia"


def test_dialogo_usa_roles():
    a = Api(); a.abrir_perfil("pruebas/sinteticos/perfil_prueba_sintetico.json")
    d = a.dialogo_modulo("ANOVA")["dialogo"]
    c = {x["campo"]: x.get("defecto") for x in d["campos"]}
    assert a.hoja.columnas[c["y"]].nombre == "tiempo" and a.hoja.columnas[c["factor"]].nombre == "condicion"


def test_resultado_se_guarda_y_exporta(api, tmp_path):
    r = ok(api, "BOXPLOT", {"y": 4, "grupo": 2})
    ruta = str(tmp_path / "caja.xlsx"); api.exportar_resultado(r["uid"], ruta)
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Reporte", "Datos"] and len(wb["Reporte"]._images) == 1
    assert [c.value for c in wb["Datos"][1]] == ["tiempo", "condicion"]
    api.exportar_resultado(r["uid"], str(tmp_path / "caja.png"))
    assert (tmp_path / "caja.png").read_bytes()[:4] == b"\x89PNG"
    api._elegir = lambda *a, **k: str(tmp_path / "p.dmike"); api.guardar()
    b = Api(); b.abrir_proyecto(str(tmp_path / "p.dmike"))
    assert b.resultados()[0]["titulo"] == r["titulo"]


def test_avance_automatico():
    a = Api(); a.abrir_perfil("pruebas/sinteticos/perfil_prueba_sintetico.json")
    a.pegar(0, 6, "\n".join(str(10 + i % 7) for i in range(20)))
    r = a.ejecutar_modulo("CORRIDA", {"y": 6})
    assert a.proyecto.modulos["CORRIDA"]["avance"] == "en_curso" and "proyecto" in r


# ---------- formatos ----------
def test_hoja_datos_plan_de_muestreo(tmp_path):
    a = Api(); a.fijar_datos({"nombre_proyecto": "Proyecto sintético"})
    d = a.abrir_formato("HOJA-DATOS")["datos"]
    d["cabecera"].update(tipo_dato="discreto", poblacion="400", p="0.09", d="0.01", confianza="90", tiempo="480")
    r = a.guardar_formato("HOJA-DATOS", d)
    assert r["calc"]["cab"]["n_calc"].startswith("339 (sin ajuste 2,203")      # 15_GB, MUE-DISCRETA
    d["cabecera"].update(tipo_dato="continuo", poblacion="3200", s="0.458", d="0.1", confianza="95")
    r = a.guardar_formato("HOJA-DATOS", d)
    assert r["calc"]["cab"]["n_calc"] == "81" and r["calc"]["cab"]["frecuencia_calc"] == "5.93"   # MUE-CONTINUA
    a.pegar(-1, 0, "fecha\ttiempo\n2026-10-05\t3")
    ruta = str(tmp_path / "h.xlsx"); a.exportar_formato("HOJA-DATOS", ruta)
    wb = load_workbook(ruta)
    assert "Recoleccion" in wb.sheetnames and wb["HOJA-DATOS"]["A4"].value == "Proyecto sintético"


def test_amef_rpn_y_alertas(tmp_path):
    a = Api(); d = a.abrir_formato("AMEF")["datos"]
    d["tablas"]["filas"][0].update(operacion="Operación sintética", falla="Falla 1", sev="9", occ="2", det="3")
    d["tablas"]["filas"][1].update(falla="Falla 2", sev="5", occ="6", det="4", sev2="5", occ2="3", det2="2")
    d["tablas"]["filas"][2].update(falla="Falla 3", sev="11", occ="2", det="2")
    c = a.guardar_formato("AMEF", d)["calc"]
    assert c["calc"]["filas"][0]["rpn"] == 54 and c["calc"]["filas"][1]["rpn"] == 120 and c["calc"]["filas"][1]["rpn2"] == 30
    assert c["calc"]["filas"][2]["rpn"] == "" and c["clases"]["filas"][2]["sev"] == "mal"
    assert c["clases"]["filas"][0]["sev"] == "alerta" and any("Severidad 9 a 10" in x for x in c["avisos"])
    ruta = str(tmp_path / "amef.xlsx"); a.exportar_formato("AMEF", ruta)
    ws = load_workbook(ruta)["AMEF"]
    vals = [c.value for fila in ws.iter_rows() for c in fila]
    assert 120 in vals and 54 in vals and "Operación sintética" in vals


def test_plan_control_avisos():
    a = Api(); d = a.abrir_formato("PLAN-CONTROL")["datos"]
    d["tablas"]["filas"][0].update(paso="Paso sintético", que="Temperatura", critico="Sí")
    c = a.guardar_formato("PLAN-CONTROL", d)["calc"]
    assert any("sin acción correctiva" in x for x in c["avisos"]) and any("aprueba" in x for x in c["avisos"])
    assert c["clases"]["filas"][0]["accion"] == "falta"


def test_tablero_manual_y_desde_hoja(api):
    d = api.abrir_formato("TABLERO")["datos"]
    d["cabecera"].update(semanas="4", inicio="05/10/2026", col_fecha="fecha")
    d["tablas"]["kpi"][0].update(indicador="Tiempo medio", sentido="<=", meta="14", col="tiempo", calculo="media")
    d["tablas"]["kpi"][1].update(indicador="Manual", sentido=">=", meta="25", s1="21", s2="26", s3="x")
    c = api.guardar_formato("TABLERO", d)["calc"]
    assert c["semanas"] == 4 and c["calc"]["kpi"][0]["s1"] != "" and "<svg" in c["calc"]["kpi"][0]["tendencia"]
    k = c["clases"]["kpi"][1]
    assert k["s1"] == "no" and k["s2"] == "ok" and k["s3"] == "mal"


def test_formatos_persisten(api, tmp_path):
    d = api.abrir_formato("AMEF")["datos"]; d["tablas"]["filas"][0]["falla"] = "Falla sintética"
    api.guardar_formato("AMEF", d)
    api._elegir = lambda *a, **k: str(tmp_path / "f.dmike"); api.guardar()
    b = Api(); b.abrir_proyecto(str(tmp_path / "f.dmike"))
    assert b.abrir_formato("AMEF")["datos"]["tablas"]["filas"][0]["falla"] == "Falla sintética"
