"""Plantillas propias en Excel: reporte de un resultado de módulo y formatos (cabecera + tablas).
Identidad: encabezados naranja #dd5813, Montserrat (Excel usa otra fuente si no la tiene), celdas numéricas a la derecha."""
import base64, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from ..nucleo.hoja import a_numero

FUENTE = "Montserrat"
NARANJA = "DD5813"
T_TITULO = Font(name=FUENTE, size=14, bold=True, color=NARANJA)
T_ETQ = Font(name=FUENTE, size=9, color="6F675B")
T_VAL = Font(name=FUENTE, size=10, bold=True, color="2B2620")
T_CAB = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
T_NORMAL = Font(name=FUENTE, size=10, color="2B2620")
T_NOTA = Font(name=FUENTE, size=9, italic=True, color="6F675B")
RELLENO = PatternFill("solid", fgColor=NARANJA)
CLARO = PatternFill("solid", fgColor="FBE8DD")
LINEA = Border(bottom=Side(style="thin", color="E2DCD0"))
COLOR = {"ok": PatternFill("solid", fgColor="E3F1E0"), "no": PatternFill("solid", fgColor="FDECEA"),
         "advertencia": PatternFill("solid", fgColor="FFF4D6"), "falta": PatternFill("solid", fgColor="FFFAF2"),
         "alerta": PatternFill("solid", fgColor="FBE8DD"), "mal": PatternFill("solid", fgColor="FDECEA")}


def _cabecera(ws, titulo, pares, fila=1):
    ws.cell(fila, 1, titulo).font = T_TITULO
    fila += 2
    for k, (etq, val) in enumerate(pares):
        col = 1 + (k % 3) * 3
        ws.cell(fila, col, etq).font = T_ETQ
        c = ws.cell(fila + 1, col, val or ""); c.font = T_VAL
        if k % 3 == 2: fila += 3
    return fila + (3 if len(pares) % 3 else 0) + 1


def _tabla(ws, fila, titulo, cabeza, filas, clases=None, anchos=None, crudas=None):
    if titulo:
        ws.cell(fila, 1, titulo).font = Font(name=FUENTE, size=11, bold=True, color="2B2620"); fila += 1
    for j, h in enumerate(cabeza, 1):
        c = ws.cell(fila, j, h); c.font = T_CAB; c.fill = RELLENO
        c.alignment = Alignment(wrap_text=True, vertical="center")
    fila += 1
    for i, f in enumerate(filas):
        for j, v in enumerate(f, 1):
            cr = crudas[i][j - 1] if crudas and i < len(crudas) and j - 1 < len(crudas[i]) else None
            c = ws.cell(fila, j, cr if cr is not None else _celda(v))          # A4: número con precisión completa; Excel lo redondea solo al mostrar
            c.font = T_NORMAL; c.border = LINEA
            c.alignment = Alignment(horizontal="right" if isinstance(c.value, (int, float)) else "left", vertical="top", wrap_text=not isinstance(c.value, (int, float)))
            k = (clases[i] if clases and i < len(clases) else {}).get(j - 1)
            if k in COLOR: c.fill = COLOR[k]
        fila += 1
    return fila + 1


def _celda(v):
    """Texto numérico -> número (para que Excel lo trate como número); lo demás se queda como texto."""
    if isinstance(v, str):
        x = a_numero(v) if v.strip() else None
        return x if x is not None else v
    return v


def _datos_fijos(proyecto):
    f = {k: v["valor"] for k, v in proyecto.fijos.items()}
    return [("Proyecto", f.get("nombre_proyecto")), ("Cliente o área", f.get("cliente_area")), ("Responsable", f.get("responsable"))]


def _anchos(ws, n, ancho=16):
    for j in range(1, n + 1): ws.column_dimensions[get_column_letter(j)].width = ancho


def reporte_resultado(ruta, res, proyecto, hoja_datos):
    """Reporte con cabecera de datos fijos, tablas, textos e imagen; hoja «Datos» plana para Minitab."""
    wb = Workbook(); ws = wb.active; ws.title = "Reporte"
    fila = _cabecera(ws, res["titulo"], _datos_fijos(proyecto) + [("Fecha del análisis", res["fecha"]), ("Módulo", res["modulo"])])
    for a in res.get("avisos", []):
        ws.cell(fila, 1, "Aviso: " + a).font = T_NOTA; fila += 1
    if res.get("avisos"): fila += 1
    img_fila = None
    for b in res["bloques"]:
        if b["tipo"] == "tabla":
            fila = _tabla(ws, fila, b["titulo"], b["cabeza"], b["filas"], crudas=b.get("crudas"))
        elif b["tipo"] == "texto":
            c = ws.cell(fila, 1, b["texto"]); c.font = T_VAL if b["clase"] == "conclusion" else T_NOTA
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=9)
            c.alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[fila].height = 15 * max(1, len(b["texto"]) // 110 + 1)
            fila += 2
        elif b["tipo"] == "grafica" and img_fila is None:
            img_fila = fila; fila += 22
    if res.get("png") and img_fila:
        from openpyxl.drawing.image import Image
        im = Image(io.BytesIO(base64.b64decode(res["png"]))); im.width, im.height = im.width * 0.5, im.height * 0.5
        ws.add_image(im, f"A{img_fila}")
    _anchos(ws, 12)
    if hoja_datos:
        wd = wb.create_sheet("Datos")
        for j, col in enumerate(hoja_datos, 1):
            wd.cell(1, j, col["nombre"]).font = Font(bold=True)
            for i, v in enumerate(col["valores"], 2):
                if v is None: continue
                c = wd.cell(i, j, v)
                if col["tipo"] == "fecha": c.number_format = "DD/MM/YYYY"
    wb.save(ruta)


def reporte_formato(ruta, cls, datos, calc, proyecto, hoja=None):
    wb = Workbook(); ws = wb.active; ws.title = cls.ID[:31]
    fijos = {k: v["valor"] for k, v in proyecto.fijos.items()}
    pares = []
    for c in cls.CABECERA:
        if c.get("fijo"): v = fijos.get(c["fijo"])
        elif c.get("calc"): v = calc["cab"].get(c["campo"], "")
        else:
            v = datos["cabecera"].get(c["campo"], "")
            if c.get("tipo") == "opcion": v = dict(map(tuple, c["opciones"])).get(v, v)
        pares.append((c["etiqueta"], v))
    fila = _cabecera(ws, cls.TITULO, pares)
    for a in calc.get("avisos", []):
        ws.cell(fila, 1, "Aviso: " + a).font = T_NOTA; fila += 1
    fila += 1
    ns = calc.get("semanas")
    for t in cls.TABLAS:
        cols = [c for c in t["columnas"] if c.get("tipo") != "svg" and not c.get("oculto") and (not c.get("semana") or ns is None or c["semana"] <= ns)]
        filas, clases = [], []
        cc, kk = calc["calc"].get(t["id"], []), calc["clases"].get(t["id"], [])
        for i, f in enumerate(datos["tablas"][t["id"]]):
            if not any(f.values()) and not any((cc[i] if i < len(cc) else {}).values()): continue
            fila_v, cl = [], {}
            for j, c in enumerate(cols):
                v = (cc[i] if i < len(cc) else {}).get(c["campo"]) if c.get("calc") or (i < len(cc) and c["campo"] in cc[i]) else f.get(c["campo"], "")
                if c.get("tipo") == "opcion" and (not isinstance(v, str) or a_numero(v) is None): v = dict(map(tuple, c["opciones"])).get(v, v)
                fila_v.append("" if v is None else v)
                k = (kk[i] if i < len(kk) else {}).get(c["campo"])
                if k: cl[j] = k
            filas.append(fila_v); clases.append(cl)
        fila = _tabla(ws, fila, t["titulo"], [c["etiqueta"] for c in cols], filas, clases)
        for j, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(8, min(40, c.get("ancho", 110) / 7))
    for b in calc.get("bloques", []):
        if b["tipo"] == "texto":
            c = ws.cell(fila, 1, b["texto"]); c.font = T_VAL; fila += 2
    for t in calc.get("tarjetas", []):
        ws.cell(fila, 1, t["indicador"]).font = T_VAL
        ws.cell(fila, 2, f"{t['valor']} {t['unidad']}".strip()).font = T_NORMAL
        ws.cell(fila, 3, t["meta"]).font = T_ETQ
        c = ws.cell(fila, 4, {"ok": "Cumple", "advertencia": "Advertencia", "no": "No cumple"}.get(t["estado"], "")); c.font = T_NORMAL
        if t["estado"] in COLOR: c.fill = COLOR[t["estado"]]
        fila += 1
    if calc.get("png"):
        from openpyxl.drawing.image import Image
        im = Image(io.BytesIO(base64.b64decode(calc["png"]))); im.width, im.height = im.width * 0.55, im.height * 0.55
        ws.add_image(im, f"A{fila + 1}"); fila += 30
    if cls.NOTA:
        c = ws.cell(fila + 1, 1, cls.NOTA); c.font = T_NOTA
    if cls.ID == "HOJA-DATOS" and hoja is not None:
        _hoja_recoleccion(wb, pares, hoja, calc)
    wb.save(ruta)


def _hoja_recoleccion(wb, pares, hoja, calc):
    """Hoja de recolección imprimible: cabecera + columnas de la hoja central, con los datos ya capturados y
    renglones en blanco hasta completar la muestra calculada."""
    ws = wb.create_sheet("Recoleccion")
    sel = [p for p in pares if p[0] in ("Proyecto", "Responsable", "Cliente o área", "Inicio de captura (DD/MM/AAAA)",
                                        "Instrumento", "Qué se cuenta o mide (unidades)")]
    fila = _cabecera(ws, "Hoja de recolección de datos", sel)
    cols = [c for c in hoja.columnas if c.nombre]
    n_calc = a_numero((calc["cab"].get("n_calc") or "").split(" ")[0])
    total = max(hoja.n, int(n_calc) if n_calc else 0, 20)
    ws.cell(fila, 1, "#").font = T_CAB; ws.cell(fila, 1).fill = RELLENO
    for j, c in enumerate(cols, 2):
        x = ws.cell(fila, j, c.nombre); x.font = T_CAB; x.fill = RELLENO
    for i in range(total):
        ws.cell(fila + 1 + i, 1, i + 1).font = T_ETQ
        for j, c in enumerate(cols, 2):
            v = c.valor(i) if i < hoja.n else ("", True, None)
            cel = ws.cell(fila + 1 + i, j, v[2] if v[1] and v[2] is not None else (v[0] or None)); cel.border = LINEA
            if c.tipo == "fecha" and v[2] is not None: cel.number_format = "DD/MM/YYYY"
    _anchos(ws, len(cols) + 1, 15)
