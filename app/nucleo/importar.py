"""Importar datos a la hoja desde CSV, TSV, TXT delimitado o Excel (.xlsx). Primera fila = nombres de columna.

Decisiones (E2, no respaldadas por la base; formatos de archivo):
- CSV: se detecta el separador entre tabulador, punto y coma y coma (el que aparece de forma más constante).
- Con separador punto y coma (CSV de Excel en español), los números con coma decimal («12,5») se convierten a punto.
- Codificación: UTF-8 (con o sin BOM); si falla, Windows-1252.
- Excel: primera hoja del libro, valores (no fórmulas); fechas a AAAA-MM-DD.
"""
import csv, io, os, re
from datetime import date, datetime

COMA_DECIMAL = re.compile(r"^[+-]?\d+,\d+$")


def _texto(ruta):
    with open(ruta, "rb") as f: crudo = f.read()
    for cod in ("utf-8-sig", "cp1252"):
        try: return crudo.decode(cod)
        except UnicodeDecodeError: continue
    raise ValueError("no se pudo leer la codificación del archivo")


def _separador(texto):
    lineas = [l for l in texto.splitlines() if l.strip()][:20]
    mejor, puntaje = ",", -1
    for s in ("\t", ";", ","):
        cuentas = [l.count(s) for l in lineas]
        if not cuentas or min(cuentas) == 0: continue
        p = min(cuentas) * (2 if len(set(cuentas)) == 1 else 1)
        if p > puntaje: mejor, puntaje = s, p
    return mejor


def leer(ruta):
    """Devuelve (filas como listas de texto, descripción del formato)."""
    ext = os.path.splitext(ruta)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = []
        for fila in ws.iter_rows(values_only=True):
            filas.append(["" if v is None else v.date().isoformat() if isinstance(v, datetime) else v.isoformat() if isinstance(v, date)
                          else (str(int(v)) if isinstance(v, float) and v.is_integer() else repr(v)) if isinstance(v, float) else str(v) for v in fila])
        wb.close()
        return _limpiar(filas), f"Excel, hoja «{ws.title}»"
    if ext in (".xls",):
        raise ValueError("el formato .xls antiguo no se puede leer; guárdalo como .xlsx o .csv desde Excel")
    texto = _texto(ruta)
    sep = _separador(texto)
    filas = [list(r) for r in csv.reader(io.StringIO(texto), delimiter=sep)]
    if sep == ";":
        filas = [[c.replace(",", ".") if COMA_DECIMAL.match(c.strip()) else c for c in f] for f in filas]
    nombre_sep = {"\t": "tabulador", ";": "punto y coma", ",": "coma"}[sep]
    return _limpiar(filas), f"texto separado por {nombre_sep}"


def _limpiar(filas):
    filas = [[str(c).replace("\t", " ").replace("\n", " ").strip() for c in f] for f in filas]
    while filas and not any(filas[-1]): filas.pop()
    if not filas: raise ValueError("el archivo no tiene datos")
    ancho = max(len(f) for f in filas)
    while ancho and all(len(f) < ancho or not f[ancho - 1] for f in filas): ancho -= 1
    return [f[:ancho] + [""] * (ancho - len(f)) for f in filas]
