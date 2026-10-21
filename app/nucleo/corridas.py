"""Archivo de corridas de un diseño (Entrega 3, E3-4): el diseño sale sin respuestas para correr los experimentos y vuelve con ellas por run_id.

Columnas, en el orden en que se corre: RunId, version_diseno, OrdenCorrida, TipoPunto, Bloque, los factores (niveles reales) y la respuesta en blanco.
CSV UTF-8 con BOM, coma y punto decimal; o Excel (hoja «Corridas» y hoja «Diseño» con la identidad y los niveles). En un proyecto de demostración
se agrega la columna «marca» con la etiqueta sintética; en un proyecto real nunca. El archivo no trae respuestas aunque la hoja ya las tenga, salvo
que se pida (con_respuestas). Se vuelve a importar con E3-3 (Importar respuestas del diseño): trae los factores, así que sus niveles se verifican.
"""
import csv
from . import disenos as DIS, modo as MD


def _numero(t):
    """Valor para escribir: entero si lo es, flotante si no, y el texto tal cual si no es número."""
    try: x = float(t)
    except (TypeError, ValueError): return t
    return int(x) if x.is_integer() else x


def preparar(proyecto, hoja, design_id=None, con_respuestas=False):
    """(cabecera, filas, meta) del archivo de corridas. Lanza ValueError con el motivo si no se puede exportar."""
    ds = DIS.descriptores(proyecto, hoja)
    if not ds: raise ValueError("Este proyecto no tiene diseños de experimentos: crea uno primero (Estadística > Diseño de experimentos).")
    if design_id is None:
        if len(ds) > 1: raise ValueError("Hay varios diseños (" + ", ".join(d["design_id"] for d in ds) + "): elige uno en «Diseños del proyecto».")
        design_id = ds[0]["design_id"]
    d = next((x for x in ds if x["design_id"] == design_id), None)
    if d is None: raise ValueError(f"El diseño {design_id} no existe en este proyecto.")
    info = proyecto.formatos["_disenos"][d["clave"]]
    cols = DIS.columnas(hoja, info)
    if cols["orden_est"] is None or any(j is None for j in cols["factores"]) or cols["respuesta"] is None:
        raise ValueError(f"Faltan en la hoja columnas de {design_id} (OrdenEst, factores o respuesta): no se puede exportar.")
    ids, prob = DIS.run_ids(hoja, info, d["clave"])
    if prob["repetidos"]: raise ValueError("La hoja tiene OrdenEst repetido (" + ", ".join(sorted(prob["repetidos"])[:5]) + "): corrígelo antes de exportar.")
    marca = MD.etiqueta(proyecto)
    cab = ["RunId", "version_diseno", "OrdenCorrida", "TipoPunto", "Bloque"] + [f["nombre"] for f in info["factores"]] + [info["respuesta"]] + (["marca"] if marca else [])
    filas = []
    for i, r in enumerate(ids):
        if not r: continue
        def v(j): return None if j is None else _numero(hoja.columnas[j].valor(i)[0]) if hoja.columnas[j].valor(i)[0] != "" else None
        resp = v(cols["respuesta"]) if con_respuestas else None
        orden = v(cols["orden_corrida"]) if cols["orden_corrida"] is not None else None
        filas.append((orden if orden is not None else float("inf"), v(cols["orden_est"]),
                      [r, d["version_diseno"], orden, v(cols["tipo_punto"]), v(cols["bloque"])] + [v(j) for j in cols["factores"]] + [resp] + ([marca] if marca else [])))
    filas.sort(key=lambda t: (t[0], t[1]))
    meta = {"design_id": design_id, "version_diseno": d["version_diseno"], "tipo": info.get("tipo"), "respuesta": info["respuesta"], "corridas": len(filas),
            "factores": [{"nombre": f["nombre"], "bajo": f.get("bajo"), "alto": f.get("alto")} for f in info["factores"]], "marca": marca,
            "huella_definicion": DIS.huella_definicion(hoja, info), "con_respuestas": bool(con_respuestas)}
    return cab, [f[2] for f in filas], meta


def escribir_csv(ruta, cab, filas):
    with open(ruta, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(cab)
        for fila in filas: w.writerow(["" if x is None else x for x in fila])


def escribir_xlsx(ruta, cab, filas, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = "Corridas"
    ws.append(cab)
    for c in ws[1]: c.font = Font(bold=True)
    for fila in filas: ws.append(fila)
    for j, c in enumerate(cab, 1): ws.column_dimensions[ws.cell(1, j).column_letter].width = max(12, len(str(c)) + 2)
    ws.freeze_panes = "A2"
    d = wb.create_sheet("Diseño")
    d.append(["Concepto", "Valor"])
    for c in d[1]: c.font = Font(bold=True)
    for k, v in (("design_id", meta["design_id"]), ("version_diseno", meta["version_diseno"]), ("Tipo", meta["tipo"]), ("Corridas", meta["corridas"]),
                 ("Respuesta", meta["respuesta"]), ("Huella de la definición (SHA-256)", meta["huella_definicion"]),
                 ("Respuestas incluidas", "sí" if meta["con_respuestas"] else "no: se capturan al correr los experimentos")):
        d.append([k, v])
    d.append([]); d.append(["Factor", "Nivel bajo", "Nivel alto"])
    for c in d[d.max_row]: c.font = Font(bold=True)
    for f in meta["factores"]: d.append([f["nombre"], f["bajo"], f["alto"]])
    if meta["marca"]: d.append([]); d.append(["Marca", meta["marca"]])
    d.column_dimensions["A"].width = 36; d.column_dimensions["B"].width = 44
    wb.save(ruta)
