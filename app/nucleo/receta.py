"""Recetas del diagnóstico (contrato v1 y v2; ver ESPEC_receta_v2.md; 0.27: variables, tareas, estados y revisiones, secciones 12 a 15). Una receta es un archivo DECLARATIVO (datos y parámetros, nunca código) que genera el
artefacto en línea «Diagnóstico de proyecto D-MIKE». D-MIKE la valida contra su catálogo y la ejecuta: datos fijos, perfil,
columnas y datos, diseños de experimentos con sus respuestas, formatos, análisis y avance.
La receta viaja sola (.json) o dentro de un libro de Excel (.xlsx) en la hoja oculta «_receta», junto con las hojas de datos
que el usuario llenó. En un parámetro, «@nombre» es una columna de la hoja y «#id» el resultado de un análisis anterior."""
import hashlib, json, re
from datetime import datetime

TIPO, VERSION_SOPORTADA = "dmike-receta", 2
HOJA_RECETA = "_receta"
_UUID4 = re.compile(r"[0-9a-f]{8}-[0-9a-f]{4}-4[0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}")
ORIGEN_DATOS = ("real", "vacia", "sintetico")
ESTADOS_ANALISIS = ("pendiente", "disponible", "analizado", "aprobado")
ESTADOS_TAREA = ("abierta", "cerrada")
EVIDENCIA = ("REAL_VERIFICADO", "REAL_NO_VALIDADO", "PENDIENTE_RECOLECCION", "DESCONOCIDO", "NO_APLICA", "SINTETICO_DEMO")


class RecetaInvalida(ValueError):
    pass


def _version(v):
    try: return tuple(int(x) for x in str(v).split(".")[:3])
    except ValueError: return (0,)


def leer(ruta):
    """Devuelve (receta, datos) donde datos = {hoja: [[encabezados], [fila], ...]} (vacío si es .json)."""
    if ruta.lower().endswith(".json"):
        with open(ruta, encoding="utf-8-sig") as f: return json.load(f), {}
    import openpyxl
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    if HOJA_RECETA not in wb.sheetnames: raise RecetaInvalida("El libro no trae la hoja «_receta»: no es una plantilla generada por el diagnóstico.")
    trozos = [r[0] for r in wb[HOJA_RECETA].iter_rows(values_only=True) if r and r[0]]
    receta = json.loads("".join(str(t) for t in trozos))
    datos = {}
    for h in wb.sheetnames:
        if h.startswith("_") or h.upper() == "LEEME": continue
        filas = [list(r) for r in wb[h].iter_rows(values_only=True)]
        while filas and all(v in (None, "") for v in filas[-1]): filas.pop()
        if filas: datos[h] = filas
    return receta, datos


def es_receta(ruta):
    try:
        r, _ = leer(ruta); return isinstance(r, dict) and r.get("tipo") == TIPO
    except Exception:
        return False


def huella_receta(receta):
    """SHA-256 del JSON canónico (claves ordenadas, sin espacios, UTF-8) sin el campo «huella». Formato «sha256:<hex>»."""
    base = {k: v for k, v in receta.items() if k != "huella"}
    return "sha256:" + hashlib.sha256(json.dumps(base, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode("utf-8")).hexdigest()


def huella_efectiva(receta):
    """Huella con que se reconoce una receta aplicada (12.5): la del campo «huella» si la trae; si no, la calculada sin folio (el folio nunca se guarda)."""
    h = receta.get("huella")
    if h: h = str(h); return h if h.startswith("sha256:") else "sha256:" + h
    return huella_receta({k: v for k, v in receta.items() if k != "folio"})


def entero(v):
    """El valor si es un entero (no booleano); si no, None."""
    return v if isinstance(v, int) and not isinstance(v, bool) else None


def huella_sintetica(receta):
    """Huella del CONTENIDO de una receta, sin huella, fecha de generación ni folio: identifica las recetas sintéticas que distribuye el repositorio."""
    base = {k: v for k, v in receta.items() if k not in ("huella", "generado_en", "folio")}
    return "sha256:" + hashlib.sha256(json.dumps(base, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode("utf-8")).hexdigest()


def errores_modo(receta, modo):
    """E1-R2/R6: con un modo distinto de «demostracion» no se aceptan hojas sintéticas ni variables SINTETICO_DEMO."""
    err = []
    if modo == "demostracion": return err
    for h in receta.get("hojas") or []:
        if isinstance(h, dict) and h.get("origen_datos") == "sintetico":
            err.append(f"la hoja «{h.get('hoja')}» es sintética y el proyecto es real: en modo real no se cargan datos sintéticos (no se cargó nada)")
    for v in receta.get("variables") or []:
        if isinstance(v, dict) and v.get("estado_evidencia") == "SINTETICO_DEMO":
            err.append(f"la variable «{v.get('nombre')}» es SINTETICO_DEMO y el proyecto es real: en modo real no se cargan datos sintéticos (no se cargó nada)")
    return err


EXIGE_LIMITES_DESDE = "0.29"


def errores_limites(receta, modo):
    """Entrega 3, Decisión 5: disenos[].limites_seguros por factor ({"A": {"min": 0, "max": 1}}). Si vienen, deben ser números con mínimo menor que máximo para cada
    factor del diseño. Se exigen a los diseños de recetas con dmike_minimo 0.29 o mayor en modo real; en demostración son opcionales y las recetas anteriores
    se aceptan sin ellos (el diseño queda con región «no declarada»)."""
    err = []; exige = modo == "real" and _version(receta.get("dmike_minimo") or "0") >= _version(EXIGE_LIMITES_DESDE)
    for i, d in enumerate(receta.get("disenos") or [], 1):
        if not isinstance(d, dict) or d.get("modulo") != "DISENO-RSM": continue
        nom = d.get("id") or i; ls = d.get("limites_seguros")
        nombres = [x.strip() for x in str((d.get("parametros") or {}).get("nombres") or "").replace(";", ",").split(",") if x.strip()]
        if ls is None:
            if exige: err.append(f"el diseño «{nom}» no trae limites_seguros: en modo real son obligatorios desde D-MIKE {EXIGE_LIMITES_DESDE} (un mínimo y un máximo por factor)")
            continue
        if not isinstance(ls, dict): err.append(f"limites_seguros del diseño «{nom}» debe ser un objeto con un {{min, max}} por factor"); continue
        for n in nombres:
            v = ls.get(n)
            if not isinstance(v, dict) or any(isinstance(v.get(c), bool) or not isinstance(v.get(c), (int, float)) for c in ("min", "max")):
                err.append(f"limites_seguros del diseño «{nom}»: falta un {{min, max}} numérico para el factor «{n}»")
            elif v["min"] >= v["max"]: err.append(f"limites_seguros del diseño «{nom}»: el mínimo de «{n}» debe ser menor que su máximo")
    return err


def errores_v2(receta):
    """12 / E2-R1: estructura de variables, tareas, revision, revision_base y analisis[].estado / requiere. Las claves desconocidas se toleran."""
    err = []
    vs = receta.get("variables")
    if vs is not None and not isinstance(vs, list): err.append("variables debe ser una lista")
    for i, v in enumerate(vs if isinstance(vs, list) else [], 1):
        if not isinstance(v, dict): err.append(f"la variable {i} debe ser un objeto con nombre y estado_evidencia"); continue
        if not isinstance(v.get("nombre"), str) or not v["nombre"].strip(): err.append(f"la variable {i} no tiene nombre (debe ser un texto no vacío)")
        if v.get("ensayo_destructivo") is not None and not isinstance(v["ensayo_destructivo"], bool):
            err.append(f"ensayo_destructivo de la variable «{v.get('nombre')}» debe ser verdadero o falso")
    ts = receta.get("tareas")
    if ts is not None and not isinstance(ts, list): err.append("tareas debe ser una lista")
    vistos = set()
    for i, t in enumerate(ts if isinstance(ts, list) else [], 1):
        if not isinstance(t, dict): err.append(f"la tarea {i} debe ser un objeto con id y descripcion"); continue
        tid = t.get("id")
        if not isinstance(tid, str) or not tid.strip(): err.append(f"la tarea {i} no tiene id (debe ser un texto no vacío)")
        elif tid in vistos: err.append(f"el id de tarea «{tid}» está repetido")
        else: vistos.add(tid)
        if t.get("estado") is not None and t["estado"] not in ESTADOS_TAREA:
            err.append(f"estado «{t['estado']}» no válido en la tarea «{tid}» (abierta o cerrada)")
    rev, base = receta.get("revision"), receta.get("revision_base")
    if rev is not None and (entero(rev) is None or rev < 1): err.append("revision debe ser un entero mayor o igual a 1")
    if base is not None:
        if entero(base) is None: err.append("revision_base debe ser nula o un entero menor que revision")
        elif entero(rev) is None: err.append("revision_base solo tiene sentido con una revision válida")
        elif base >= rev: err.append(f"revision_base ({base}) debe ser menor que revision ({rev})")
    for i, a in enumerate(receta.get("analisis") or [], 1):
        if not isinstance(a, dict): continue
        nom = a.get("id") or a.get("modulo") or i
        if a.get("estado") is not None and a["estado"] not in ESTADOS_ANALISIS:
            err.append(f"estado «{a['estado']}» no válido en el análisis «{nom}» (pendiente, disponible, analizado o aprobado)")
        rq = a.get("requiere")
        if rq is not None and (not isinstance(rq, list) or any(not (isinstance(x, str) and x.startswith("@") and len(x) > 1) for x in rq)):
            err.append(f"requiere del análisis «{nom}» debe ser una lista de columnas con el prefijo @ (por ejemplo \"@Espesor\")")
    return err


def validar(receta, registro, version_programa, modo=None):
    """Errores que impiden ejecutar (lista vacía = válida). modo: el modo con que se ejecutará (por omisión, el de la receta)."""
    from . import modo as MD
    err = []
    if receta.get("tipo") != TIPO: err.append("no es una receta de D-MIKE")
    try: rv = int(receta.get("receta_version") or 0)
    except (TypeError, ValueError): rv = 0
    if rv > VERSION_SOPORTADA: err.append(f"versión de receta {receta.get('receta_version')} más nueva que la que entiende este programa (soporta hasta la {VERSION_SOPORTADA}); actualiza D-MIKE")
    if _version(receta.get("dmike_minimo") or "0") > _version(version_programa): err.append(f"requiere D-MIKE {receta.get('dmike_minimo')} o posterior")
    declarado = receta.get("modo_proyecto")
    if declarado is not None and declarado not in MD.MODOS: err.append(f"modo_proyecto «{declarado}» no válido (real o demostracion)")
    if rv >= 2 and declarado is None: err.append("la receta v2 debe declarar modo_proyecto (real o demostracion)")
    idp = receta.get("id_proyecto")
    if idp is not None and not (isinstance(idp, str) and _UUID4.fullmatch(idp)): err.append("id_proyecto no es un UUID v4 en minúsculas")
    if receta.get("huella") is not None:
        h = str(receta["huella"]); h = h if h.startswith("sha256:") else "sha256:" + h
        if h != huella_receta(receta): err.append("la huella no coincide: la receta fue alterada o está dañada")
    for h in receta.get("hojas") or []:
        if isinstance(h, dict) and h.get("origen_datos") is not None and h["origen_datos"] not in ORIGEN_DATOS:
            err.append(f"origen_datos «{h['origen_datos']}» no válido en la hoja «{h.get('hoja')}»")
    for v in receta.get("variables") if isinstance(receta.get("variables"), list) else []:
        if isinstance(v, dict) and v.get("estado_evidencia") is not None and v["estado_evidencia"] not in EVIDENCIA:
            err.append(f"estado_evidencia «{v['estado_evidencia']}» no válido en la variable «{v.get('nombre')}»")
    err += errores_v2(receta)
    modo_efectivo = modo or MD.detectar(receta, huella_sintetica(receta))[0]
    err += errores_modo(receta, modo_efectivo)
    err += errores_limites(receta, modo_efectivo)
    for paso in receta.get("disenos", []) + receta.get("analisis", []):
        if isinstance(paso, dict) and paso.get("modulo") not in registro: err.append(f"módulo desconocido: {paso.get('modulo')}")
    from ..modulos.formatos import FORMATOS
    for f in receta.get("formatos", []):
        if f.get("id") not in FORMATOS: err.append(f"formato desconocido: {f.get('id')}")
    return err


def guardable(receta, id_proyecto=None):
    """Registro «_receta» que se guarda en el proyecto. El folio de una receta v1 se ignora: no se escribe en ninguna parte nueva (E1-R8)."""
    sin = {k: v for k, v in receta.items() if k != "folio"}
    return {"ids": {}, "json": json.dumps(sin, ensure_ascii=False), "origen": receta.get("origen") or "", "id_proyecto": id_proyecto,
            "revision": entero(receta.get("revision")), "huella": huella_efectiva(receta), "aplicada_en": datetime.now().isoformat(timespec="seconds")}


def resumen(receta, datos):  # noqa: C901
    hojas = [(h["hoja"], len(datos.get(h["hoja"], [])) - 1 if datos.get(h["hoja"]) else 0) for h in receta.get("hojas", [])]
    return {"nombre": (receta.get("datos_fijos") or {}).get("nombre_proyecto") or "(sin nombre)", "origen": receta.get("origen", ""),
            "hojas": [{"hoja": h, "filas": max(n, 0)} for h, n in hojas], "disenos": len(receta.get("disenos", [])),
            "formatos": len(receta.get("formatos", [])), "analisis": len(receta.get("analisis", [])),
            "sin_datos": [h for h, n in hojas if n <= 0]}


def _tsv(filas):
    def v(x):
        if x is None: return ""
        if hasattr(x, "strftime"): return x.strftime("%d/%m/%Y")
        if isinstance(x, float) and x.is_integer(): return str(int(x))
        return str(x)
    return "\n".join("\t".join(v(x) for x in f) for f in filas)


def _columnas_citadas(params):
    out = []
    for v in params.values():
        for x in (v if isinstance(v, list) else [v]):
            if isinstance(x, str) and x.startswith("@"): out.append(x)
    return out


def _tiene_datos(a, nombre):
    j = a.hoja.indice(nombre)
    return j is not None and any(a.hoja.columnas[j].valor(i)[2] is not None for i in range(a.hoja.n))


def solo_datos(api, receta, datos, id_proyecto=None):
    """Carga las hojas de la plantilla como columnas, sin formatos ni análisis (para trabajar a mano). La receta queda guardada."""
    n = 0
    if receta.get("datos_fijos"): api.fijar_datos({k: v or "" for k, v in receta["datos_fijos"].items() if k != "logo"})
    for h in receta.get("hojas", []):
        filas = datos.get(h["hoja"])
        if not filas: continue
        cab = [str(x).strip() if x is not None else "" for x in filas[0]]
        api.pegar(-1, len(api.hoja.columnas), _tsv([cab] + filas[1:])); n += 1
    api.proyecto.formatos["_receta"] = guardable(receta, id_proyecto)
    return n


def _vacio(v): return v in (None, "", [], {})


def _parametros_con_limites(d):
    """Parámetros del diseño con su región autorizada (Entrega 3, Decisión 5). Si la receta trae limites_seguros por factor, se pasan como «limites»; si no, el diseño
    se crea sin región declarada (recetas anteriores a 0.29: compatibilidad) y queda marcado «región no declarada»."""
    p = dict(d.get("parametros") or {}); ls = d.get("limites_seguros")
    nombres = [x.strip() for x in str(p.get("nombres") or "").replace(";", ",").split(",") if x.strip()]
    if isinstance(ls, dict) and nombres and all(isinstance(ls.get(n), dict) and ls[n].get("min") is not None and ls[n].get("max") is not None for n in nombres):
        p["limites"] = "; ".join(f"{ls[n]['min']}, {ls[n]['max']}" for n in nombres)
    elif not p.get("limites"): p["limites_opcionales"] = True
    return p


def ejecutar(api, receta, datos, actualizar=False, id_proyecto=None):
    """Arma el proyecto en la sesión actual. Con actualizar=True agrega solo lo nuevo al proyecto abierto: no sobrescribe columnas
    con datos ni campos de formatos ya escritos, y no repite análisis ya corridos con esta receta (se reconocen por su id)."""
    from .perfil import validar as validar_perfil
    from .hoja import Hoja
    from . import modo as MD, seguimiento as SG
    malos = errores_modo(receta, MD.valor(api.proyecto))
    if malos: raise RecetaInvalida("; ".join(malos))
    log, avisos, pend, pend_t = [], [], [], []
    a = api

    def pendiente(texto, tipo, objeto, accion, liga):
        """El pendiente sigue saliendo como texto (mismo texto en el registro) y además queda como tarea auto con id estable (12.3)."""
        pend.append(texto); pend_t.append({"tipo": tipo, "objeto": objeto, "texto": texto, "accion": accion, "liga": liga})
    previo = a.proyecto.formatos.get("_receta") or {} if actualizar else {}
    hechos_antes = dict(previo.get("ids") or {})
    if receta.get("datos_fijos"):
        fijos = {k: v or "" for k, v in receta["datos_fijos"].items() if k != "logo"}
        if actualizar: fijos = {k: v for k, v in fijos.items() if not (a.proyecto.fijos.get(k) or {}).get("valor")}
        if fijos: a.fijar_datos(fijos); log.append("Datos fijos del proyecto.")
    if receta.get("perfil") and not (actualizar and a.proyecto.modulos):
        perfil, av = validar_perfil(json.dumps(receta["perfil"], ensure_ascii=False), a.catalogo)
        a.proyecto.aplicar_perfil(perfil, a.hoja); avisos += av
        log.append(f"Perfil aplicado: {sum(1 for m in a.proyecto.modulos.values() if m['estado'] == 'activo')} módulos activos.")
    if receta.get("hoja_sugerida") is False and not actualizar: a.hoja = Hoja()
    for mid, fase in receta.get("complementarios", []):
        if not (actualizar and a.proyecto.modulos.get(mid, {}).get("estado") == "activo"): a.proyecto.activar_modulo(mid, fase, a.catalogo)
    dis_hojas = {d.get("hoja") for d in receta.get("disenos", [])}
    ref = {}
    # 1) diseños de experimentos primero (así las columnas quedan donde el usuario las espera) de experimentos y sus respuestas
    for d in receta.get("disenos", []):
        if actualizar and (d.get("id") in hechos_antes or any(a.hoja.indice(x) is not None for x in d.get("respuestas", []))):
            if d.get("id") in hechos_antes: ref[d["id"]] = hechos_antes[d["id"]]
            info = {"factores": [{"nombre": f} for f in (d["parametros"].get("nombres") or "").replace(";", ",").split(",") if f.strip()]}
            for f in info["factores"]: f["nombre"] = f["nombre"].strip()
            if not info["factores"] or a.hoja.indice(info["factores"][0]["nombre"]) is None:
                log.append(f"Diseño «{d.get('id')}»: ya estaba en el proyecto."); continue
            log.append(f"Diseño «{d.get('id')}»: ya estaba en el proyecto; se revisan sus respuestas.")
        else:
            r = a.ejecutar_modulo(d["modulo"], _parametros_con_limites(d))
            if r["error"]: avisos.append(f"Diseño «{d.get('id')}»: {r['sesion'][0]['texto']}"); continue
            ref[d["id"]] = r["resultado"]["uid"]; info = r["resultado"]["datos"]["diseno"]
        filas = datos.get(d.get("hoja") or "")
        factores = [f["nombre"] for f in info["factores"]]
        jf = [a.hoja.indice(n) for n in factores]
        filas_dis = [i for i in range(a.hoja.n) if a.hoja.columnas[jf[0]].valor(i)[2] is not None]
        if not filas or len(filas) <= 1:
            pendiente(f"Correr el diseño «{d.get('id')}» en el orden de la hoja y capturar {', '.join(d.get('respuestas', []))}.", "diseno", d.get("id"),
                      "Correr el diseño en el orden de la hoja y capturar sus respuestas", f"diseno:{d.get('id')}"); log.append(f"Diseño «{d.get('id')}» creado; respuestas por capturar.")
            continue
        cab = [str(x).strip() if x is not None else "" for x in filas[0]]; cuerpo = filas[1:]
        ok = len(cuerpo) == len(filas_dis)
        for q, n in enumerate(factores):
            if not ok: break
            if n in cab:
                c = cab.index(n)
                for k, i in enumerate(filas_dis):
                    v = a.hoja.columnas[jf[q]].valor(i)[2]
                    if v is None or cuerpo[k][c] is None or abs(float(cuerpo[k][c]) - float(v)) > 1e-6 * max(1, abs(float(v))): ok = False; break
        if not ok:
            avisos.append(f"Diseño «{d.get('id')}»: las corridas de la plantilla no coinciden con el diseño creado; captura las respuestas a mano."); continue
        n_resp = 0
        for resp in d.get("respuestas", []):
            if resp not in cab: pendiente(f"Capturar {resp} del diseño «{d.get('id')}».", "respuesta", f"{d.get('id')}:{resp}", "Capturar la respuesta del diseño", f"diseno:{d.get('id')}"); continue
            c = cab.index(resp)
            if a.hoja.indice(resp) is None: a.pegar(-1, len(a.hoja.columnas), resp)
            j = a.hoja.indice(resp); n = 0
            if actualizar and _tiene_datos(a, resp): continue
            for k, i in enumerate(filas_dis):
                v = cuerpo[k][c]
                if v not in (None, ""): a.hoja.editar(i, j, str(v)); n += 1
            n_resp += n
            if n < len(filas_dis): pendiente(f"Diseño «{d.get('id')}»: faltan {len(filas_dis) - n} valores de {resp}.", "respuesta", f"{d.get('id')}:{resp}", "Capturar la respuesta del diseño", f"diseno:{d.get('id')}")
        log.append(f"Diseño «{d.get('id')}»: {len(filas_dis)} corridas" + (" con sus respuestas." if n_resp else "; respuestas por capturar."))
    # 2) datos capturados por el usuario (por nombre de columna)
    for h in receta.get("hojas", []):
        if h["hoja"] in dis_hojas: continue
        nombres = [c["nombre"] for c in h["columnas"]]
        filas = datos.get(h["hoja"])
        if actualizar and any(a.hoja.indice(n) is not None for n in nombres):
            cab = [str(x).strip() if x is not None else "" for x in filas[0]] if filas else []
            puestas = 0
            for n in nombres:
                j = a.hoja.indice(n); col = [f[cab.index(n)] if n in cab and cab.index(n) < len(f) else None for f in (filas[1:] if filas else [])]
                hay = any(v not in (None, "") for v in col)
                if j is None: a.pegar(-1, len(a.hoja.columnas), _tsv([[n]] + [[v] for v in col])); puestas += hay
                elif not _tiene_datos(a, n) and hay: a.pegar(-1, j, _tsv([[n]] + [[v] for v in col])); puestas += 1
                elif hay and _tiene_datos(a, n): avisos.append(f"La columna {n} ya tenía datos: no se sobrescribió.")
            log.append(f"Hoja «{h['hoja']}»: {puestas} columnas con datos nuevos."); continue
        if filas:
            cab = [str(x).strip() if x is not None else "" for x in filas[0]]
            faltan = [n for n in nombres if n not in cab]
            if faltan: avisos.append(f"Hoja «{h['hoja']}»: faltan las columnas {', '.join(faltan)}; se crean vacías.")
            idx = [cab.index(n) if n in cab else None for n in nombres]
            cuerpo = [[(f[i] if i is not None and i < len(f) else None) for i in idx] for f in filas[1:]]
            a.pegar(-1, len(a.hoja.columnas), _tsv([nombres] + cuerpo))
            log.append(f"Hoja «{h['hoja']}»: {len(cuerpo)} filas en {len(nombres)} columnas.")
            if not cuerpo: pendiente(f"Llenar la hoja «{h['hoja']}» ({', '.join(nombres)}).", "hoja", h["hoja"], "Llenar la hoja con las mediciones", f"hoja:{h['hoja']}")
        else:
            a.pegar(-1, len(a.hoja.columnas), "\t".join(nombres)); pendiente(f"Llenar la hoja «{h['hoja']}» ({', '.join(nombres)}).", "hoja", h["hoja"], "Llenar la hoja con las mediciones", f"hoja:{h['hoja']}")
    # 3) formatos
    for f in receta.get("formatos", []):
        base = a.abrir_formato(f["id"])["datos"]
        if actualizar:
            for k, v in (f.get("cabecera") or {}).items():
                if _vacio(base["cabecera"].get(k)): base["cabecera"][k] = v
                elif not _vacio(v) and str(base["cabecera"].get(k)) != str(v): avisos.append(f"{f['id']}: «{k}» ya estaba escrito; se conservó.")
            for t, filas in (f.get("tablas") or {}).items():
                actual = base["tablas"].get(t) or []
                if not any(any(not _vacio(v) for v in x.values()) for x in actual): base["tablas"][t] = [dict(x) for x in filas]
                elif [dict(x) for x in actual if any(not _vacio(v) for v in x.values())] != [dict(x) for x in filas if any(not _vacio(v) for v in x.values())]:
                    avisos.append(f"{f['id']}: la tabla «{t}» ya tenía renglones; se conservó.")
        else:
            base["cabecera"].update(f.get("cabecera") or {})
            for t, filas in (f.get("tablas") or {}).items(): base["tablas"][t] = [dict(x) for x in filas]
        r = a.guardar_formato(f["id"], base)
        if r["error"]: avisos.append(f"Formato {f['id']}: {r['sesion'][0]['texto']}")
    if receta.get("formatos"): log.append(f"Formatos llenados: {len(receta['formatos'])}.")
    # 4) análisis
    hechos = 0
    def resolver(v):
        if isinstance(v, str) and v.startswith("@"):
            j = a.hoja.indice(v[1:])
            if j is None: raise KeyError(v[1:])
            return j
        if isinstance(v, str) and v.startswith("#"):
            if v[1:] not in ref: raise KeyError(v)
            return ref[v[1:]]
        if isinstance(v, list): return [resolver(x) for x in v]
        return v
    for k, v in hechos_antes.items(): ref.setdefault(k, v)
    for p in receta.get("analisis", []):
        aid, req = p.get("id") or p["modulo"], SG.requiere_de(p)
        if actualizar and p.get("id") in hechos_antes:
            SG.registrar_analisis(a.proyecto, aid, p["modulo"], "analizado", req, "", hechos_antes[p["id"]]); continue

        def sin_correr(texto, motivo, bloqueado, como_pendiente=True):
            """12.4: el análisis queda «pendiente» con su motivo; no se corre sin los datos mínimos."""
            if como_pendiente: pendiente(texto, "analisis", aid, "Reunir los datos y correr el análisis", f"analisis:{aid}")
            else: avisos.append(texto)
            SG.registrar_analisis(a.proyecto, aid, p["modulo"], "pendiente", req, motivo, None, bloqueado)
        try: params = {k: resolver(v) for k, v in p["parametros"].items()}
        except KeyError as e:
            x = e.args[0]
            sin_correr(f"Análisis «{p.get('id')}» ({p['modulo']}): falta {x}.", f"Falta el resultado del análisis «{x[1:]}»." if x.startswith("#") else (SG.motivo_columnas(*SG.sin_datos(a.hoja, req)) or f"Falta la columna {x}."), x.startswith("#")); continue
        vacias = [v[1:] for v in _columnas_citadas(p["parametros"]) if not _tiene_datos(a, v[1:])]
        if vacias: sin_correr(f"Análisis «{p.get('id')}» ({p['modulo']}): sin datos en {', '.join(vacias)}.", SG.motivo_columnas(*SG.sin_datos(a.hoja, req)), False); continue
        faltan, vacias = SG.sin_datos(a.hoja, req)          # columnas que la receta declara en «requiere» aunque los parámetros no las citen
        if faltan or vacias:
            m = SG.motivo_columnas(faltan, vacias); sin_correr(f"Análisis «{p.get('id')}» ({p['modulo']}): {m[0].lower() + m[1:]}", m, False); continue
        try: r = a.ejecutar_modulo(p["modulo"], params)
        except Exception as e: sin_correr(f"Análisis «{p.get('id')}»: no se pudo correr ({e}).", f"No se pudo correr: {e}", True); continue
        if r["error"]:
            msg = r["sesion"][0]["texto"]; sin_correr(f"Análisis «{p.get('id')}»: {msg}", msg, True, p.get("si_faltan_datos", "dejar_pendiente") == "dejar_pendiente"); continue
        ref[aid] = r["resultado"]["uid"]; hechos += 1
        SG.registrar_analisis(a.proyecto, aid, p["modulo"], "analizado", req, "", ref[aid])
    log.append(f"Análisis corridos: {hechos} de {len(receta.get('analisis', []))}.")
    ORDEN = {"pendiente": 0, "en_curso": 1, "hecho": 2}
    for mid, av in (receta.get("avance") or {}).items():
        actual = (a.proyecto.modulos.get(mid) or {}).get("avance", "pendiente")
        if actualizar and ORDEN.get(actual, 0) >= ORDEN.get(av, 0): continue
        try: a.proyecto.marcar_avance(mid, av)
        except ValueError: pass
    # seguimiento (0.27, secciones 12.2 a 12.5): variables, tareas y revisión
    avisos += SG.fusionar_variables(a.proyecto, receta.get("variables"))
    SG.fusionar_tareas_receta(a.proyecto, receta.get("tareas"))
    SG.registrar_pendientes(a.proyecto, pend_t)
    guardada = a.proyecto.formatos.get("_receta") if isinstance(a.proyecto.formatos.get("_receta"), dict) else {}
    mismo = SG.aplicada(a.proyecto)["huella"] == huella_efectiva(receta)
    avisos += SG.archivar_receta_previa(a.proyecto, receta)         # la receta anterior pasa a «_revisiones» antes de reemplazarse
    nuevo = guardable(receta, id_proyecto if id_proyecto is not None else previo.get("id_proyecto")); nuevo["ids"] = ref
    if mismo and guardada.get("aplicada_en"): nuevo["aplicada_en"] = guardada["aplicada_en"]
    if "folio" in previo: nuevo["folio"] = previo["folio"]          # E1-R8: el folio que ya traía un .dmike anterior se conserva intacto
    a.proyecto.formatos["_receta"] = nuevo
    try: anterior = json.loads(guardada.get("json") or "{}")
    except (ValueError, TypeError): anterior = {}
    cerradas = SG.cerrar_cumplidas(a.proyecto, a.hoja, receta, anterior)
    return {"log": log, "avisos": avisos, "pendientes": pend, "tareas_cerradas": cerradas}
